"""Read a final workbook by configuration; never calculate or trust XLSX caches."""
from __future__ import annotations

import hashlib
import re
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string

from .contracts import ValidationError, decimal_text, nonempty, require, schema_fields, validate_collection


def normalized(value) -> str:
    return " ".join(str(value or "").split())


def column(value) -> int:
    require(isinstance(value, str) and bool(re.fullmatch(r"[A-Z]{1,3}", value)), "Invalid column")
    number = column_index_from_string(value)
    require(number <= 16384, "Column exceeds Excel limit")
    return number


def positive_row(value) -> int:
    require(type(value) is int and 1 <= value <= 1048576, "Invalid Excel row")
    return value


def snapshot_cells(snapshot: dict | None, digest: str) -> dict:
    if snapshot is None:
        return {}
    require(isinstance(snapshot, dict) and snapshot.get("status") == "succeeded",
            "Recalculation did not succeed")
    require(snapshot.get("resultSha256") == digest, "Recalculation snapshot is stale")
    require(nonempty(snapshot.get("engine")), "Recalculation engine is required")
    require(isinstance(snapshot.get("cells"), list), "Snapshot cells are required")
    cells = {}
    for item in snapshot["cells"]:
        require(isinstance(item, dict) and nonempty(item.get("sheet")), "Invalid snapshot cell")
        ref = item.get("cell")
        require(isinstance(ref, str) and bool(re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", ref)),
                "Invalid snapshot cell reference")
        require(item.get("type") in {"number", "string", "blank", "error"} and "value" in item,
                "Snapshot cell type and value are required")
        require((item["sheet"], ref) not in cells, "Duplicate snapshot cell")
        if item["type"] == "number":
            decimal_text(item["value"])
        elif item["type"] == "blank":
            require(item["value"] is None, "Blank snapshot cell must be null")
        else:
            require(isinstance(item["value"], str), "Snapshot text/error must be a string")
        cells[item["sheet"], ref] = item
    return cells


def read_value(cell, sheet: str, snapshot: dict):
    if cell.data_type == "f":
        value = snapshot.get((sheet, cell.coordinate))
        if value is None:
            return None, "FORMULA_RESULT_UNAVAILABLE"
        if value["type"] == "error":
            return None, "FORMULA_ERROR"
        return value["value"], None
    if cell.data_type == "e":
        return None, "CELL_ERROR"
    return cell.value, None


def read_code_collection(workbook_path: str | Path, schema: dict, binding: dict, manifest: dict,
                         *, calculation_snapshot: dict | None = None, allow_draft: bool = False) -> dict:
    """Manifest/snapshot must come from trusted server orchestration, not model output.

    Caller supplies explicit employee rows and identity expectations. Template ancestry
    is asserted by that manifest; this reader cannot infer ancestry from XLSX bytes.
    This offline API never authorizes submission or declares a Code-vs-AI PASS.
    """
    fields = schema_fields(schema)
    require(isinstance(binding, dict) and isinstance(manifest, dict), "Binding and manifest must be objects")
    require(binding.get("formatVersion") == 0, "Unsupported binding formatVersion")
    for key in ("schemaId", "schemaVersion"):
        require(binding.get(key) == schema[key], "Binding " + key + " mismatch")
    for config in (schema, binding):
        require(config.get("status") in {"draft_not_enabled", "enabled"}, "Unsupported configuration status")
        require(config["status"] == "enabled" or allow_draft, "Draft configuration requires allow_draft")
        require(config.get("automaticPassEnabled") is False, "Automatic PASS is not supported")
    rules = binding.get("fields")
    require(isinstance(rules, list) and all(isinstance(r, dict) for r in rules), "Binding fields are required")
    names = [r.get("field") for r in rules]
    require(all(isinstance(n, str) for n in names) and len(names) == len(set(names))
            and set(names) == set(fields), "Binding fields must match schema exactly")
    raw = Path(workbook_path).read_bytes()
    digest = hashlib.sha256(raw).hexdigest()
    require(manifest.get("resultSha256") == digest, "Result manifest is stale")
    template = binding.get("template")
    require(isinstance(template, dict) and bool(re.fullmatch(r"[0-9a-f]{64}", str(template.get("sha256", "")))),
            "Binding template SHA-256 is required")
    require(manifest.get("templateSha256") == template["sha256"], "Template binding mismatch")
    require(nonempty(manifest.get("period")), "Manifest period is required")
    require(bool(re.fullmatch(r"[A-Z]{3}", str(manifest.get("currency", "")))), "Manifest currency is required")
    employees = manifest.get("employees")
    require(isinstance(employees, list) and bool(employees), "Explicit employee manifest is required")
    row_ids, entity_ids = set(), set()
    header_row = positive_row(binding.get("headerRow"))
    for employee in employees:
        require(isinstance(employee, dict), "Invalid employee manifest")
        row = positive_row(employee.get("row"))
        require(row > header_row and row not in row_ids, "Duplicate or invalid employee row")
        require(nonempty(employee.get("entityKey")) and employee["entityKey"] not in entity_ids,
                "Duplicate or missing entity key")
        require(nonempty(employee.get("expectedName")), "Expected employee name is required")
        row_ids.add(row)
        entity_ids.add(employee["entityKey"])
    snapshot = snapshot_cells(calculation_snapshot, digest)
    entity = binding.get("entity")
    require(isinstance(entity, dict), "Entity binding is required")
    name_column = column(entity.get("latinNameColumn"))
    sheet_name = binding.get("resultSheet")
    require(nonempty(sheet_name), "Result sheet is required")
    # Work from immutable bytes, not a path which can change between reads.
    wb = load_workbook(BytesIO(raw), data_only=False, read_only=False, keep_links=False)
    try:
        require(sheet_name in wb.sheetnames, "Result sheet is missing")
        ws = wb[sheet_name]
        require(max(row_ids) <= ws.max_row, "Employee row exceeds worksheet bounds")
        headers = {}
        for cell in ws[header_row]:
            if cell.value is not None:
                headers.setdefault(normalized(cell.value), []).append(cell.column)
        resolved = {}
        for rule in rules:
            require(nonempty(rule.get("expectedHeader")), "Expected field header is required")
            matches = headers.get(normalized(rule["expectedHeader"]), [])
            require(len(matches) == 1, "Missing or ambiguous header: " + rule["field"])
            require(matches[0] not in resolved.values(), "Multiple fields bind to the same column")
            # Column hints are observations, not the lookup key; unique headers may move.
            resolved[rule["field"]] = matches[0]
        result = {"engine": "code", "schemaId": schema["schemaId"], "schemaVersion": schema["schemaVersion"],
                  "resultSha256": digest, "templateSha256": template["sha256"],
                  "period": manifest["period"], "currency": manifest["currency"],
                  "records": [], "issues": [], "automaticPassEnabled": False}
        seen_names = set()
        def issue(code, employee, cell, field=None):
            item = {"code": code, "entityKey": employee["entityKey"], "sheet": sheet_name, "cell": cell}
            if field is not None:
                item["field"] = field
            result["issues"].append(item)
        for employee in employees:
            row = employee["row"]
            name_cell = ws.cell(row, name_column)
            name, error = read_value(name_cell, sheet_name, snapshot)
            if error or not nonempty(name):
                issue(error or "IDENTITY_MISSING", employee, name_cell.coordinate)
            else:
                key = normalized(name).casefold()
                if key != normalized(employee["expectedName"]).casefold():
                    issue("IDENTITY_MISMATCH", employee, name_cell.coordinate)
                if key in seen_names:
                    issue("AMBIGUOUS_IDENTITY", employee, name_cell.coordinate)
                seen_names.add(key)
            record = {"entityKey": employee["entityKey"], "observedName": name if isinstance(name, str) else None,
                      "row": row, "fields": []}
            for rule in rules:
                field = rule["field"]
                cell = ws.cell(row, resolved[field])
                value, error = read_value(cell, sheet_name, snapshot)
                state, text_value = "found", None
                if error:
                    state = "unreadable"
                elif value is None or value == "":
                    state, error = "missing", "FIELD_MISSING"
                else:
                    try:
                        text_value = decimal_text(value)
                    except ValidationError:
                        state, error = "unreadable", "INVALID_DECIMAL"
                if error:
                    issue(error, employee, cell.coordinate, field)
                record["fields"].append({"field": field, "status": state, "value": text_value,
                                         "currency": manifest["currency"],
                                         "source": {"fileSha256": digest, "sheet": sheet_name, "cell": cell.coordinate}})
            result["records"].append(record)
        result["readiness"] = "review_required" if result["issues"] else "ready_for_comparison"
        validate_collection(result, schema)
        return result
    finally:
        wb.close()
