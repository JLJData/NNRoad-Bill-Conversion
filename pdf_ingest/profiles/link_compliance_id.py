# -*- coding: utf-8 -*-
"""
Link Compliance（Indonesia）Tax Invoice PDF

主转换请用工资明细 Excel（indonesia_payroll_calc）。
本 PDF 仅作旁路：抽取 Cash Advance → 写入 PN（见 vendor_plugins.link_compliance_cash_advance）。

vendor-to-source：
- convert_excels / convert_sources：工资明细 Excel 原样（或已是 Indonesia-L）作为引擎输入
- Tax Invoice PDF：拆入 artifact_facts，不写员工表
"""
from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any

from bill_convert.vendor_plugins.link_compliance_cash_advance import (
    looks_like_link_compliance_invoice,
    parse_cash_advance_pdf,
)

_EXCEL_SUFFIXES = (".xlsx", ".xlsm", ".xls")
PROFILE_ID = "link_compliance_id"


def convert_pdf(
    pdf_path: Path,
    output_path: Path,
    *,
    template_path: Path | None = None,
    pn_meta: Any = None,
    registry_dir: Path | None = None,
    fill_fx: bool = True,
    convert_mapping: dict | None = None,
) -> dict[str, Any]:
    raise ValueError(
        "Link Compliance 主账单请上传工资明细 Excel；"
        "Tax Invoice PDF 仅用于抽取 Cash Advance（作旁路附件 / artifactBatch），"
        "勿单独走 pdf-to-source"
    )


def convert_pdfs(
    pdf_paths: list[Path],
    output_path: Path,
    **kwargs: Any,
) -> dict[str, Any]:
    return convert_pdf(pdf_paths[0] if pdf_paths else Path("."), output_path, **kwargs)


def _looks_like_indonesia_l_workbook(path: Path) -> bool:
    from openpyxl import load_workbook
    from profiles.indonesia_payroll_calc.convert import looks_like_indonesia_l

    wb = load_workbook(path, data_only=False, read_only=True)
    try:
        if "Indonesia-L" not in wb.sheetnames:
            return False
        # read_only 下 looks_like 需要可随机访问；改用非 read_only
    finally:
        wb.close()
    wb2 = load_workbook(path, data_only=False)
    try:
        return looks_like_indonesia_l(wb2["Indonesia-L"])
    finally:
        wb2.close()


def convert_excels(
    excel_paths: list[Path],
    output_path: Path,
    *,
    template_path: Path | None = None,
    pn_meta: Any = None,
    registry_dir: Path | None = None,
    fill_fx: bool = True,
    convert_mapping: dict | None = None,
) -> dict[str, Any]:
    """工资明细 Excel（或已是 Indonesia-L）→ 原样输出，供 indonesia_payroll_calc 消费。"""
    del template_path, pn_meta, registry_dir, fill_fx
    from convert_mapping import resolve_convert_mapping
    from profiles.indonesia_payroll_calc import convert as id_mod

    paths = [Path(p).resolve() for p in excel_paths]
    if not paths:
        raise ValueError("未提供 Excel")
    for p in paths:
        if not p.is_file():
            raise FileNotFoundError(f"Excel 不存在: {p}")
    if len(paths) > 1:
        raise ValueError("Link Compliance 请一次只传一份工资明细 Excel")

    src = paths[0]
    output_path = Path(output_path).resolve()
    mapping_in = dict(convert_mapping) if isinstance(convert_mapping, dict) else {}
    mapping_in.setdefault("pdfProfileId", PROFILE_ID)
    mapping = resolve_convert_mapping("indonesia_payroll_calc", mapping_in)

    warnings: list[str] = []
    id_mod._ACTIVE_MAPPING = mapping
    try:
        employees = id_mod.parse_source_workbook(src, warnings)
    finally:
        id_mod._ACTIVE_MAPPING = None

    if not employees:
        raise ValueError(
            f"未能从「{src.name}」解析到员工行；请确认是 Link Compliance 工资明细或 Indonesia-L"
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, output_path)
    kind = "excel_indonesia_l" if _looks_like_indonesia_l_workbook(src) else "excel_payroll"
    if kind == "excel_indonesia_l":
        warnings.append("源表已是 Indonesia-L，已原样用作转换输入")
    else:
        warnings.append("已采用工资明细 Excel 作为转换主源（Tax Invoice PDF 请作旁路）")

    return {
        "ok": True,
        "profile_id": PROFILE_ID,
        "region": "Indonesia",
        "source_kind": kind,
        "output": str(output_path),
        "employee_count": len(employees),
        "parsed": [
            {
                "employee_name": e.get("Name of Employee") or e.get("Employee Name"),
                "base_salary": e.get("Base Salary") or e.get("BASIC PAY (IDR)"),
            }
            for e in employees
        ],
        "warnings": warnings,
        "fx_rate": None,
        "pn_meta": None,
    }


def convert_sources(
    source_paths: list[Path],
    output_path: Path,
    *,
    template_path: Path | None = None,
    pn_meta: Any = None,
    registry_dir: Path | None = None,
    fill_fx: bool = True,
    convert_mapping: dict | None = None,
) -> dict[str, Any]:
    """Excel 主源 → 引擎输入；Tax Invoice PDF → Cash Advance artifact_facts。"""
    from bill_convert.vendor_plugins.runtime import parse_artifact_facts, split_main_and_artifacts

    paths = [Path(p).resolve() for p in source_paths]
    if not paths:
        raise ValueError("未提供源文件")

    main_paths, artifact_paths, split_warnings = split_main_and_artifacts(
        paths, pdf_profile_id=PROFILE_ID
    )
    excels = [p for p in main_paths if p.suffix.lower() in _EXCEL_SUFFIXES]
    leftover_pdfs = [p for p in main_paths if p.suffix.lower() == ".pdf"]
    other = [p for p in main_paths if p not in excels and p not in leftover_pdfs]
    if other:
        raise ValueError(f"不支持的文件类型: {[p.name for p in other]}")

    rescued: list[Path] = []
    still: list[Path] = []
    for p in leftover_pdfs:
        try:
            if looks_like_link_compliance_invoice(p):
                rescued.append(p)
                continue
        except Exception:
            pass
        still.append(p)
    artifact_paths = list(artifact_paths) + rescued
    if still:
        raise ValueError(
            "link_compliance_id 主源仅支持工资明细 Excel；"
            f"无法识别的 PDF: {[p.name for p in still]}。"
            "Tax Invoice 请标为旁路附件，或与 Excel 同批上传。"
        )
    if not excels:
        raise ValueError(
            "请至少上传一份 Link Compliance 工资明细 Excel。"
            "Tax Invoice PDF 仅用于抽取 Cash Advance，不能单独转换。"
        )

    artifact_facts, artifact_warnings = parse_artifact_facts(
        artifact_paths, pdf_profile_id=PROFILE_ID
    )
    # 插件未命中时仍尝试直接解析（同批 CONVERT 的发票）
    if not artifact_facts and artifact_paths:
        for p in artifact_paths:
            try:
                parsed = parse_cash_advance_pdf(p) or {}
            except Exception:
                parsed = {}
            if parsed:
                artifact_facts.update(parsed)
                break

    result = convert_excels(
        excels,
        output_path,
        template_path=template_path,
        pn_meta=pn_meta,
        registry_dir=registry_dir,
        fill_fx=fill_fx,
        convert_mapping=convert_mapping,
    )
    warnings = list(result.get("warnings") or [])
    warnings.extend(split_warnings)
    warnings.extend(artifact_warnings)
    if artifact_paths and not artifact_facts:
        warnings.append("已识别 Tax Invoice PDF 但未解析到 Cash Advance")
    elif artifact_facts:
        amt = artifact_facts.get("link_compliance.cash_advance.amount_idr")
        name = artifact_facts.get("link_compliance.cash_advance.employee_name")
        warnings.append(f"已解析 Cash Advance：{name or '?'} IDR={amt}")

    result["warnings"] = warnings
    result["artifact_facts"] = artifact_facts or {}
    return result
