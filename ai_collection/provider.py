"""Provider abstraction and the OpenAI Responses API adapter."""
from __future__ import annotations

import base64
import copy
import hashlib
import json
import mimetypes
import os
import re
import ssl
import time
import urllib.error
import urllib.request
from abc import ABC, abstractmethod
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from bill_validation.contracts import ValidationError, nonempty, require

from .contracts import validate_ai_collection, validate_collection_request


class AIProviderError(RuntimeError):
    """Provider transport/model failure, distinct from business disagreement."""


class AIProvider(ABC):
    @property
    @abstractmethod
    def provider_id(self) -> str:
        raise NotImplementedError

    @abstractmethod
    def collect(self, request: dict) -> dict:
        raise NotImplementedError


class MockAIProvider(AIProvider):
    """Deterministic test provider. It performs no network or file access."""

    def __init__(self, response: dict | Callable[[dict], dict], provider_id: str = "mock"):
        require(nonempty(provider_id), "Mock provider id is required")
        self._response = response
        self._provider_id = provider_id
        self.calls = []

    @property
    def provider_id(self) -> str:
        return self._provider_id

    def collect(self, request: dict) -> dict:
        validate_collection_request(request)
        safe_request = copy.deepcopy(request)
        self.calls.append(safe_request)
        try:
            result = self._response(copy.deepcopy(safe_request)) if callable(self._response) else self._response
            return copy.deepcopy(result)
        except AIProviderError:
            raise
        except Exception as exc:
            raise AIProviderError(str(exc)) from exc


def _strict_object(properties: dict, required: list[str] | None = None) -> dict:
    return {
        "type": "object",
        "properties": properties,
        "required": required if required is not None else list(properties),
        "additionalProperties": False,
    }


def _collection_output_schema(request: dict, model: str) -> dict:
    field_names = [item["field"] for item in request["schema"]["fields"]]
    file_ids = [item["fileId"] for item in request["documents"]]
    source = _strict_object({
        "fileId": {"type": "string", "enum": file_ids},
        "location": {"type": "string"},
        "page": {"type": ["integer", "null"], "minimum": 1},
        "rawText": {"type": "string"},
    })
    field = _strict_object({
        "field": {"type": "string", "enum": field_names},
        "status": {"type": "string", "enum": ["found", "missing", "unreadable", "not_applicable"]},
        "value": {"type": ["string", "null"]},
        "currency": {"type": "string", "enum": [request["currency"]]},
        "confidence": {"type": "number", "minimum": 0, "maximum": 1},
        "source": source,
    })
    record = _strict_object({
        "sourceEntity": _strict_object({
            "employeeId": {"type": ["string", "null"]},
            "name": {"type": "string"},
        }),
        "fields": {"type": "array", "items": field},
    })
    unknown = _strict_object({
        "label": {"type": "string"},
        "value": {"type": ["string", "null"]},
        "currency": {"type": ["string", "null"]},
        "source": source,
    })
    issue = _strict_object({"code": {"type": "string"}, "message": {"type": "string"}})
    return _strict_object({
        "engine": {"type": "string", "enum": ["ai"]},
        "provider": {"type": "string", "enum": ["openai"]},
        "model": {"type": "string", "enum": [model]},
        "runId": {"type": "string", "enum": [request["runId"]]},
        "schemaId": {"type": "string", "enum": [request["schema"]["schemaId"]]},
        "schemaVersion": {"type": "string", "enum": [request["schema"]["schemaVersion"]]},
        "period": {"type": "string", "enum": [request["period"]]},
        "currency": {"type": "string", "enum": [request["currency"]]},
        "automaticPassEnabled": {"type": "boolean", "enum": [False]},
        "records": {"type": "array", "items": record},
        "unknownItems": {"type": "array", "items": unknown},
        "issues": {"type": "array", "items": issue},
        "readiness": {"type": "string", "enum": ["ready_for_association", "review_required"]},
    })


def _template_plan_output_schema(template_manifest: dict, collection: dict, schema: dict) -> dict:
    field_names = [item["field"] for item in schema["fields"]]
    write = _strict_object({
        "targetCell": {"type": "string", "pattern": "^[A-Z]{1,3}[1-9][0-9]*$"},
        "recordIndex": {"type": "integer", "minimum": 0, "maximum": max(0, len(collection["records"]) - 1)},
        "sourceKind": {"type": "string", "enum": ["entity_name", "field"]},
        "field": {"type": ["string", "null"], "enum": field_names + [None]},
    })
    issue = _strict_object({"code": {"type": "string"}, "message": {"type": "string"}})
    return _strict_object({
        "planVersion": {"type": "integer", "enum": [1]},
        "templateSha256": {"type": "string", "enum": [template_manifest["templateSha256"]]},
        "sheetName": {"type": "string", "enum": [template_manifest["sheetName"]]},
        "automaticWriteEnabled": {"type": "boolean", "enum": [False]},
        "writes": {"type": "array", "items": write},
        "issues": {"type": "array", "items": issue},
    })


def _dynamic_template_plan_output_schema(template_manifest: dict, documents: list[dict], model: str) -> dict:
    """Schema for template-driven extraction without a predeclared payroll field list."""
    file_ids = [item["fileId"] for item in documents]
    source = _strict_object({
        "fileId": {"type": "string", "enum": file_ids},
        "location": {"type": "string"},
        "page": {"type": ["integer", "null"], "minimum": 1},
        "rawText": {"type": "string"},
    })
    write = _strict_object({
        "targetCell": {"type": "string", "pattern": "^[A-Z]{1,3}[1-9][0-9]*$"},
        "semanticLabel": {"type": "string"},
        "sourceLabel": {"type": "string"},
        "valueType": {"type": "string", "enum": ["text", "decimal", "date"]},
        "value": {"type": "string"},
        "source": source,
        "confidence": {"type": "number", "minimum": 0, "maximum": 1},
    })
    issue = _strict_object({"code": {"type": "string"}, "message": {"type": "string"}})
    return _strict_object({
        "planVersion": {"type": "integer", "enum": [3]},
        "templateSha256": {"type": "string", "enum": [template_manifest["templateSha256"]]},
        "sheetName": {"type": "string", "enum": [template_manifest["sheetName"]]},
        "model": {"type": "string", "enum": [model]},
        "automaticWriteEnabled": {"type": "boolean", "enum": [False]},
        "writes": {"type": "array", "items": write},
        "issues": {"type": "array", "items": issue},
    })


def _openai_ssl_context() -> ssl.SSLContext:
    """Build a Windows-safe TLS client context (system CA path is often missing)."""
    try:
        import certifi
        context = ssl.create_default_context(cafile=certifi.where())
    except Exception:
        context = ssl.create_default_context()
    # Prefer modern TLS; some VPN/proxy middleboxes reset weaker handshakes mid-flight.
    if hasattr(ssl, "TLSVersion"):
        context.minimum_version = ssl.TLSVersion.TLSv1_2
    return context


def _is_transient_openai_transport_error(exc: BaseException) -> bool:
    text = str(exc).casefold()
    markers = (
        "eof occurred in violation of protocol",
        "connection reset",
        "connection aborted",
        "broken pipe",
        "timed out",
        "temporarily unavailable",
        "remote end closed connection",
        "server disconnected",
        "disconnected without sending",
        "ssl",
        "wrong version number",
    )
    return any(item in text for item in markers)


def _default_http_post(url: str, headers: dict[str, str], payload: dict, timeout_seconds: float) -> dict:
    body = json.dumps(payload).encode("utf-8")
    ssl_context = _openai_ssl_context()
    last_error: BaseException | None = None
    attempts = 3
    print(
        f"[ai] calling OpenAI ({len(body)} bytes, timeout={timeout_seconds}s)",
        flush=True,
    )
    for attempt in range(1, attempts + 1):
        started = time.time()
        try:
            print(f"[ai] attempt {attempt}/{attempts} …", flush=True)
            # httpx handles large TLS uploads more reliably than urllib on Windows.
            try:
                import certifi
                import httpx
                verify = certifi.where()
                with httpx.Client(timeout=timeout_seconds, verify=verify, http2=False) as client:
                    response = client.post(url, headers=headers, content=body)
                if response.status_code >= 400:
                    raise AIProviderError(
                        f"OpenAI HTTP {response.status_code}: {response.text[:1000]}"
                    )
                print(
                    f"[ai] OpenAI OK in {time.time() - started:.1f}s "
                    f"(HTTP {response.status_code})",
                    flush=True,
                )
                return response.json()
            except ImportError:
                request = urllib.request.Request(url, data=body, headers=headers, method="POST")
                with urllib.request.urlopen(
                    request, timeout=timeout_seconds, context=ssl_context,
                ) as response:
                    raw = response.read().decode("utf-8")
                print(f"[ai] OpenAI OK in {time.time() - started:.1f}s (urllib)", flush=True)
                return json.loads(raw)
        except AIProviderError:
            raise
        except urllib.error.HTTPError as exc:
            err_body = exc.read().decode("utf-8", errors="replace")[:1000]
            raise AIProviderError(f"OpenAI HTTP {exc.code}: {err_body}") from exc
        except (urllib.error.URLError, TimeoutError, OSError, json.JSONDecodeError, ssl.SSLError) as exc:
            last_error = exc
            print(f"[ai] attempt {attempt} failed after {time.time() - started:.1f}s: {exc}", flush=True)
            if attempt >= attempts or not _is_transient_openai_transport_error(exc):
                break
            time.sleep(min(2 ** (attempt - 1), 4))
        except Exception as exc:
            # httpx network/timeout errors land here.
            last_error = exc
            print(f"[ai] attempt {attempt} failed after {time.time() - started:.1f}s: {exc}", flush=True)
            if attempt >= attempts or not _is_transient_openai_transport_error(exc):
                break
            time.sleep(min(2 ** (attempt - 1), 4))
    raise AIProviderError("OpenAI request failed: " + str(last_error)) from last_error


class OpenAIResponsesProvider(AIProvider):
    """Independent structured extraction using original documents only.

    ``document_resolver`` must resolve an opaque document descriptor to a local
    original file.  The adapter verifies the SHA-256 before the file can leave
    the conversion service.  It never accepts a converted workbook or code
    result in the request contract.
    """

    def __init__(self, *, document_resolver: Callable[[dict], str | Path], api_key: str | None = None,
                 model: str = "gpt-5.6-luna", base_url: str = "https://api.openai.com/v1",
                 reasoning_effort: str = "medium", timeout_seconds: float = 180,
                 max_total_file_bytes: int = 50 * 1024 * 1024,
                 http_post: Callable[[str, dict[str, str], dict, float], dict] | None = None):
        require(callable(document_resolver), "OpenAI document resolver is required")
        require(nonempty(model), "OpenAI model is required")
        require(reasoning_effort in {"none", "low", "medium", "high", "xhigh", "max"},
                "Unsupported OpenAI reasoning effort")
        require(timeout_seconds > 0 and max_total_file_bytes > 0, "Invalid OpenAI provider limits")
        self._document_resolver = document_resolver
        self._api_key = api_key
        self._model = model
        self._base_url = base_url.rstrip("/")
        self._reasoning_effort = reasoning_effort
        self._timeout_seconds = timeout_seconds
        self._max_total_file_bytes = max_total_file_bytes
        self._http_post = http_post or _default_http_post

    @property
    def provider_id(self) -> str:
        return "openai"

    @property
    def model(self) -> str:
        return self._model

    def _file_parts(self, request: dict) -> list[dict]:
        result = []
        total = 0
        for document in request["documents"]:
            path = Path(self._document_resolver(copy.deepcopy(document))).resolve()
            require(path.is_file(), "Original document is missing: " + document["fileId"])
            raw = path.read_bytes()
            require(hashlib.sha256(raw).hexdigest() == document["sha256"],
                    "Original document hash mismatch: " + document["fileId"])
            total += len(raw)
            require(total <= self._max_total_file_bytes, "Original documents exceed OpenAI input limit")
            media_type = document["mediaType"] or mimetypes.guess_type(path.name)[0]
            require(nonempty(media_type), "Original document media type is required")
            item = {
                "type": "input_file",
                "filename": path.name,
                "file_data": "data:" + media_type + ";base64," + base64.b64encode(raw).decode("ascii"),
            }
            if media_type == "application/pdf":
                item["detail"] = "high"
            result.append(item)
        return result

    def _payload(self, request: dict) -> dict:
        semantic_request = copy.deepcopy(request)
        for document in semantic_request["documents"]:
            document.pop("sourceRef", None)
        prompt = {
            "task": "Extract supplier-reported facts independently from the attached original documents.",
            "rules": [
                "Return JSON that matches the supplied schema.",
                "Emit every requested field once per employee.",
                "Keep missing, unreadable, not applicable, and explicit zero distinct.",
                "Do not calculate payroll, infer absent amounts, or use knowledge of a code conversion result.",
                "Preserve signs and decimal text. Record precise source evidence.",
                "Report unexpected financial items in unknownItems.",
            ],
            "request": semantic_request,
        }
        content = self._file_parts(request)
        content.append({"type": "input_text", "text": json.dumps(prompt, ensure_ascii=False)})
        return {
            "model": self._model,
            "reasoning": {"effort": self._reasoning_effort},
            "store": False,
            "input": [{"role": "user", "content": content}],
            "text": {"format": {
                "type": "json_schema",
                "name": "ai_bill_collection",
                "strict": True,
                "schema": _collection_output_schema(request, self._model),
            }},
        }

    def _call_structured(self, payload: dict) -> dict:
        api_key = self._api_key or os.environ.get("OPENAI_API_KEY")
        if not nonempty(api_key):
            raise AIProviderError("OPENAI_API_KEY is not configured")
        response = self._http_post(
            self._base_url + "/responses",
            {"Authorization": "Bearer " + api_key, "Content-Type": "application/json"},
            payload,
            self._timeout_seconds,
        )
        self._last_response_metadata = {key: response.get(key) for key in ("id", "model", "status", "usage")}
        if response.get("status") != "completed":
            raise AIProviderError("OpenAI response did not complete: " + str(response.get("status")))
        for output in response.get("output") or []:
            if output.get("type") != "message":
                continue
            for content in output.get("content") or []:
                if content.get("type") == "refusal":
                    raise AIProviderError("OpenAI refused the extraction request")
                if content.get("type") == "output_text" and isinstance(content.get("text"), str):
                    try:
                        return json.loads(content["text"])
                    except json.JSONDecodeError as exc:
                        raise AIProviderError("OpenAI returned invalid structured JSON") from exc
        raise AIProviderError("OpenAI response contains no structured output")

    def _trust_xlsx_source_evidence(self, plan: dict, documents: list[dict],
                                    template_manifest: dict | None = None) -> dict:
        """Replace model-reported XLSX labels/evidence with cells from the original file."""
        from .template_fill import (
            _column_contexts,
            filter_inconsistent_employee_source_writes,
            inspect_source_employee_layout,
        )

        trusted = copy.deepcopy(plan)
        by_id = {item["fileId"]: item for item in documents}
        workbooks = {}
        source_layouts: dict[str, dict[str, dict]] = {}
        header_contexts: dict[tuple[str, str], list[dict]] = {}
        try:
            for item in trusted.get("writes") or []:
                source = item.get("source") if isinstance(item, dict) else None
                if not isinstance(source, dict):
                    continue
                document = by_id.get(source.get("fileId"))
                if not document:
                    continue
                path = Path(self._document_resolver(copy.deepcopy(document))).resolve()
                if path.suffix.lower() not in {".xlsx", ".xlsm"}:
                    continue
                location = str(source.get("location") or "").strip()
                if "!" not in location:
                    continue
                sheet_name, coordinate = location.rsplit("!", 1)
                sheet_name = sheet_name.strip().strip("'").replace("''", "'")
                coordinate = coordinate.strip().replace("$", "")
                if not re.fullmatch(r"[A-Za-z]{1,3}[1-9][0-9]*", coordinate):
                    continue
                if document["fileId"] not in workbooks:
                    workbooks[document["fileId"]] = (
                        load_workbook(path, data_only=False, read_only=False, keep_links=False),
                        load_workbook(path, data_only=True, read_only=False, keep_links=False),
                    )
                formula_wb, value_wb = workbooks[document["fileId"]]
                if sheet_name not in formula_wb.sheetnames or sheet_name not in value_wb.sheetnames:
                    continue
                formula_ws = formula_wb[sheet_name]
                value_ws = value_wb[sheet_name]
                layouts = source_layouts.setdefault(document["fileId"], {})
                if sheet_name not in layouts:
                    layouts[sheet_name] = inspect_source_employee_layout(formula_ws)
                cell = formula_ws[coordinate]
                actual = value_ws[coordinate].value
                header_key = (document["fileId"], sheet_name)
                if header_key not in header_contexts:
                    _header_row, sheet_contexts = _column_contexts(formula_ws)
                    header_contexts[header_key] = sheet_contexts
                source_ctx = next(
                    (ctx for ctx in header_contexts[header_key]
                     if isinstance(ctx, dict) and ctx.get("column") == cell.column),
                    None,
                )
                if source_ctx and nonempty(source_ctx.get("pathLabel")):
                    path_label = str(source_ctx["pathLabel"])
                    label = path_label if " / " in path_label else str(
                        source_ctx.get("primaryLabel") or path_label
                    )
                else:
                    label = ""
                    for row in range(cell.row - 1, 0, -1):
                        candidate = formula_ws.cell(row, cell.column)
                        if candidate.data_type == "f" or not isinstance(candidate.value, str):
                            continue
                        label = " ".join(candidate.value.split()).strip()
                        if label:
                            break
                if not label or actual is None:
                    continue
                if hasattr(actual, "isoformat"):
                    actual_text = actual.isoformat()[:10]
                    item["valueType"] = "date"
                    item["value"] = actual_text
                elif isinstance(actual, (int, float)) and not isinstance(actual, bool):
                    from decimal import Decimal
                    number = Decimal(str(actual))
                    actual_text = format(number, "f")
                    if "." in actual_text:
                        actual_text = actual_text.rstrip("0").rstrip(".")
                    item["valueType"] = "decimal"
                    item["value"] = actual_text
                else:
                    actual_text = str(actual).strip()
                    item["valueType"] = "text"
                    item["value"] = actual_text
                item["sourceLabel"] = label
                source["rawText"] = f"{label}: {actual_text}"
            # Also inspect every sheet so omitted employees are still discoverable.
            for document in documents:
                try:
                    path = Path(self._document_resolver(copy.deepcopy(document))).resolve()
                except Exception:
                    continue
                if path.suffix.lower() not in {".xlsx", ".xlsm"}:
                    continue
                if document["fileId"] not in workbooks:
                    workbooks[document["fileId"]] = (
                        load_workbook(path, data_only=False, read_only=False, keep_links=False),
                        load_workbook(path, data_only=True, read_only=False, keep_links=False),
                    )
                formula_wb, _value_wb = workbooks[document["fileId"]]
                layouts = source_layouts.setdefault(document["fileId"], {})
                for sheet_name in formula_wb.sheetnames:
                    if sheet_name not in layouts:
                        layouts[sheet_name] = inspect_source_employee_layout(formula_wb[sheet_name])
            filtered = filter_inconsistent_employee_source_writes(
                trusted,
                source_layouts=source_layouts,
                template_manifest=template_manifest,
            )
            if template_manifest is not None and self._template_has_prefilled_identities(template_manifest):
                # CODE (or another step) already anchored person names on -L.
                # Keep summary-row filters, but do not invent extra identity rows / missing-name issues.
                issues = [
                    item for item in (filtered.get("issues") or [])
                    if not (isinstance(item, dict) and item.get("code") == "MISSING_SOURCE_EMPLOYEES")
                ]
                filtered["issues"] = issues
                return filtered
            if template_manifest is not None:
                from .template_fill import seed_missing_employee_identity_writes
                return seed_missing_employee_identity_writes(
                    filtered, template_manifest, source_layouts,
                )
            return filtered
        finally:
            for formula_wb, value_wb in workbooks.values():
                formula_wb.close()
                value_wb.close()

    @staticmethod
    def _template_has_prefilled_identities(template_manifest: dict) -> bool:
        from .template_fill import _prefilled_identity_rows
        return bool(_prefilled_identity_rows(template_manifest))

    def _discover_source_employee_roster(self, documents: list[dict]) -> list[dict]:
        """Program-detect named employees from original XLSX bills before prompting."""
        from .template_fill import (
            _column_contexts, _is_service_fee_label,
            inspect_source_employee_layout, list_named_source_employees,
        )

        source_layouts: dict[str, dict[str, dict]] = {}
        workbooks = []
        value_workbooks = {}
        contexts = {}
        try:
            for document in documents:
                try:
                    path = Path(self._document_resolver(copy.deepcopy(document))).resolve()
                except Exception:
                    continue
                if path.suffix.lower() not in {".xlsx", ".xlsm"}:
                    continue
                wb = load_workbook(path, data_only=False, read_only=False, keep_links=False)
                workbooks.append(wb)
                values = load_workbook(path, data_only=True, read_only=False, keep_links=False)
                workbooks.append(values)
                value_workbooks[document["fileId"]] = values
                layouts = source_layouts.setdefault(document["fileId"], {})
                for sheet_name in wb.sheetnames:
                    layouts[sheet_name] = inspect_source_employee_layout(wb[sheet_name])
                    if layouts[sheet_name]["nameColumns"]:
                        contexts[(document["fileId"], sheet_name)] = _column_contexts(wb[sheet_name])[1]
            employees = list_named_source_employees(source_layouts)
            for employee in employees:
                ws = value_workbooks[employee["fileId"]][employee["sheetName"]]
                facts = []
                for context in contexts[(employee["fileId"], employee["sheetName"])]:
                    if _is_service_fee_label(context["pathLabel"]):
                        continue
                    cell = ws.cell(employee["sourceRow"], context["column"])
                    value = cell.value
                    if value is None or value == "" or value == 0:
                        continue
                    value_type = "text"
                    if hasattr(value, "isoformat"):
                        value = value.isoformat()[:10]
                        value_type = "date"
                    elif isinstance(value, (int, float)) and not isinstance(value, bool):
                        value_type = "decimal"
                    facts.append({
                        "sourceLabel": context["pathLabel"],
                        "location": f"{employee['sheetName']}!{cell.coordinate}",
                        "value": str(value),
                        "valueType": value_type,
                    })
                employee["inputFacts"] = facts
            return employees
        finally:
            for wb in workbooks:
                wb.close()

    def collect(self, request: dict) -> dict:
        validate_collection_request(request)
        if not nonempty(self._api_key or os.environ.get("OPENAI_API_KEY")):
            raise AIProviderError("OPENAI_API_KEY is not configured")
        return self._call_structured(self._payload(request))

    def plan_template_fill(self, collection: dict, schema: dict, documents: list[dict],
                           template_manifest: dict) -> dict:
        """Map validated AI facts to blank cells in the inspected last ``-L`` sheet."""
        validate_ai_collection(collection, schema, documents, provider_id=self.provider_id)
        require(isinstance(template_manifest, dict)
                and nonempty(template_manifest.get("templateSha256"))
                and nonempty(template_manifest.get("sheetName")), "Invalid template manifest")
        prompt = {
            "task": "Map the validated AI facts into blank input cells of the current last -L worksheet.",
            "rules": [
                "Use worksheet meaning, labels, row context, units and employee layout; do not rely on old coordinates.",
                "Map every employee name and available requested field only when the target is unambiguous.",
                "Never target formulas, headers, merged placeholders, aggregate summary rows, or another worksheet. A named employee's Total column is an input field and may be filled from that employee's source Total.",
                "Do not include values in the plan; reference only recordIndex and field.",
                "Put ambiguity or missing target structure in issues and leave the cell unwritten.",
            ],
            "schema": copy.deepcopy(schema),
            "aiCollection": copy.deepcopy(collection),
            "template": copy.deepcopy(template_manifest),
        }
        payload = {
            "model": self._model,
            "reasoning": {"effort": self._reasoning_effort},
            "store": False,
            "input": [{"role": "user", "content": [{
                "type": "input_text", "text": json.dumps(prompt, ensure_ascii=False),
            }]}],
            "text": {"format": {
                "type": "json_schema",
                "name": "ai_template_fill_plan",
                "strict": True,
                "schema": _template_plan_output_schema(template_manifest, collection, schema),
            }},
        }
        plan = self._call_structured(payload)
        from .template_fill import validate_template_fill_plan
        validate_template_fill_plan(plan, template_manifest, collection, schema, documents)
        return plan

    def _original_workbook_cells(self, documents):
        """Lossless cell-address view of original XLSX, without payroll inference.

        File ingestion may present tables without empty leading rows/columns.
        Supplying original coordinates avoids inventing source addresses; no CODE
        data, employee matching, header mapping or calculated values are added.
        """
        result = []
        characters = 0
        for document in documents:
            path = Path(self._document_resolver(copy.deepcopy(document))).resolve()
            if path.suffix.lower() not in {".xlsx", ".xlsm"}:
                continue
            wb = load_workbook(path, data_only=False, keep_links=False)
            cached = load_workbook(path, data_only=True, keep_links=False)
            try:
                sheets = []
                for ws in wb:
                    cells = {}
                    for row in ws.iter_rows():
                        for cell in row:
                            if cell.value is None:
                                continue
                            value = cell.value
                            if cell.data_type == "f":
                                value = {"formula": value, "cachedValue": cached[ws.title][cell.coordinate].value}
                            # JSON conversion only: no source-value normalization or inference.
                            encoded = json.dumps(value, ensure_ascii=False, default=str)
                            characters += len(encoded) + len(cell.coordinate) + 8
                            require(characters <= 750000, "Original XLSX coordinate view is too large; split the input bill")
                            cells[cell.coordinate] = json.loads(encoded)
                    sheets.append({"sheetName": ws.title, "cells": cells})
                result.append({"fileId": document["fileId"], "sheets": sheets})
            finally:
                wb.close()
                cached.close()
        return result

    def plan_independent_template_fill(self, *, documents, template_manifest, template_path,
                                       run_id, period, currency, validate, on_attempt,
                                       instructions: list[str] | None = None):
        """Independent proposals; deterministic code can reject but never rewrite."""
        from .independent import plan_sha256
        schema = _dynamic_template_plan_output_schema(template_manifest, documents, self._model)
        schema["properties"]["planVersion"]["enum"] = [4]
        write_schema = schema["properties"]["writes"]["items"]
        write_schema["properties"]["reason"] = {"type": "string"}
        write_schema["required"].append("reason")
        invoice_backfill_rule = (
            "When the supplier invoice bills several lines with the same payroll meaning for one "
            "employee on the same invoice—especially Medical Insurance Cover / Medical Insurance, "
            "including backfill or arrears that mention prior months (e.g. "
            "\"June, July and August ... Medical Insurance Cover 552.33\" under that employee)—sum "
            "every such nonzero line into the single Medical Insurance (or matching) target cell. "
            "Do not omit invoice backfill merely because the line text names earlier months; if it "
            "is charged on this invoice for that employee, include it in the total written value."
        )
        vertical = str(template_manifest.get("layout") or "").strip() == "vertical_label_amount"
        if vertical:
            rules = [
                "This master uses a vertical label|amount layout (UK/EOR single-person -L). template.rowFields lists each input amount cell and its left-hand field label.",
                "The blank template defines output fields, not facts. Use only original bills as evidence for the employee identity and amounts.",
                "Write the employee display name to template.employeeNameCell with semanticLabel Employee Name. A bare person name is enough; the service restores the Salary Calculation title when needed.",
                "For each reported nonzero payroll amount, write to the matching rowFields.amountCell. semanticLabel must exactly equal that row's label/primaryLabel/pathLabel.",
                invoice_backfill_rule,
                "Do not clear or overwrite left-hand labels, formulas, headers, or other sheets. Never invent salary splits or derive payroll from invoice totals.",
                "Source labels need not resemble template labels. Semantic equivalents with different wording are valid when supported by context.",
                "If a reported item is ambiguous or has no target row field, add an issue. Do not quietly omit it.",
                "Never write formulas, occupied fixed cells, merged placeholders or duplicate target cells. Use metadataInputs only with their declared type.",
                "Every write must contain original source evidence and a placement reason. XLSX evidence uses Sheet!Cell. PDF evidence uses the actual page plus item/line and quotation.",
                "sourceWorkbookCells is the original XLSX cell-address view. Use those exact coordinates for evidence; it is not CODE output.",
                "CODE-owned columns listed in template.codeOwnedColumns are excluded. Service/Agency/Management fees are outside AI payroll filling.",
                "Blank and zero can be treated as equivalent for comparison. A discrepancy with CODE is allowed; never reshape output to agree with CODE.",
            ]
        else:
            rules = [
                "The blank template defines output fields, not facts. Use only original bills as evidence for employee identities and amounts.",
                "Identify every employee independently. Choose your own employee row from template.employeeRows; never assume source row order equals target order.",
                "Use the complete header path, units and bill context to select each target column. Explain your reason. Repeated Basic Salary, EE/ER, adjustments and parent headers are distinct.",
                "Source labels need not resemble template labels. Semantic equivalents with different wording are valid when supported by context.",
                "The program does not move your columns, replace your amounts, seed names, or fill missing ordinary payroll items from CODE.",
                "Copy every nonzero reported amount with an appropriate target. Preserve currency, sign and precision. Do not invent missing salary splits or derive payroll from invoice totals.",
                invoice_backfill_rule,
                "If a reported item is ambiguous or has no target, add an issue describing the employee, source item and reason. Do not quietly omit it.",
                "Never write formulas, occupied fixed cells, headers, merged placeholders or duplicate target cells. Use metadataInputs only with their declared type.",
                "Fill required period metadata from the original bill and requested period. Fixed template labels and formulas remain intact.",
                "semanticLabel must exactly equal the selected column's primaryLabel or full pathLabel. sourceLabel must preserve the original bill's wording.",
                "Every write must contain original source evidence and a placement reason. XLSX evidence uses Sheet!Cell. PDF evidence uses the actual page plus item/line and quotation.",
                "sourceWorkbookCells is the original XLSX cell-address view, including original row/column gaps. Use those exact coordinates for evidence; it is not an employee roster, target mapping, or CODE output.",
                "CODE-owned columns listed in template.codeOwnedColumns are excluded: do not find, invent, or write these fields. They will be separately labeled as CODE-supplied after employee matching.",
                "Service/Agency/Management fees are outside AI payroll filling. Do not move them to neighboring fields. Blank and zero can be treated as equivalent for comparison.",
                "An employee Total is not a nameless summary row. Copy an explicitly reported employee total only to an appropriate blank input, not a formula.",
                "A discrepancy with CODE is allowed. Never use knowledge of CODE or a previous comparison result to make your output agree.",
            ]
        profile_instructions = [str(item) for item in (instructions or []) if nonempty(item)]
        prompt = {
            "task": "Independently read the original supplier bills and fill the blank master template. Return planVersion 4.",
            "runId": run_id, "period": period, "currency": currency,
            "rules": rules,
            "profileInstructions": profile_instructions,
            "documents": [{k: v for k, v in d.items() if k != "sourceRef"} for d in documents],
            "sourceWorkbookCells": self._original_workbook_cells(documents),
            "template": copy.deepcopy(template_manifest),
        }
        content = self._file_parts({"documents": documents})
        raw_template = Path(template_path).read_bytes()
        require(hashlib.sha256(raw_template).hexdigest() == template_manifest["templateSha256"], "Blank template changed")
        content.append({"type": "input_file", "filename": "blank-master-template.xlsx",
                        "file_data": "data:application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;base64,"
                        + base64.b64encode(raw_template).decode("ascii")})
        content.append({"type": "input_text", "text": json.dumps(prompt, ensure_ascii=False)})
        payload = {"model": self._model, "reasoning": {"effort": self._reasoning_effort}, "store": False,
                   "input": [{"role": "user", "content": content}],
                   "text": {"format": {"type": "json_schema", "name": "independent_ai_template_plan",
                                       "strict": True, "schema": schema}}}
        for attempt in range(1, 3):
            plan = self._call_structured(payload)
            digest = plan_sha256(plan)
            error = None
            try:
                validate(plan)
            except ValidationError as exc:
                error = str(exc)
            require(plan_sha256(plan) == digest, "Independent validation must not mutate the model plan")
            on_attempt({"attempt": attempt, "plan": copy.deepcopy(plan), "planSha256": digest,
                        "validationError": error, "response": copy.deepcopy(self._last_response_metadata)})
            if error is None:
                return plan
            if attempt == 2:
                raise ValidationError("Independent AI validation failed after retry: " + error)
            payload["input"][0]["content"].append({"type": "input_text", "text": json.dumps({
                "task": "Return a corrected complete replacement plan. Preserve correct source-backed writes.",
                "validationError": error, "rejectedPlan": plan,
                "rule": "Re-read original evidence. The validator provides constraints, not substitute columns or values.",
            }, ensure_ascii=False)})

    def plan_dynamic_template_fill(self, *, documents: list[dict], template_manifest: dict,
                                   run_id: str, period: str, currency: str,
                                   instructions: list[str],
                                   column_mappings: dict[str, str] | None = None,
                                   code_anchored_employees: list[dict] | None = None,
                                   code_provenance_cells: list[dict] | None = None) -> dict:
        """Read original bills against the current template itself as the output contract."""
        require(bool(documents), "Original documents are required")
        require(isinstance(template_manifest, dict)
                and nonempty(template_manifest.get("templateSha256"))
                and nonempty(template_manifest.get("sheetName")), "Invalid template manifest")
        request = {"documents": copy.deepcopy(documents)}
        column_mappings = copy.deepcopy(column_mappings or {})
        code_anchored = copy.deepcopy(code_anchored_employees or [])
        from .template_fill import normalize_code_provenance_cells, strip_code_provenance_writes
        code_provenance = normalize_code_provenance_cells(code_provenance_cells)
        reserved_cells = {
            get_column_letter(item["col"]) + str(item["row"])
            for item in code_provenance if item["sheet"] == template_manifest["sheetName"]
        }
        detected_employees = self._discover_source_employee_roster(documents)
        # Prefer CODE-anchored roster when present; otherwise fall back to source scan.
        roster_for_prompt = [
            {
                "displayName": item.get("displayName"),
                "names": item.get("names") or [],
                "targetRow": item.get("row"),
                "source": "code_identity",
            }
            for item in code_anchored
            if isinstance(item, dict) and nonempty(item.get("displayName"))
        ] or [
            {
                "displayName": item["displayName"],
                "names": item["names"],
                "sheetName": item["sheetName"],
                "sourceRow": item["sourceRow"],
                "fileId": item["fileId"],
                "source": "original_xlsx",
            }
            for item in detected_employees
        ]
        employee_count = len(roster_for_prompt)
        prompt = {
            "task": "Read the attached original supplier bills and fill the current last -L worksheet semantically.",
            "rules": [
                "Treat the current worksheet layout and labels as the field specification; there is no fixed payroll field list.",
                "Infer employees, rows and target meanings from the current worksheet and the original bills.",
                "Use worksheet meaning, labels, row context and units; do not rely on historical coordinates or field names.",
                "Only propose writes to blank input cells in this worksheet.",
                "Choose the employee row carefully. The service will resolve a uniquely known target column from columnMappings and current template labels instead of trusting the proposed column coordinate.",
                "Never target formulas, headers, merged placeholders, aggregate summary rows or another worksheet. A named employee's Total column is an input field and must be filled when the original bill provides it and the target is blank.",
                "Formula cells (valueKind=formula in template.nonemptyCells) belong to the template. Do not propose writes to them; Excel will compute them.",
                "Blank and numeric 0 are equivalent. Never write 0 / 0.0 / 0.00 into a cell — leave that cell unwritten (blank).",
                "Copy source facts faithfully. Do not invent payroll calculations. Missing or zero amounts should stay blank, not become literal zero writes.",
                "Every write must include precise evidence from an attached original file.",
                "detectedSourceEmployees.inputFacts contains program-read non-zero facts from the ORIGINAL XLSX, with exact source locations. Use these facts alongside the attached bill; they are not CODE amounts.",
                "Identity columns are sacred: BU, CN Name, EN Name, Position, Start Date, End Date and FT/PT must map to the template column with the same meaning. Never shift them one column right/left.",
                "Do not write company/BU values into CN Name, employee Chinese names into EN Name, or English names into Position.",
                "Source and template column orders often differ (extra Monthly/Allowance columns). Always match by field label meaning, never by column letter or left-to-right position.",
                "Never copy the source total/summary row into an employee row.",
                "A source row is an employee only when CN Name or EN Name (or equivalent person-name column) is non-empty. Rows with amounts but blank names are totals/padding — never map them.",
                "When codeAnchoredEmployees is present, person names are ALREADY written on those target rows. Do not rewrite names, do not reorder rows, and do not invent extra people.",
                "Match each codeAnchoredEmployees targetRow to the same person in the original bill by CN/EN name, then fill the remaining blank non-formula pay fields on that exact row.",
                "Never put employee A's pay fields on employee B's code-anchored row. The service will drop cross-person writes.",
                "Bills commonly contain MULTIPLE employees. detectedSourceEmployees/codeAnchoredEmployees is the authoritative roster; fill every listed person.",
                f"The current roster lists {employee_count} employee(s). Filling only one person when the roster has more is incorrect.",
                "One employee maps to exactly one target -L row. Never duplicate the same person across multiple target rows.",
                "All values written to one target row must come from that same source employee row. Do not mix a person's identity with another row's hours/salary/fees.",
                "Prefer contiguous blank data rows under the -L header; fill every available unambiguous non-zero, non-formula field for each employee.",
                "semanticLabel must copy the selected target column's template.columnContexts primaryLabel exactly.",
                "sourceLabel must copy the original source field label exactly; rawText must contain that label and value.",
                "Compare the source label with every target column using parent+child pathLabel when present; identical leaf labels under different parents are different fields.",
                "For Excel originals, if the source column's parent+child header path equals a template.columnContexts pathLabel, that column is the exclusive highest-priority match. Do not pick a neighboring column that only shares a similar leaf.",
                "columnMappings contains reviewed supplier-label to target-label hints for this supplier and customer; prefer a matching hint when its target exists in the current template.",
                "If a columnMappings target is absent or stale, use the current template semantics and report ambiguity instead of inventing a target.",
                "Treat qualifiers as meaningful: employee/employer, EE/ER, normal/adjustment, tier/category, generation/base, and similar suffixes are different fields.",
                "Nearby columns often share a parent topic or look alike. Never assign a value to a neighbor just because it is adjacent or partially similar — match the full source label to the full target pathLabel/primaryLabel.",
                "When several candidate columns partially match, keep only a write whose distinguishing tokens all agree (role, plan type, adjustment vs base, fee kind). If more than one column still fits, leave it unwritten and report an issue.",
                "If a source fact or target is ambiguous, report an issue and leave it unwritten.",
                "Do not write section/group headers such as Pay Items into employee value cells.",
                "Never write Service Fee / 服务费 (or any source/template column whose meaning is service fee). Leave those cells blank; they are out of scope for AI filling.",
                "Skipping the Service Fee column does not exclude an employee's reported Total. Copy the original Total as reported, even when it includes service fees; do not recompute it or subtract the skipped fee.",
                "Do not move a skipped fee amount into an adjacent fee/charge column — omit that amount entirely.",
                "codeProvenanceCells lists CODE special-source cells (FX rate, outstanding payment, business tax, salary splits, statutory fees, etc.). Never invent or fill those coordinates from the bill; the service copies them from CODE after AI filling.",
            ],
            "runId": run_id,
            "period": period,
            "currency": currency,
            "profileInstructions": list(instructions),
            "columnMappings": column_mappings,
            "codeAnchoredEmployees": code_anchored,
            "codeProvenanceCells": [
                {
                    "sheet": item.get("sheet"),
                    "row": item.get("row"),
                    "col": item.get("col"),
                    "kind": item.get("kind"),
                    "label": item.get("label"),
                    "cell": (
                        get_column_letter(int(item["col"])) + str(int(item["row"]))
                        if item.get("col") is not None and item.get("row") is not None
                        else None
                    ),
                }
                for item in code_provenance
                if isinstance(item, dict)
            ],
            "detectedSourceEmployees": detected_employees,
            "targetEmployees": roster_for_prompt,
            "documents": [{key: value for key, value in item.items() if key != "sourceRef"}
                          for item in documents],
            "template": copy.deepcopy(template_manifest),
        }
        content = self._file_parts(request)
        content.append({"type": "input_text", "text": json.dumps(prompt, ensure_ascii=False)})
        payload = {
            "model": self._model,
            "reasoning": {"effort": self._reasoning_effort},
            "store": False,
            "input": [{"role": "user", "content": content}],
            "text": {"format": {
                "type": "json_schema",
                "name": "ai_dynamic_template_fill_plan",
                "strict": True,
                "schema": _dynamic_template_plan_output_schema(template_manifest, documents, self._model),
            }},
        }
        from .template_fill import (
            _configured_target, _exact_header_column, _is_service_fee_label,
            _name_token_set, _prefilled_identity_rows, _unambiguous_column,
            resolve_dynamic_template_fill_targets, validate_dynamic_template_fill_plan,
        )

        def _validate(plan: dict) -> None:
            validate_dynamic_template_fill_plan(
                plan, template_manifest, documents, column_mappings=column_mappings,
            )
            # Validate retained writes, after formula/zero/unsafe writes are removed.
            # Pre-filled names prove identity only, never successful data extraction.
            occupied = {item["cell"] for item in template_manifest.get("nonemptyCells") or []}
            kept_cells = {item["targetCell"] for item in plan["writes"]}
            missing = []
            missing_fields = []
            for row, identity in _prefilled_identity_rows(template_manifest).items():
                expected = set()
                exact_expected = {}
                for employee in detected_employees:
                    if not identity["tokens"] & _name_token_set(employee["names"]):
                        continue
                    for fact in employee.get("inputFacts") or []:
                        if fact["valueType"] != "decimal":
                            continue
                        contexts = template_manifest["columnContexts"]
                        exact = _exact_header_column(fact["sourceLabel"], contexts)
                        configured = None if exact else _configured_target(fact["sourceLabel"], column_mappings)
                        label = configured or fact["sourceLabel"]
                        context = exact or _unambiguous_column(label, contexts)
                        if context is None:
                            continue
                        if _is_service_fee_label(context["primaryLabel"]):
                            continue
                        cell = context["columnLetter"] + str(row)
                        if cell not in occupied and cell not in reserved_cells:
                            expected.add(cell)
                            # Only exact headers or reviewed mappings impose field-level
                            # completeness; fuzzy candidates remain subject to AI review.
                            if exact is not None or configured is not None:
                                exact_expected[cell] = context["primaryLabel"]
                if expected and not (expected & kept_cells):
                    missing.append(f"{identity['displayName']} (target row {row})")
                for cell, label in exact_expected.items():
                    if cell not in kept_cells:
                        missing_fields.append(f"{identity['displayName']}: {label} -> {cell}")
            require(not missing, "AI returned no payroll data for employees with available source values: "
                    + ", ".join(missing) + ". Fill their blank input cells using detectedSourceEmployees.inputFacts.")
            require(not missing_fields, "AI omitted available source fields with unambiguous targets: "
                    + "; ".join(missing_fields[:40])
                    + ". Fill these inputs from detectedSourceEmployees.inputFacts; employee Total columns are not summary rows.")

        def _finalize(raw_plan: dict) -> dict:
            resolved = resolve_dynamic_template_fill_targets(
                self._trust_xlsx_source_evidence(
                    raw_plan, documents, template_manifest=template_manifest,
                ),
                template_manifest, column_mappings=column_mappings,
            )
            # Resolve semantic columns first: retargeting can land on a reserved cell.
            # Remove these writes before both structural and completeness validation.
            return strip_code_provenance_writes(resolved, code_provenance, template_manifest["sheetName"])

        def _needs_employee_retry(plan: dict) -> bool:
            codes = {
                item.get("code") for item in (plan.get("issues") or []) if isinstance(item, dict)
            }
            return bool(codes & {"MISSING_SOURCE_EMPLOYEES", "EMPLOYEE_IDENTITY_SEEDED"})

        first_plan = _finalize(self._call_structured(payload))
        try:
            _validate(first_plan)
        except ValidationError as exc:
            correction = {
                "task": "Return a complete corrected replacement plan.",
                "validationError": str(exc),
                "rules": [
                    "Re-check every write, not only the first rejected write.",
                    "For every sourceLabel, compare all template.columnContexts pathLabel/primaryLabel values again.",
                    "Preserve source values and evidence; correct targetCell and semanticLabel when necessary.",
                    "Do not swap nearby similar columns; require full-label / qualifier agreement or omit the write.",
                    "Never use nameless summary/total source rows; keep one employee on one target row from one source row.",
                    "Cover EVERY employee in detectedSourceEmployees with their own target row and pay fields.",
                    "If no target is unambiguous, omit that write and add a structured issue.",
                    "Never relocate a skipped fee amount into an adjacent fee column.",
                ],
                "detectedSourceEmployees": prompt["detectedSourceEmployees"],
                "rejectedPlan": first_plan,
            }
            retry_payload = copy.deepcopy(payload)
            retry_payload["input"][0]["content"].append({
                "type": "input_text",
                "text": json.dumps(correction, ensure_ascii=False),
            })
            corrected_plan = _finalize(self._call_structured(retry_payload))
            _validate(corrected_plan)
            return corrected_plan

        if employee_count > 1 and _needs_employee_retry(first_plan):
            correction = {
                "task": "Return a complete replacement plan that covers every detected employee.",
                "validationError": (
                    "Previous plan missed or only partially covered multiple source employees. "
                    "detectedSourceEmployees is authoritative."
                ),
                "rules": [
                    "Create one target -L row per detectedSourceEmployees entry.",
                    "Fill unambiguous pay fields for each employee from that employee's own source row.",
                    "Do not stop after the first employee.",
                    "Never use nameless summary/total source rows.",
                ],
                "detectedSourceEmployees": prompt["detectedSourceEmployees"],
                "rejectedPlan": first_plan,
            }
            retry_payload = copy.deepcopy(payload)
            retry_payload["input"][0]["content"].append({
                "type": "input_text",
                "text": json.dumps(correction, ensure_ascii=False),
            })
            corrected_plan = _finalize(self._call_structured(retry_payload))
            _validate(corrected_plan)
            return corrected_plan
        return first_plan


class ProviderRegistry:
    def __init__(self):
        self._providers = {}

    def register(self, provider: AIProvider) -> None:
        require(isinstance(provider, AIProvider) and nonempty(provider.provider_id), "Invalid AI provider")
        require(provider.provider_id not in self._providers, "Duplicate AI provider: " + provider.provider_id)
        self._providers[provider.provider_id] = provider

    def get(self, provider_id: str) -> AIProvider:
        if provider_id not in self._providers:
            raise ValidationError("Unknown AI provider: " + str(provider_id))
        return self._providers[provider_id]


def run_collection(provider: AIProvider, request: dict, schema: dict) -> dict:
    validate_collection_request(request)
    try:
        result = provider.collect(copy.deepcopy(request))
    except AIProviderError:
        raise
    except Exception as exc:
        raise AIProviderError(str(exc)) from exc
    validate_ai_collection(result, schema, request["documents"], run_id=request["runId"],
                           provider_id=provider.provider_id)
    return result
