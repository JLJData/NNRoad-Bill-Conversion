"""Read-only structural validation and verbatim writing of independent AI plans.

The validator never replaces source values or scores target columns. Employee
matching is used only by the explicitly labeled CODE-special overlay after AI.
A rejected plan must be corrected by the model, never by the writer.
"""
from __future__ import annotations

import hashlib
import json
import re
import shutil
from collections import Counter
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries, get_column_letter
from openpyxl.formula.translate import Translator
from openpyxl.formula.tokenizer import Token, Tokenizer

from bill_validation.contracts import ValidationError, decimal_text, nonempty, require
from .template_fill import (
    _dynamic_value, _is_service_fee_label, _parse_sheet_coordinate,
    _is_person_name_label, _normalized_semantic, _column_contexts,
    _last_l_sheet, inspect_last_l_sheet, list_l_sheet_names, list_l_sheets_in_workbook,
)

AUDIT_SHEET = "_AI_Audit"
PROVENANCE_SHEET = "_AI_CODE_Provenance"
VERTICAL_LAYOUT = "vertical_label_amount"
_TITLE_RE = re.compile(
    r"^\s*(.+?)\s*-?\s*Salary Calculation\s+for\s+FY\s+(.+?)\s*$",
    re.IGNORECASE,
)
_VERTICAL_SKIP_LABELS = {
    "details", "amount", "amountingbp", "amountinusd", "exchangerate",
    "label", "item", "description",
}


def is_excluded_service_fee(label):
    normalized = _normalized_semantic(label)
    return _is_service_fee_label(label) or any(term in normalized for term in
        ("servicefee", "sevicefee", "agencyfee", "managementfee"))


def _date_value(value) -> bool:
    return isinstance(value, date) or bool(re.fullmatch(r"\d{4}[-/]\d{1,2}[-/]\d{1,2}", str(value or "")))


def _person(value) -> str:
    if not isinstance(value, str) or _date_value(value):
        return ""
    key = _normalized_semantic(value)
    if key in {"total", "subtotal", "grandtotal", "sum", "合计", "总计", "小计"}:
        return ""
    return key


def _cell_label(value) -> str:
    if value is None or isinstance(value, (int, float, date)):
        return ""
    return " ".join(str(value).split()).strip()


def _header(ws, layout: dict) -> int:
    if layout.get("headerRow"):
        row = int(layout["headerRow"])
        require(1 <= row <= ws.max_row, "targetL.headerRow is outside the template")
        return row
    candidates = []
    for row in ws.iter_rows(max_row=min(ws.max_row, 200)):
        count = sum(c.data_type != "f" and _is_person_name_label(c.value) for c in row)
        if count:
            candidates.append((count, sum(c.value is not None for c in row), -row[0].row))
    require(bool(candidates), "Template employee header is unclear; configure targetL.headerRow")
    return -max(candidates)[2]


def _vertical_requested(layout: dict) -> bool:
    return str(layout.get("layout") or "").strip() == VERTICAL_LAYOUT


def _resolve_vertical_spec(ws, layout: dict, *, required: bool) -> dict | None:
    """UK-style label|amount sheet: one employee, fields stacked by row."""
    try:
        label_col = int(layout.get("labelColumn") or 1)
        amount_col = int(layout.get("amountColumn") or 2)
    except (TypeError, ValueError):
        if required:
            raise ValidationError("targetL.labelColumn/amountColumn must be integers") from None
        return None
    if not (1 <= label_col <= max(1, ws.max_column) and 1 <= amount_col <= max(1, ws.max_column)):
        if required:
            raise ValidationError("targetL.labelColumn/amountColumn out of sheet range")
        return None
    if label_col == amount_col:
        if required:
            raise ValidationError("targetL.labelColumn and amountColumn must differ")
        return None

    details_row = None
    title_cell = title_fy = None
    fields = []
    for row in range(1, min(ws.max_row or 1, 200) + 1):
        label_cell = ws.cell(row, label_col)
        if label_cell.data_type == "f":
            continue
        label = _cell_label(label_cell.value)
        if not label:
            continue
        title_match = _TITLE_RE.match(label)
        if title_match:
            title_cell = label_cell.coordinate
            fy = title_match.group(2).strip()
            # Master placeholders use {FY}; keep only a concrete fiscal year.
            if fy and not (fy.startswith("{") and fy.endswith("}")):
                title_fy = fy
            continue
        key = _normalized_semantic(label)
        if key in _VERTICAL_SKIP_LABELS or key.startswith("amountin"):
            if key == "details":
                details_row = row
            continue
        amount = ws.cell(row, amount_col)
        # Field rows keep a stable left-hand label; amounts may be blank, numeric,
        # short text (e.g. Days Paid), or formula. Skip pure section banners.
        if amount.data_type != "f" and isinstance(amount.value, str) and len(_cell_label(amount.value)) > 40:
            continue
        fields.append({
            "row": row,
            "label": label,
            "labelCell": label_cell.coordinate,
            "amountCell": f"{get_column_letter(amount_col)}{row}",
        })

    # Conventional tables put "Employee Name" in a header row. UK masters put the
    # same words inside the Salary Calculation title — that is not a name column.
    def _conventional_name_header(value) -> bool:
        text = _cell_label(value)
        if not text or _TITLE_RE.match(text):
            return False
        return _is_person_name_label(text)

    has_name_header = any(
        c.data_type != "f" and _conventional_name_header(c.value)
        for row in ws.iter_rows(max_row=min(ws.max_row or 1, 40))
        for c in row
    )
    looks_vertical = (
        not has_name_header
        and (
            title_cell is not None
            or details_row is not None
            or (len(fields) >= 5 and (ws.max_column or 0) <= 6)
        )
    )
    if required:
        # Explicit vertical mapping trusts label/amount columns; do not refuse
        # because the master title placeholder mentions "Employee Name".
        require(bool(fields),
                "vertical_label_amount targetL has no usable label/amount field rows")
    elif not looks_vertical or not fields:
        return None
    return {
        "labelColumn": label_col,
        "amountColumn": amount_col,
        "headerRow": int(layout.get("headerRow") or details_row or 1),
        "fields": fields,
        "employeeNameCell": title_cell,
        "employeeNameFy": title_fy,
    }


def _clear_vertical_sample_inputs(ws, layout: dict, spec: dict) -> dict:
    """Clear sample amounts/title on a vertical sheet; return manifest extras."""
    require(isinstance(layout.get("protectedCells", []), list)
            and all(isinstance(c, str) and re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", c)
                    for c in layout.get("protectedCells", [])),
            "targetL.protectedCells must be a list of cell coordinates")
    protected = set(layout.get("protectedCells") or [])
    cleared = []
    amount_col = spec["amountColumn"]
    rows = [item["row"] for item in spec["fields"]]
    require(bool(rows), "vertical_label_amount has no field rows")
    for item in spec["fields"]:
        cell = ws.cell(item["row"], amount_col)
        if (not isinstance(cell, MergedCell) and cell.data_type != "f"
                and cell.value is not None and cell.coordinate not in protected):
            cleared.append(cell.coordinate)
            cell.value = None
    name_cell = spec.get("employeeNameCell")
    if name_cell and name_cell not in protected:
        cleared.append(name_cell)
        ws[name_cell] = None
    metadata = {}
    for row in ws.iter_rows(max_row=max(min(rows) - 1, 1)):
        for cell in row:
            if cell.coordinate in protected or isinstance(cell, MergedCell):
                continue
            if _date_value(cell.value):
                metadata[cell.coordinate] = "date"
    for ref in metadata:
        if ws[ref].value is not None:
            cleared.append(ref)
            ws[ref] = None
    row_fields = [{
        "row": item["row"],
        "label": item["label"],
        "labelCell": item["labelCell"],
        "amountCell": item["amountCell"],
        "primaryLabel": item["label"],
        "pathLabel": item["label"],
    } for item in spec["fields"]]
    return {
        "layout": VERTICAL_LAYOUT,
        "labelColumn": spec["labelColumn"],
        "amountColumn": amount_col,
        "rowFields": row_fields,
        "employeeRows": rows,
        "dataStartRow": min(rows),
        "metadataInputs": metadata,
        "nameColumns": [],
        "employeeNameCell": name_cell,
        "employeeNameFy": spec.get("employeeNameFy"),
        "clearedSampleCells": cleared,
        "headerRow": spec["headerRow"],
    }


def _ensure_named_l_sheets(wb, names: list[str]) -> None:
    existing = list_l_sheet_names(list(wb.sheetnames))
    donor = existing[0] if existing else _last_l_sheet(list(wb.sheetnames))
    for name in names:
        if name in wb.sheetnames:
            continue
        dest = wb.copy_worksheet(wb[donor])
        dest.title = name


def _prepare_l_worksheet(ws, layout: dict) -> dict:
    """Clear sample inputs on one -L sheet; return inspect extras for that sheet."""
    vertical = _resolve_vertical_spec(ws, layout, required=_vertical_requested(layout))
    if vertical:
        extras = _clear_vertical_sample_inputs(ws, layout, vertical)
        extras["layout"] = VERTICAL_LAYOUT
        return extras
    header = _header(ws, layout)
    _, contexts = _column_contexts(ws, header)
    name_cols = {c["column"] for c in contexts if _is_person_name_label(c["primaryLabel"])}
    if layout.get("nameColumns") is not None:
        require(isinstance(layout["nameColumns"], list) and all(type(c) is int and 1 <= c <= ws.max_column
                for c in layout["nameColumns"]), "targetL.nameColumns must contain valid 1-based column numbers")
        name_cols = set(layout["nameColumns"])
    require(bool(name_cols), "Template needs identifiable employee-name columns")
    input_cols = {c["column"] for c in contexts}
    require(name_cols <= input_cols, "Employee-name columns must have template headers")
    require(isinstance(layout.get("protectedCells", []), list)
            and all(isinstance(c, str) and re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", c)
                    for c in layout.get("protectedCells", [])), "targetL.protectedCells must be a list of cell coordinates")
    protected = set(layout.get("protectedCells") or [])
    rows = []
    explicit_start, explicit_end = layout.get("dataStartRow"), layout.get("dataEndRow")
    require(not explicit_end or explicit_start, "targetL.dataEndRow requires dataStartRow")
    if explicit_end:
        require(header < int(explicit_start) <= int(explicit_end) <= 20000,
                "Invalid targetL employee row range")
        rows = list(range(int(explicit_start), int(explicit_end) + 1))
    else:
        aggregate_rows, detail_rows = set(), set()
        for row in ws.iter_rows(min_row=header + 1, max_row=min(ws.max_row, 20000)):
            if any(_person(ws.cell(row[0].row, c).value) for c in name_cols):
                continue
            for cell in row:
                formula = re.fullmatch(r"=SUM\(\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)\)",
                                       str(cell.value or ""), re.IGNORECASE)
                if formula:
                    first, last = int(formula[2]), int(formula[4])
                    if header < first < last <= min(ws.max_row, 20000) and not first <= cell.row <= last:
                        aggregate_rows.add(cell.row)
                        detail_rows.update(range(first, last + 1))
        for r in range(max(header + 1, int(explicit_start or 1)), min(ws.max_row, 20000) + 1):
            cells = [ws.cell(r, c) for c in input_cols]
            names = [_person(ws.cell(r, c).value) for c in name_cols]
            literals = [c.value for c in cells if c.value is not None and c.data_type != "f"]
            summary = any(_normalized_semantic(v) in {"total", "subtotal", "grandtotal", "合计", "总计"}
                          for v in literals if isinstance(v, str))
            payroll = any(c.data_type == "f" or isinstance(c.value, (int, float)) for c in cells)
            numeric_only = all(isinstance(v, (int, float)) for v in literals)
            if not summary and r not in aggregate_rows and (
                    payroll and (any(names) or numeric_only) or r in detail_rows and not literals):
                rows.append(r)
        if not rows:
            start = int(explicit_start or header + 1)
            require(not any(ws.cell(r, c).value is not None for r in range(start, ws.max_row + 1)
                            for c in input_cols),
                    "Employee input area is ambiguous; configure targetL.dataStartRow/dataEndRow/protectedCells")
            rows = list(range(start, max(start, ws.max_row) + 1))
    require(len(rows) <= 20000, "Template employee area exceeds the row limit")
    cleared = []
    for r in rows:
        for c in input_cols:
            cell = ws.cell(r, c)
            if (not isinstance(cell, MergedCell) and cell.data_type != "f"
                    and cell.value is not None and cell.coordinate not in protected):
                cleared.append(cell.coordinate)
                cell.value = None
    metadata = {}
    for row in ws.iter_rows(max_row=min(rows) - 1):
        for cell in row:
            if cell.coordinate in protected or isinstance(cell, MergedCell):
                continue
            if _date_value(cell.value):
                metadata[cell.coordinate] = "date"
    for r in rows:
        for c in input_cols:
            cell = ws.cell(r, c)
            ref = re.fullmatch(r"=\$?([A-Z]+)\$?(\d+)", str(cell.value or ""))
            if ref and int(ref[2]) < min(rows):
                target = ws[f"{ref[1]}{ref[2]}"]
                if target.coordinate not in protected and (target.value is None or _date_value(target.value)):
                    metadata[target.coordinate] = "date"
    for ref in metadata:
        if ws[ref].value is not None:
            cleared.append(ref)
            ws[ref] = None
    return {
        "headerRow": header,
        "employeeRows": rows,
        "dataStartRow": min(rows),
        "metadataInputs": metadata,
        "nameColumns": sorted(name_cols),
        "clearedSampleCells": cleared,
    }


def _manifest_from_prepared(output, extras: dict, *, sheet_name: str) -> dict:
    header = extras.get("headerRow")
    manifest = inspect_last_l_sheet(output, header_row=header, sheet_name=sheet_name)
    extra = {k: v for k, v in extras.items() if k != "headerRow"}
    manifest.update(extra)
    return manifest


def prepare_independent_template(template_path, output_path, *, layout=None,
                                 l_sheet_names: list[str] | None = None) -> dict:
    """Clear sample inputs on identifiable employee rows, never a whole row band.

    ``l_sheet_names``：仅当 CODE 结果有多张 -L 时传入，按这些表名准备（缺的从表内已有 -L 克隆）。
    未传或只有一张时仍只处理最后一个精确 ``*-L`` 表。
    """
    layout = layout or {}
    require(isinstance(layout, dict), "targetL must be an object")
    source, output = Path(template_path).resolve(), Path(output_path).resolve()
    require(source != output and not output.exists(), "Prepared template must be a new file")
    multi = bool(l_sheet_names) and len(l_sheet_names) > 1
    wb = load_workbook(source, keep_links=False)
    extras_by_name: dict[str, dict] = {}
    try:
        require(AUDIT_SHEET not in wb.sheetnames and PROVENANCE_SHEET not in wb.sheetnames,
                "Use the master template, not an earlier AI comparison workbook")
        if multi:
            _ensure_named_l_sheets(wb, list(l_sheet_names))
            targets = list(l_sheet_names)
        else:
            targets = [_last_l_sheet(list(wb.sheetnames))]
        for name in targets:
            require(name in wb.sheetnames, f"Worksheet {name!r} is missing")
            extras_by_name[name] = _prepare_l_worksheet(wb[name], layout)
        output.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output)
    finally:
        wb.close()
    manifests = [_manifest_from_prepared(output, extras_by_name[name], sheet_name=name) for name in extras_by_name]
    primary = manifests[-1] if not multi else manifests[0]
    if multi:
        primary = dict(primary)
        primary["lSheetManifests"] = manifests
        primary["lSheetNames"] = [item["sheetName"] for item in manifests]
    return primary


def code_sheet_identity_hint(code_path, sheet_name: str) -> str:
    """CODE 某张 -L 上可辨认的人名（竖表标题或姓名列），仅作分表提示。"""
    if not code_path:
        return ""
    path = Path(code_path)
    if not path.is_file():
        return ""
    wb = load_workbook(path, data_only=True, keep_links=False)
    try:
        if sheet_name not in wb.sheetnames:
            return ""
        ws = wb[sheet_name]
        for row in ws.iter_rows(max_row=min(ws.max_row or 1, 40)):
            for cell in row:
                if cell.data_type == "f":
                    continue
                match = _TITLE_RE.match(_cell_label(cell.value))
                if match:
                    return match.group(1).strip()
        _, contexts = _column_contexts(ws)
        name_cols = [c["column"] for c in contexts if _is_person_name_label(c.get("primaryLabel"))]
        for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row or 1, 80)):
            names = [_cell_label(ws.cell(row[0].row, col).value) for col in name_cols]
            names = [item for item in names if _person(item)]
            if names:
                return " / ".join(names)
    finally:
        wb.close()
    return ""


def plan_sha256(plan: dict) -> str:
    return hashlib.sha256(json.dumps(plan, ensure_ascii=False, sort_keys=True,
                                    separators=(",", ":")).encode("utf-8")).hexdigest()


def _is_vertical_manifest(manifest: dict) -> bool:
    return str(manifest.get("layout") or "").strip() == VERTICAL_LAYOUT


def validate_independent_template(manifest: dict) -> None:
    if _is_vertical_manifest(manifest):
        name_cell = manifest.get("employeeNameCell")
        if name_cell:
            require(not any(item.get("cell") == name_cell and item.get("valueKind") != "formula"
                            and _person(item.get("value"))
                            for item in manifest.get("nonemptyCells") or []),
                    "Independent AI requires a blank employee title, not a prefilled sample name")
        return
    names = set(manifest.get("nameColumns") or [])
    rows = set(manifest.get("employeeRows") or [])
    require(not any(item["row"] in rows and item["column"] in names and _person(item["value"])
                    for item in manifest["nonemptyCells"] if item["valueKind"] != "formula"),
            "Independent AI requires a blank employee table, not prefilled employee identities")


def validate_independent_plan(plan: dict, manifest: dict, documents: list[dict], *,
                              source_paths: dict[str, Path] | None = None) -> None:
    """Validate without mutating the plan or suggesting target columns/values."""
    require(isinstance(plan, dict) and plan.get("planVersion") == 4,
            "Independent AI requires planVersion 4")
    require(plan.get("templateSha256") == manifest.get("templateSha256"), "AI plan template hash is stale")
    require(plan.get("sheetName") == manifest.get("sheetName"), "AI plan targets the wrong worksheet")
    require(plan.get("automaticWriteEnabled") is False, "AI cannot authorize its own writes")
    require(nonempty(plan.get("model")), "AI plan model is required")
    require(isinstance(plan.get("issues"), list) and all(
        isinstance(issue, dict) and nonempty(issue.get("code")) and nonempty(issue.get("message"))
        for issue in plan["issues"]), "AI issues must contain a code and message")
    writes = plan.get("writes")
    require(isinstance(writes, list) and bool(writes),
            "AI returned no writes; re-read the original bill and blank template")
    require(len(writes) <= 20000, "AI plan exceeds the write limit")
    known = {item["fileId"] for item in documents}
    occupied = {item["cell"] for item in manifest.get("nonemptyCells") or []}
    contexts = {item["column"]: item for item in manifest.get("columnContexts") or []}
    row_fields = {item["amountCell"]: item for item in manifest.get("rowFields") or []
                  if isinstance(item, dict) and item.get("amountCell")}
    vertical = _is_vertical_manifest(manifest)
    name_cell = manifest.get("employeeNameCell") if vertical else None
    merged = [range_boundaries(item) for item in manifest.get("mergedRanges") or []]
    expected = {"targetCell", "semanticLabel", "sourceLabel", "valueType", "value",
                "source", "confidence", "reason"}
    seen = set()
    workbooks = {}
    pdfs = {}
    employee_rows = set(manifest.get("employeeRows") or [])
    name_rows = set()
    used_rows = set()
    named_employee = False
    source_errors = []
    write_errors = []
    try:
        for item in writes:
            try:
                require(isinstance(item, dict) and set(item) == expected, "Invalid independent AI write structure")
                cell = item.get("targetCell")
                require(isinstance(cell, str) and re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", cell) is not None,
                        "Invalid target cell coordinate")
                row, column = coordinate_to_tuple(cell)
                metadata_kind = (manifest.get("metadataInputs") or {}).get(cell)
                field = row_fields.get(cell) if vertical else None
                is_name_target = bool(name_cell and cell == name_cell)
                if vertical:
                    require((field is not None or is_name_target or metadata_kind is not None)
                            and 1 <= column <= min(16384, manifest["maxColumn"]),
                            f"{cell}: target is outside the vertical label/amount inputs")
                else:
                    require((row in employee_rows or metadata_kind is not None)
                            and 1 <= column <= min(16384, manifest["maxColumn"]),
                            f"{cell}: target is outside the employee input table")
                require(cell not in seen, f"{cell}: duplicate target write")
                seen.add(cell)
                require(cell not in occupied, f"{cell}: cannot overwrite a formula or non-empty template cell")
                require(not any(c1 <= column <= c2 and r1 <= row <= r2 and (column, row) != (c1, r1)
                                for c1, r1, c2, r2 in merged), f"{cell}: cannot write a merged placeholder")
                context = contexts.get(column)
                if vertical:
                    if is_name_target:
                        require(item.get("semanticLabel") in {"Employee Name", "Employee", name_cell},
                                f"{cell}: semanticLabel for the title name cell must be Employee Name")
                    elif metadata_kind is None:
                        require(field is not None and item.get("semanticLabel") in {
                            field.get("label"), field.get("primaryLabel"), field.get("pathLabel")},
                            f"{cell}: semanticLabel must match the row field label")
                else:
                    require(metadata_kind is not None or context is not None and item.get("semanticLabel") in {
                        context.get("primaryLabel"), context.get("pathLabel")},
                        f"{cell}: semanticLabel must describe the selected template column")
                require(nonempty(item.get("sourceLabel")) and nonempty(item.get("reason")),
                        f"{cell}: source label and placement reason are required")
                require(not any(is_excluded_service_fee(label) for label in (
                    item["sourceLabel"], item["semanticLabel"],
                    (field or context or {}).get("primaryLabel") or (field or {}).get("label"))),
                    f"{cell}: Service Fee is excluded; omit this write")
                kind, value = item.get("valueType"), item.get("value")
                require(kind in {"text", "decimal", "date"} and nonempty(value), f"{cell}: invalid typed value")
                if kind == "decimal":
                    require(Decimal(decimal_text(value)).is_finite(), f"{cell}: invalid decimal")
                elif kind == "date":
                    try:
                        date.fromisoformat(value)
                    except (TypeError, ValueError):
                        raise ValidationError(f"{cell}: date must use YYYY-MM-DD") from None
                require(metadata_kind is None or kind == metadata_kind, f"{cell}: metadata requires {metadata_kind}")
                if is_name_target:
                    title = _TITLE_RE.match(str(value)) if kind == "text" else None
                    person = _person(title.group(1) if title else value)
                    require(kind == "text" and bool(person), f"{cell}: employee name must be text")
                    named_employee = True
                elif vertical and field is not None:
                    used_rows.add(row)
                    require(column not in set(manifest.get("codeOwnedColumns") or []),
                            f"{cell}: CODE-owned special field; omit this write")
                elif row in employee_rows:
                    used_rows.add(row)
                    if column in set(manifest.get("nameColumns") or []):
                        require(kind == "text" and bool(_person(value)), f"{cell}: employee name must be text")
                        name_rows.add(row)
                    require(column not in set(manifest.get("codeOwnedColumns") or []),
                            f"{cell}: CODE-owned special field; omit this write")
                confidence = item.get("confidence")
                require(type(confidence) in {int, float} and 0 <= confidence <= 1, f"{cell}: invalid confidence")
                evidence = item.get("source")
                require(isinstance(evidence, dict) and set(evidence) == {"fileId", "location", "page", "rawText"}
                        and evidence.get("fileId") in known and nonempty(evidence.get("location"))
                        and nonempty(evidence.get("rawText")), f"{cell}: invalid source evidence")
                require(evidence["page"] is None or type(evidence["page"]) is int and evidence["page"] >= 1,
                        f"{cell}: invalid source page")
                # Check the existence of a cited XLSX cell, not its semantic mapping
                # or equality to the proposed amount. Differences must remain observable.
                path = (source_paths or {}).get(evidence["fileId"])
                if path is not None and path.suffix.lower() == ".pdf":
                    from pypdf import PdfReader
                    if evidence["fileId"] not in pdfs:
                        pdfs[evidence["fileId"]] = PdfReader(path)
                    if not (type(evidence["page"]) is int and 1 <= evidence["page"] <= len(pdfs[evidence["fileId"]].pages)):
                        source_errors.append(f"{cell}: PDF evidence must reference an existing page")
                    continue
                if path is None or path.suffix.lower() not in {".xlsx", ".xlsm"}:
                    continue
                parsed = _parse_sheet_coordinate(evidence["location"])
                if parsed is None:
                    source_errors.append(f"{cell}: XLSX evidence must cite one Sheet!Cell")
                    continue
                if evidence["fileId"] not in workbooks:
                    workbooks[evidence["fileId"]] = load_workbook(path, data_only=False, keep_links=False)
                wb = workbooks[evidence["fileId"]]
                sheet_name, source_cell = parsed
                if sheet_name not in wb.sheetnames:
                    source_errors.append(f"{cell}: source worksheet does not exist")
                    continue
                ws = wb[sheet_name]
                sr, sc = coordinate_to_tuple(source_cell)
                if not (sr <= ws.max_row and sc <= ws.max_column):
                    source_errors.append(f"{cell}: source cell is outside the source sheet ({evidence['location']})")
                elif ws[source_cell].value is None:
                    source_errors.append(f"{cell}: cited source cell is empty ({evidence['location']})")
            except ValidationError as exc:
                write_errors.append(str(exc))
    finally:
        for wb in workbooks.values():
            wb.close()
    require(not write_errors, "; ".join(write_errors[:50]))
    if vertical:
        require(named_employee or not name_cell,
                "AI must identify the employee name for the vertical title cell")
        require(bool(used_rows) or named_employee,
                "AI must fill at least one vertical amount field or the employee name")
    else:
        require(bool(name_rows), "AI must identify employees independently from the original bill")
        require(used_rows <= name_rows, "Every populated employee row needs a source-backed employee name")
    require(not source_errors, "Invalid source references: " + "; ".join(source_errors[:50]))


def write_independent_template_copy(template_path: str | Path, output_path: str | Path,
                                    plan: dict, documents: list[dict], *, source_paths=None,
                                    manifest=None, into_existing: bool = False) -> dict:
    source, output = Path(template_path).resolve(), Path(output_path).resolve()
    if into_existing:
        require(source != output and output.is_file(), "Independent AI output must already exist")
    else:
        require(source != output and not output.exists(), "Independent AI output must be a new file")
    manifest = manifest or inspect_last_l_sheet(source)
    require(hashlib.sha256(source.read_bytes()).hexdigest() == manifest["templateSha256"], "Prepared template changed")
    validate_independent_template(manifest)
    validate_independent_plan(plan, manifest, documents, source_paths=source_paths)
    output.parent.mkdir(parents=True, exist_ok=True)
    if not into_existing:
        shutil.copy2(source, output)
    try:
        wb = load_workbook(output, data_only=False, keep_links=False)
        try:
            require(manifest["sheetName"] in wb.sheetnames, "AI output is missing the target -L sheet")
            ws = wb[manifest["sheetName"]]
            name_cell = manifest.get("employeeNameCell") if _is_vertical_manifest(manifest) else None
            name_fy = manifest.get("employeeNameFy")
            for item in plan["writes"]:
                cell = ws[item["targetCell"]]
                require(not isinstance(cell, MergedCell) and cell.value is None,
                        "Cannot overwrite a template cell")
                value = _dynamic_value(item)
                if (name_cell and item["targetCell"] == name_cell and item["valueType"] == "text"
                        and name_fy and not _TITLE_RE.match(str(value))):
                    value = f"{value} Salary Calculation for FY {name_fy}"
                cell.value = value
                if item["valueType"] == "text":
                    cell.data_type = "s"  # Literal model text must never become a formula.
            for sheet in wb.worksheets:
                sheet.sheet_view.tabSelected = False
            if manifest["sheetName"] in wb.sheetnames:
                wb.active = wb.sheetnames.index(manifest["sheetName"])
                ws.sheet_view.tabSelected = True
            wb.save(output)
        finally:
            wb.close()
    except Exception:
        if not into_existing:
            output.unlink(missing_ok=True)
        raise
    return {"templateSha256": manifest["templateSha256"],
            "outputSha256": hashlib.sha256(output.read_bytes()).hexdigest(),
            "sheetName": manifest["sheetName"], "writeCount": len(plan["writes"]),
            "writtenPlanSha256": plan_sha256(plan)}


def _names_by_row(ws, header, name_columns):
    result = {}
    for row in range(header + 1, min(ws.max_row, 20000) + 1):
        names = {_person(ws.cell(row, col).value) for col in name_columns if ws.cell(row, col).data_type != "f"}
        names.discard("")
        if names:
            result[row] = names
    return result


def code_owned_columns(code_path, manifest, provenance):
    """Only field definitions reach the model, never CODE names, rows or values.

    A column is globally excluded only when it belongs to CODE for every CODE
    employee. Per-person exceptions are resolved after independent extraction.
    """
    if not code_path or not provenance:
        return []
    wb = load_workbook(code_path, keep_links=False)
    try:
        require(manifest["sheetName"] in wb.sheetnames, "CODE and template -L sheet names differ")
        ws = wb[manifest["sheetName"]]
        roster = _names_by_row(ws, manifest["headerRow"], manifest["nameColumns"])
        if not roster:
            return []
        _, contexts = _column_contexts(ws, manifest["headerRow"])
        code_fields = {c["column"]: c["pathLabel"] for c in contexts}
        target_fields = {c["column"]: c["pathLabel"] for c in manifest["columnContexts"]}
        per_column = {}
        for p in provenance:
            if p["sheet"] == ws.title and p["row"] in roster:
                per_column.setdefault(p["col"], set()).add(p["row"])
        return [col for col, rows in per_column.items() if rows == set(roster)
                and code_fields.get(col) == target_fields.get(col)
                and col not in manifest["nameColumns"]]
    finally:
        wb.close()


def copy_independent_special_cells(code_path, output_path, provenance, manifest):
    """Join special fields by unique employee identity and full header path.

    Unmatched/ambiguous employees and fields stay unresolved. Ordinary payroll,
    non-L sheets, and headers are never copied wholesale.
    """
    records, issues = [], []
    if not code_path or not provenance:
        return {"records": records, "issues": issues}
    code = load_workbook(code_path, keep_links=False)
    ai = load_workbook(output_path, keep_links=False)
    try:
        sheet = manifest["sheetName"]
        require(sheet in code.sheetnames and sheet in ai.sheetnames, "CODE/AI -L sheet mismatch")
        header = manifest["headerRow"]
        vertical = _is_vertical_manifest(manifest)
        _, cc = _column_contexts(code[sheet], header)
        _, ac = _column_contexts(ai[sheet], header)
        code_names = {} if vertical else _names_by_row(code[sheet], header, manifest["nameColumns"])
        ai_names = {} if vertical else _names_by_row(ai[sheet], header, manifest["nameColumns"])
        candidates = {r: [a for a, names in ai_names.items() if names & cn] for r, cn in code_names.items()}
        target_counts = Counter(a for options in candidates.values() for a in options)
        code_contexts = {c["column"]: c for c in cc}
        for item in provenance:
            src_sheet, sr, sc = item["sheet"], item["row"], item["col"]
            base = {"sourceSheet": src_sheet, "sourceCell": f"{get_column_letter(sc)}{sr}",
                    "kind": item.get("kind"), "origin": "CODE_SPECIAL", "includedInIndependentAccuracy": False}
            if src_sheet not in code.sheetnames or src_sheet not in ai.sheetnames:
                issues.append({"code": "CODE_SPECIAL_UNRESOLVED", "message": f"Special sheet {src_sheet!r} is missing"})
                continue
            tr, tc = sr, sc
            # Vertical single-person sheets keep special coordinates; conventional
            # tables remap employee-band specials by unique name + field path.
            if src_sheet == sheet and sr > header and not vertical:
                options = candidates.get(sr, [])
                context = code_contexts.get(sc)
                matches = [c for c in ac if context and c["pathLabel"] == context["pathLabel"]]
                # Mirrored India-L blocks often share the same pathLabel
                # (Gross Salary / Basic salary on G and S). Prefer the CODE
                # source column when it is one of the path matches.
                if len(matches) > 1:
                    same_col = [c for c in matches if c["column"] == sc]
                    matches = same_col if len(same_col) == 1 else []
                if len(options) != 1 or target_counts[options[0]] != 1 or len(matches) != 1:
                    issues.append({"code": "CODE_SPECIAL_UNRESOLVED",
                                   "message": f"{src_sheet}!{base['sourceCell']}: employee or full field path is not uniquely matched"})
                    continue
                tr, tc = options[0], matches[0]["column"]
                if tc in manifest["nameColumns"]:
                    issues.append({"code": "CODE_SPECIAL_UNRESOLVED", "message": "Special copy cannot overwrite employee identity"})
                    continue
            src, dst = code[src_sheet].cell(sr, sc), ai[src_sheet].cell(tr, tc)
            if isinstance(dst, MergedCell):
                issues.append({"code": "CODE_SPECIAL_UNRESOLVED", "message": "Special target is a merged placeholder"})
                continue
            value = src.value
            if src.data_type == "f":
                # Cross-sheet dependencies may refer to CODE-owned employee order.
                # Require a reviewed literal instead of guessing that relationship.
                if "!" in value or "[" in value:
                    issues.append({"code": "CODE_SPECIAL_UNRESOLVED",
                                   "message": f"{src_sheet}!{src.coordinate}: cross-sheet special formula needs review"})
                    continue
                value = Translator(value, origin=src.coordinate).translate_formula(dst.coordinate)
            records.append({**base, "targetSheet": src_sheet, "targetCell": dst.coordinate,
                            "previousAiValue": dst.value, "copiedValue": value,
                            "employeeMatched": src_sheet == sheet and sr > header})
            dst.value = value
        ai.save(output_path)
    finally:
        code.close()
        ai.close()
    return {"records": records, "issues": issues}


def _formula_literal(value) -> str:
    if value is None:
        return "0"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float, Decimal)):
        return f"({value})" if value < 0 else str(value)
    return '"' + str(value).replace('"', '""') + '"'


def _inline_removed_sheet_refs(wb, kept: list[str], removed: list[str], code_values=None) -> list[dict]:
    """Replace references to dropped sheets (e.g. =PN!B8 title text) with their values.

    Only the referenced context cell is inlined; the rest of the -L formula still
    evaluates the AI-written amounts. Values prefer the CODE result's calculated
    PN context, then a literal in the AI workbook. Unresolvable references fail
    explicitly instead of publishing #NAME?/#REF! after the sheet is gone.
    """
    removed_names = {name.casefold(): name for name in removed}
    inlined = []

    def referenced_value(sheet: str, ref: str, where: str):
        if code_values is not None and sheet in code_values.sheetnames:
            value = code_values[sheet][ref].value
            if value is not None:
                return value
        source = wb[sheet][ref]
        require(source.data_type != "f",
                f"{where}: references {sheet}!{ref}, which has no calculated value after {sheet} is dropped")
        return source.value

    for name in kept:
        for row in wb[name].iter_rows():
            for cell in row:
                if cell.data_type != "f" or not isinstance(cell.value, str):
                    continue
                tokens = Tokenizer(cell.value)
                values = []
                for token in tokens.items:
                    if token.type != Token.OPERAND or token.subtype != Token.RANGE or "!" not in token.value:
                        continue
                    sheet_part, ref = token.value.rsplit("!", 1)
                    sheet = (sheet_part[1:-1].replace("''", "'")
                             if sheet_part.startswith("'") and sheet_part.endswith("'") else sheet_part)
                    if sheet.casefold() not in removed_names:
                        continue
                    where = f"{name}!{cell.coordinate}"
                    require(":" not in ref, f"{where}: range reference to dropped sheet {sheet} needs review")
                    value = referenced_value(removed_names[sheet.casefold()], ref.replace("$", ""), where)
                    values.append(value)
                    token.value = _formula_literal(value)
                    token.subtype = Token.TEXT if isinstance(value, str) else Token.NUMBER
                if not values:
                    continue
                formula = cell.value
                if len(tokens.items) == 1:
                    cell.value = values[0]
                    if isinstance(values[0], str):
                        cell.data_type = "s"
                else:
                    cell.value = tokens.render()
                inlined.append({"sheet": name, "cell": cell.coordinate, "formula": formula, "value": cell.value})
    return inlined


def keep_ai_comparison_sheets(path: str | Path, code_path: str | Path | None = None,
                              keep_l_sheets: list[str] | None = None) -> list[str]:
    """Drop PN and other template sheets; keep selected -L plus independent audit sheets.

    默认只留最后一个精确 ``*-L``。CODE 有多张 -L 时传入 ``keep_l_sheets``。
    """
    output = Path(path)
    wb = load_workbook(output, keep_links=False)
    code_values = load_workbook(code_path, data_only=True, keep_links=False) if code_path else None
    try:
        last_l = _last_l_sheet(list(wb.sheetnames))
        keep_l = [name for name in (keep_l_sheets or [last_l]) if name in wb.sheetnames]
        require(bool(keep_l), "AI comparison has no -L sheet to keep")
        keep = set(keep_l) | {AUDIT_SHEET, PROVENANCE_SHEET}
        removed = [name for name in wb.sheetnames if name not in keep]
        if removed:
            _inline_removed_sheet_refs(wb, [n for n in wb.sheetnames if n in keep], removed, code_values)
        for name in removed:
            del wb[name]
        if removed:
            active = keep_l[0]
            wb.active = wb.sheetnames.index(active)
            for sheet in wb.worksheets:
                sheet.sheet_view.tabSelected = False
            wb[active].sheet_view.tabSelected = True
            wb.save(output)
        return removed
    finally:
        wb.close()
        if code_values is not None:
            code_values.close()


def embed_independent_audit(output_path, audit):
    """Keep complete model attempts in the delivered workbook, not HTTP headers."""
    wb = load_workbook(output_path, keep_links=False)
    try:
        require(AUDIT_SHEET not in wb.sheetnames and PROVENANCE_SHEET not in wb.sheetnames,
                "Audit sheets already exist")
        ws = wb.create_sheet(AUDIT_SHEET)
        ws.sheet_state = "hidden"
        raw = json.dumps(audit, ensure_ascii=False, sort_keys=True, default=str)
        ws.append(["Independent AI audit JSON; concatenate column A from row 2"])
        for index in range(0, len(raw), 15000):
            ws.append([raw[index:index + 15000]])
            ws.cell(ws.max_row, 1).data_type = "s"
        provenance = wb.create_sheet(PROVENANCE_SHEET)
        provenance.sheet_state = "hidden"
        provenance.append(["targetSheet", "targetCell", "sourceSheet", "sourceCell", "origin"])
        for record in audit.get("codeSpecialCopies") or []:
            provenance.append([record[key] for key in ("targetSheet", "targetCell", "sourceSheet", "sourceCell", "origin")])
        wb.save(output_path)
    finally:
        wb.close()
