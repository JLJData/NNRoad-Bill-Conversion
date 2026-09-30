"""Opt-in live evaluation against an existing CODE reference; never changes it.

python -m ai_collection.evaluate_independent --original bill.pdf --template master.xlsx
    --code reference.xlsx --period 2026-09 --currency AED --output output/ai_evaluation/run.xlsx

This invokes the configured model. The report compares raw first and final plans
separately. CODE is used for evaluation/special ownership only, not model filling.
"""
from __future__ import annotations

import argparse
import json
import os
import uuid
from decimal import Decimal, InvalidOperation
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple

from .independent import AUDIT_SHEET, _names_by_row, _person, is_excluded_service_fee
from .runner import run_ai_comparison_workbook


def compare_plan(plan, code_path, manifest, provenance):
    wb = load_workbook(code_path, keep_links=False)
    try:
        ws = wb[manifest["sheetName"]]
        roster = _names_by_row(ws, manifest["headerRow"], manifest["nameColumns"])
        identities, writes = {}, {}
        for item in plan.get("writes") or []:
            try:
                row, col = coordinate_to_tuple(item["targetCell"])
            except (ValueError, KeyError):
                continue
            writes[row, col] = item
            if col in manifest["nameColumns"] and row in manifest["employeeRows"]:
                identities.setdefault(row, set()).add(_person(item.get("value")))
        paired = {}
        for ar, names in identities.items():
            options = [cr for cr, cn in roster.items() if (names - {""}) & cn]
            if len(options) == 1 and sum(bool((other - {""}) & roster[options[0]]) for other in identities.values()) == 1:
                paired[ar] = options[0]
        excluded = {(p["row"], p["col"]) for p in provenance if p["sheet"] == ws.title}
        contexts = {c["column"]: c for c in manifest["columnContexts"]}
        differences, correct, missing = [], 0, []
        for ar, cr in paired.items():
            for col, ctx in contexts.items():
                if col in manifest["nameColumns"] or (cr, col) in excluded or is_excluded_service_fee(ctx["primaryLabel"]):
                    continue
                reference = ws.cell(cr, col)
                if reference.data_type == "f":
                    continue
                item = writes.get((ar, col))
                expected, actual = reference.value, item.get("value") if item else None
                if isinstance(expected, date):
                    expected_date = expected.date() if isinstance(expected, datetime) else expected
                    equal = expected_date.isoformat() == str(actual or "").split("T")[0]
                else:
                    try:
                        equal = Decimal(str(expected or 0)) == Decimal(str(actual or 0))
                    except InvalidOperation:
                        equal = str(expected or "").strip() == str(actual or "").strip()
                if equal:
                    if item:
                        correct += 1
                elif item is None:
                    missing.append({"codeCell": reference.coordinate, "aiRow": ar, "field": ctx["pathLabel"], "expected": expected})
                else:
                    differences.append({"codeCell": reference.coordinate, "aiCell": item["targetCell"],
                                        "field": ctx["pathLabel"], "expected": expected, "actual": actual})
        return {"matchedEmployees": len(paired), "unmatchedAiEmployees": len(identities) - len(paired),
                "missingCodeEmployees": len(roster) - len(set(paired.values())),
                "matchingOrdinaryWrites": correct, "differentValueOrPosition": differences,
                "missingOrdinaryInputs": missing, "modelIssueCount": len(plan.get("issues") or []),
                "note": "CODE-reference agreement, not a universal accuracy guarantee; formulas and declared CODE specials excluded."}
    finally:
        wb.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--original", action="append", required=True)
    parser.add_argument("--template", required=True)
    parser.add_argument("--code", required=True)
    parser.add_argument("--period", required=True)
    parser.add_argument("--currency", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--hints", help="JSON with targetL layout and special coordinates; ordinary mappings are ignored")
    parser.add_argument("--env-file", help="Optional existing local convert.env; credentials are never printed")
    args = parser.parse_args()
    if args.env_file:
        for raw in Path(args.env_file).read_text(encoding="utf-8-sig").splitlines():
            line = raw.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            key = key.strip()
            if key in {"OPENAI_API_KEY", "OPENAI_BASE_URL", "AI_VALIDATION_MODEL", "AI_VALIDATION_REASONING_EFFORT"}:
                os.environ[key] = value.strip().strip('"').strip("'")
    hints = json.loads(Path(args.hints).read_text(encoding="utf-8")) if args.hints else {}
    hints["comparisonMode"] = "independent"
    options = {"model": os.environ.get("AI_VALIDATION_MODEL", "gpt-5.6-luna"),
               "reasoning_effort": os.environ.get("AI_VALIDATION_REASONING_EFFORT", "medium")}
    if os.environ.get("OPENAI_BASE_URL"):
        options["base_url"] = os.environ["OPENAI_BASE_URL"]
    metadata = run_ai_comparison_workbook(original_paths=args.original, template_path=args.template,
                 output_path=args.output, code_result_path=args.code, run_id="evaluation-" + str(uuid.uuid4()),
                 period=args.period, currency=args.currency, provider_options=options, source_hints=hints)
    wb = load_workbook(args.output, read_only=True)
    try:
        audit = json.loads("".join(r[0].value for r in wb[AUDIT_SHEET].iter_rows(min_row=2)))
    finally:
        wb.close()
    report = {"metadata": metadata, "firstAttempt": compare_plan(audit["attempts"][0]["plan"], args.code,
                  audit["template"], hints.get("codeProvenanceCells") or []),
              "finalAttempt": compare_plan(audit["finalPlan"], args.code, audit["template"],
                  hints.get("codeProvenanceCells") or [])}
    destination = Path(args.output).with_suffix(".evaluation.json")
    destination.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    for stage in ("firstAttempt", "finalAttempt"):
        result = report[stage]
        print(stage, "employees=", result["matchedEmployees"], "matching=", result["matchingOrdinaryWrites"],
              "different=", len(result["differentValueOrPosition"]), "missing=", len(result["missingOrdinaryInputs"]))
    print("Evaluation report:", destination)


if __name__ == "__main__":
    main()
