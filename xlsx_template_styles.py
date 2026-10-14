# -*- coding: utf-8 -*-
"""Preserve row/column fills before converters populate template placeholders."""
from __future__ import annotations

from copy import copy

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


def materialize_inherited_fills(wb) -> int:
    """Make implicit fills explicit inside each sheet's existing used area.

    Excel displays omitted cells using row/column defaults. openpyxl creates
    those cells without inheriting the defaults; writing an employee code then
    masks the row background. Do this BEFORE writing/copying employee rows.
    Keep every explicit cell style, including deliberate white/no-fill cells.
    """
    changed = 0
    for ws in wb.worksheets:
        max_row, max_col = ws.max_row, ws.max_column
        column_fills = {}
        for dim in ws.column_dimensions.values():
            if not dim.has_style or not dim.fill.fill_type:
                continue
            for col in range(dim.min, min(dim.max, max_col) + 1):
                column_fills[col] = dim.fill

        for row in range(1, max_row + 1):
            dim = ws.row_dimensions.get(row)
            row_fill = dim.fill if dim is not None and dim.has_style else None
            columns = range(1, max_col + 1) if row_fill and row_fill.fill_type else column_fills
            for col in columns:
                # Inspect without creating a cell: explicit no-fill styles must
                # remain distinguishable from omitted template placeholders.
                cell = ws._cells.get((row, col))
                if cell is not None and (isinstance(cell, MergedCell) or cell._style is not None):
                    continue
                fill = row_fill if row_fill is not None else column_fills.get(col)
                if fill is None or not fill.fill_type:
                    continue
                ws.cell(row, col).fill = copy(fill)
                changed += 1
    return changed


def load_conversion_template(path, **kwargs):
    wb = load_workbook(path, **kwargs)
    materialize_inherited_fills(wb)
    return wb
