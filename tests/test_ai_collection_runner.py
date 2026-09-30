import copy
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from ai_collection import MockAIProvider, run_ai_comparison_workbook


class PlanningMockProvider(MockAIProvider):
    def __init__(self, name="Alice", amount="1234.50"):
        super().__init__({})
        self.name = name
        self.amount = amount
        self.plan_inputs = []

    def plan_dynamic_template_fill(self, *, documents, template_manifest, run_id, period, currency,
                                   instructions, column_mappings, **_kwargs):
        self.plan_inputs.append({
            "documents": copy.deepcopy(documents), "manifest": copy.deepcopy(template_manifest),
            "runId": run_id, "period": period, "currency": currency,
            "instructions": copy.deepcopy(instructions),
            "columnMappings": copy.deepcopy(column_mappings),
            "codeAnchoredEmployees": copy.deepcopy(_kwargs.get("code_anchored_employees")),
        })
        return {
            "planVersion": 3,
            "templateSha256": template_manifest["templateSha256"],
            "sheetName": template_manifest["sheetName"],
            "model": "fixture",
            "automaticWriteEnabled": False,
            "writes": [
                {"targetCell": "D11", "semanticLabel": "Employee", "sourceLabel": "Employee", "valueType": "text",
                 "value": self.name, "confidence": 0.99,
                 "source": {"fileId": "source-1", "location": "employee row", "page": None,
                            "rawText": "Employee: " + self.name}},
                {"targetCell": "E11", "semanticLabel": "Basic Salary", "sourceLabel": "Basic Salary", "valueType": "decimal",
                 "value": self.amount, "confidence": 0.95,
                 "source": {"fileId": "source-1", "location": "salary row", "page": None,
                            "rawText": "Basic Salary: " + self.amount}},
            ],
            "issues": [],
        }


class AICollectionRunnerTests(unittest.TestCase):
    def test_original_document_and_changed_current_template_produce_ai_copy(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            original = root / "coral-sea.xlsx"
            original.write_bytes(b"original supplier bytes")
            template = root / "current-template.xlsx"
            wb = Workbook()
            wb.active.title = "Old-L"
            target = wb.create_sheet("Redesigned Payroll-L")
            target["D10"] = "Employee"
            target["E10"] = "Basic Salary"
            target["F11"] = "=E11*2"
            wb.save(template)
            wb.close()

            fields = [
                {"field": name, "status": "missing", "value": None, "currency": "TWD", "confidence": 0.8,
                 "source": {"fileId": "source-1", "location": "not present", "page": None, "rawText": ""}}
                for name in (
                    "basic_salary", "phone_allowance", "transport_allowance", "back_pay", "meal_allowance",
                    "sick_leave_deduction", "variable_bonus", "incentive", "expense_reimbursement",
                    "unused_leave_payment", "severance_payment", "overtime_payment", "employer_insurance_total",
                )
            ]
            fields[0].update(status="found", value="1234.50")
            provider = PlanningMockProvider()
            output = root / "ai-comparison.xlsx"
            metadata = run_ai_comparison_workbook(
                original_paths=[original], template_path=template, output_path=output,
                run_id="run-9", period="2026-09", currency="twd", provider=provider,
                source_hints={"comparisonMode": "legacy"},
            )

            self.assertFalse(metadata["formalResult"])
            self.assertFalse(metadata["automaticPassEnabled"])
            self.assertEqual(metadata["sheetName"], "Redesigned Payroll-L")
            self.assertEqual(metadata["planVersion"], 3)
            self.assertNotIn("writeEvidence", metadata)
            self.assertEqual(provider.plan_inputs[0]["documents"][0]["sourceRef"],
                             "ai-input://run-9/source-1")
            self.assertNotIn("codeResult", provider.plan_inputs[0])
            result = load_workbook(output, data_only=False)
            try:
                self.assertEqual(result["Redesigned Payroll-L"]["D11"].value, "Alice")
                self.assertEqual(result["Redesigned Payroll-L"]["E11"].value, 1234.5)
                self.assertEqual(result["Redesigned Payroll-L"]["F11"].value, "=E11*2")
            finally:
                result.close()

    def test_office_can_supply_versioned_supplier_customer_configuration(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            original = root / "bill.pdf"
            original.write_bytes(b"pdf fixture")
            template = root / "template.xlsx"
            wb = Workbook()
            wb.active.title = "Customer Layout-L"
            wb.active["D10"] = "Employee"
            wb.active["E10"] = "Basic Salary"
            wb.save(template)
            wb.close()
            hints = {"schemaId": "customer-fields", "schemaVersion": "7", "instructions": ["Read source."],
                     "comparisonMode": "legacy",
                     "columnRename": {"Vendor Salary": "Basic Salary"}}
            provider = PlanningMockProvider("Bob", "88")
            output = root / "ai.xlsx"
            metadata = run_ai_comparison_workbook(
                original_paths=[original], template_path=template, output_path=output,
                run_id="run-db-config", period="2026-09", currency="TWD", profile_id="42",
                provider=provider, source_hints=hints,
            )
            self.assertEqual(metadata["profileId"], "42")
            self.assertEqual(metadata["schemaId"], "current-template-last-L")
            self.assertEqual(provider.plan_inputs[0]["instructions"], ["Read source."])
            self.assertEqual(provider.plan_inputs[0]["columnMappings"],
                             {"Vendor Salary": "Basic Salary"})
            self.assertEqual(metadata["columnMappingHintCount"], 1)

    def test_data_start_row_from_hints_and_india_profile(self):
        from ai_collection.runner import _data_start_row

        self.assertEqual(_data_start_row({"targetL": {"dataStartRow": 10}}, "taiwan-coral-sea"), 10)
        self.assertEqual(_data_start_row({"dataStartRow": 9}, "taiwan-coral-sea"), 9)
        self.assertEqual(_data_start_row({}, "india-payroll"), 10)
        self.assertEqual(_data_start_row({"engineId": "india_payroll_calc"}, "11"), 10)
        self.assertEqual(_data_start_row({}, "11", sheet_name="India-L"), 10)
        self.assertIsNone(_data_start_row({}, "taiwan-coral-sea"))

    def test_code_provenance_cells_copied_after_ai_write(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            original = root / "bill.xlsx"
            original.write_bytes(b"original")
            template = root / "template.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Payroll-L"
            ws["D10"] = "Employee Name"
            ws["E10"] = "Basic Salary"
            ws["G10"] = "Business Tax"
            wb.create_sheet("PN")
            wb.save(template)
            wb.close()

            code = root / "code.xlsx"
            cw = Workbook()
            cws = cw.active
            cws.title = "Payroll-L"
            cws["D10"] = "Employee Name"
            cws["E10"] = "Basic Salary"
            cws["G10"] = "Business Tax"
            cws["D11"] = "Alice"
            cws["G11"] = 55
            cpn = cw.create_sheet("PN")
            cpn["B5"] = 6.5
            cpn["B6"] = 100
            cw.save(code)
            cw.close()

            class ProvProvider(PlanningMockProvider):
                def plan_dynamic_template_fill(self, *, documents, template_manifest, run_id, period,
                                               currency, instructions, column_mappings, **kwargs):
                    self.plan_inputs.append({
                        "columnMappings": copy.deepcopy(column_mappings),
                        "codeAnchoredEmployees": copy.deepcopy(kwargs.get("code_anchored_employees")),
                    })
                    self.last_provenance = kwargs.get("code_provenance_cells")
                    # Name already anchored from CODE; only propose pay + a forbidden provenance write.
                    return {
                        "planVersion": 3,
                        "templateSha256": template_manifest["templateSha256"],
                        "sheetName": template_manifest["sheetName"],
                        "model": "fixture",
                        "automaticWriteEnabled": False,
                        "writes": [
                            {
                                "targetCell": "E11", "semanticLabel": "Basic Salary",
                                "sourceLabel": "Basic Salary", "valueType": "decimal",
                                "value": self.amount, "confidence": 0.95,
                                "source": {"fileId": "source-1", "location": "salary", "page": None,
                                           "rawText": "Basic Salary: " + self.amount},
                            },
                            {
                                "targetCell": "G11", "semanticLabel": "Business Tax",
                                "sourceLabel": "Business Tax", "valueType": "decimal",
                                "value": "1", "confidence": 0.9,
                                "source": {"fileId": "source-1", "location": "tax", "page": None,
                                           "rawText": "1"},
                            },
                        ],
                        "issues": [],
                    }

            provider = ProvProvider("Alice", "10")
            output = root / "ai.xlsx"
            hints = {
                "comparisonMode": "legacy",
                "codeProvenanceCells": [
                    {"sheet": "PN", "row": 5, "col": 2, "kind": "pn_fx"},
                    {"sheet": "Payroll-L", "row": 11, "col": 7, "kind": "business_tax"},
                ],
            }
            metadata = run_ai_comparison_workbook(
                original_paths=[original], template_path=template, output_path=output,
                run_id="run-prov", period="2026-09", currency="INR", profile_id="taiwan-coral-sea",
                provider=provider, source_hints=hints, code_result_path=code,
            )
            self.assertEqual(metadata["codeProvenanceCellCount"], 2)
            self.assertEqual(metadata["codeProvenanceCopyCount"], 2)
            self.assertEqual(metadata["codeOwnedCopyCount"], 3)
            self.assertEqual(metadata["codeOwnedNonLCopyCount"], 1)
            self.assertEqual(len(provider.last_provenance), 2)
            out = load_workbook(output, data_only=False)
            try:
                self.assertEqual(out["PN"]["B5"].value, 6.5)
                self.assertEqual(out["Payroll-L"]["G11"].value, 55)  # CODE wins over AI write
                self.assertEqual(out["Payroll-L"]["E11"].value, 10)
                self.assertEqual(out["Payroll-L"]["D11"].value, "Alice")
            finally:
                out.close()


if __name__ == "__main__":
    unittest.main()
