from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook
from xlsx_date_cell_fix import (
    _ensure_date_header_column_formats,
    _is_excel_date_format,
    normalize_template_date_cells,
)


class DateCellFixTests(unittest.TestCase):
    def test_inline_period_to_does_not_reformat_salary_column(self):
        wb = Workbook()
        ws = wb.active
        ws['B2'] = 'Payroll period:'
        ws['C2'] = datetime(2026, 9, 1)
        ws['D2'] = 'to'
        ws['E2'] = datetime(2026, 9, 30)
        ws['D7'] = 'OT'
        ws['D8'] = 1250
        ws['D8'].number_format = '#,##0.00'
        ws['D9'] = '=D8*2'
        ws['D9'].number_format = '#,##0.00'
        self.assertEqual(_ensure_date_header_column_formats(wb), 0)
        self.assertEqual(ws['D8'].number_format, '#,##0.00')
        self.assertEqual(ws['D9'].number_format, '#,##0.00')
        self.assertEqual(ws['D9'].value, '=D8*2')

    def test_real_date_column_headers_still_apply(self):
        wb = Workbook()
        ws = wb.active
        ws.append(['Employee', 'From', 'To', 'Salary'])
        ws['B1'].number_format = 'yyyy/m/d'  # Header text can share the date column style.
        ws.append(['Example', 46266, '=B2+29', 46266])
        ws['D2'].number_format = '#,##0.00'
        self.assertEqual(_ensure_date_header_column_formats(wb), 2)
        self.assertEqual(ws['B2'].number_format, 'yyyy/m/d')
        self.assertEqual(ws['C2'].number_format, 'yyyy/m/d')
        self.assertEqual(ws['C2'].value, '=B2+29')
        self.assertEqual(ws['D2'].number_format, '#,##0.00')

    def test_quoted_units_are_not_dates(self):
        for fmt in ['0 "days"', '0 "元/月"', '0 "year"', '0.00 "USD"']:
            with self.subTest(fmt=fmt):
                self.assertFalse(_is_excel_date_format(fmt))
        for fmt in ['yyyy/mm', 'yyyy"年"m"月"', 'mm:ss', 'yyyy/m/d']:
            with self.subTest(fmt=fmt):
                self.assertTrue(_is_excel_date_format(fmt))

    def test_postprocess_preserves_numeric_value_with_unit(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / 'units.xlsx'
            wb = Workbook()
            wb.active['A1'] = 46266
            wb.active['A1'].number_format = '0 "days"'
            wb.save(path)
            normalize_template_date_cells(path)
            saved = load_workbook(path)
            self.assertEqual(saved.active['A1'].value, 46266)
            self.assertEqual(saved.active['A1'].number_format, '0 "days"')
            saved.close()


if __name__ == '__main__':
    unittest.main()
