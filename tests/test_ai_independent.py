"""Independent model proposals must survive unchanged, including wrong amounts."""
import copy
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from pypdf import PdfWriter

from ai_collection import OpenAIResponsesProvider, run_ai_comparison_workbook
from ai_collection.independent import (
    AUDIT_SHEET, PROVENANCE_SHEET, prepare_independent_template,
    validate_independent_plan, copy_independent_special_cells,
)
from bill_validation.contracts import ValidationError


class IndependentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.env = patch.dict(os.environ, {"AI_VALIDATION_AUDIT_DIR": str(self.root / "audit")})
        self.env.start()
        self.addCleanup(self.env.stop)
        self.template = self.root / "template.xlsx"
        self.source = self.root / "source.xlsx"
        self.code = self.root / "code.xlsx"
        self.output = self.root / "ai.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Payroll-L"
        ws.append(["Employee Name", "Basic Salary", "Additional Expenses", "Special", "Total"])
        ws.append(["Old sample A", 999, 888, 777, "=SUM(B2:D2)"])
        ws.append(["Old sample B", 999, 888, 777, "=SUM(B3:D3)"])
        ws["A5"] = "Notes: retain this fixed instruction"
        wb.create_sheet("PN")["A1"] = "Fixed template label"
        wb.save(self.template)
        ws["A2"], ws["A3"] = "Alice", "Bob"
        ws["B2"], ws["B3"] = 100, 200
        ws["D2"], ws["D3"] = 11, 22
        wb.save(self.code)
        wb.close()
        wb = Workbook()
        wb.active.title = "Source"
        wb.active.append(["Name", "Monthly Payroll", "Reimbursement"])
        wb.active.append(["Alice", 100, 10])
        wb.active.append(["Bob", 200, 20])
        wb.save(self.source)
        wb.close()
        self.calls = []

    def plan(self, payload):
        prompt = json.loads(payload["input"][0]["content"][-1]["text"])
        if "template" not in prompt:
            prompt = json.loads(payload["input"][0]["content"][-2]["text"])
        manifest = prompt["template"]
        result = {"planVersion": 4, "templateSha256": manifest["templateSha256"],
                  "sheetName": manifest["sheetName"], "model": "fixture", "automaticWriteEnabled": False,
                  "writes": [], "issues": []}
        # Independent order reverses CODE; Bob's model amount is deliberately wrong.
        for row, sr, name, salary, expense in [(2, 3, "Bob", "201", "20"), (3, 2, "Alice", "100", "10")]:
            for col, label, source_label, source_col, value, kind in [
                ("A", "Employee Name", "Name", "A", name, "text"),
                ("B", "Basic Salary", "Monthly Payroll", "B", salary, "decimal"),
                ("C", "Additional Expenses", "Reimbursement", "C", expense, "decimal"),
            ]:
                result["writes"].append({"targetCell": f"{col}{row}", "semanticLabel": label,
                    "sourceLabel": source_label, "valueType": kind, "value": value, "confidence": .95,
                    "reason": "Original employee item fits the target field by meaning, not word similarity.",
                    "source": {"fileId": "source-1", "location": f"Source!{source_col}{sr}",
                               "page": None, "rawText": f"{name} {source_label} {value}"}})
        return result

    def run_ai(self, make_plan=None, hints=None):
        def post(url, headers, payload, timeout):
            self.calls.append(copy.deepcopy(payload))
            plan = make_plan(payload, len(self.calls)) if make_plan else self.plan(payload)
            return {"id": f"response-{len(self.calls)}", "model": "fixture", "status": "completed",
                    "output": [{"type": "message", "content": [{"type": "output_text", "text": json.dumps(plan)}]}]}
        provider = OpenAIResponsesProvider(document_resolver=lambda d: self.source,
                                           api_key="fixture", http_post=post)
        return run_ai_comparison_workbook(original_paths=[self.source], template_path=self.template,
                    output_path=self.output, code_result_path=self.code, run_id="independent-test",
                    period="2026-09", currency="AED", provider=provider, source_hints=hints or {})

    def test_default_keeps_model_order_semantic_targets_and_wrong_amount(self):
        hints = {"columnRename": {"Monthly Payroll": "Special"},
                 "codeProvenanceCells": [{"sheet": "Payroll-L", "row": r, "col": 4} for r in (2, 3)]}
        with patch("ai_collection.provider.OpenAIResponsesProvider._trust_xlsx_source_evidence", side_effect=AssertionError("must not rewrite")), \
             patch("ai_collection.template_fill.resolve_dynamic_template_fill_targets", side_effect=AssertionError("must not retarget")):
            meta = self.run_ai(hints=hints)
        self.assertEqual(meta["comparisonMode"], "independent")
        self.assertFalse(meta["codeIdentityAnchored"])
        self.assertEqual(meta["codeProvenanceCopyCount"], 2)
        wb = load_workbook(self.output)
        self.addCleanup(wb.close)
        ws = wb["Payroll-L"]
        self.assertEqual([ws["A2"].value, ws["A3"].value], ["Bob", "Alice"])
        self.assertEqual(ws["B2"].value, 201)  # never replaced by source=200 or CODE=200
        self.assertEqual(ws["C2"].value, 20)  # Reimbursement survives dissimilar target wording
        self.assertEqual([ws["D2"].value, ws["D3"].value], [22, 11])
        self.assertEqual(ws["A5"].value, "Notes: retain this fixed instruction")
        self.assertEqual(ws["E2"].value, "=SUM(B2:D2)")
        raw = "".join(r[0].value for r in wb[AUDIT_SHEET].iter_rows(min_row=2))
        audit = json.loads(raw)
        self.assertEqual(audit["attempts"][0]["plan"], audit["finalPlan"])
        self.assertEqual(audit["attempts"][0]["response"]["id"], "response-1")
        self.assertFalse(audit["codeSpecialCopies"][0]["includedInIndependentAccuracy"])
        prompt = json.loads(self.calls[0]["input"][0]["content"][-1]["text"])
        self.assertNotIn("columnMappings", prompt)
        self.assertNotIn("codeAnchoredEmployees", prompt)
        self.assertNotIn("Old sample", json.dumps(prompt))
        original_cells = prompt["sourceWorkbookCells"][0]["sheets"][0]["cells"]
        self.assertEqual(original_cells["B3"], 200)
        self.assertNotIn("CODE", json.dumps(prompt["sourceWorkbookCells"]))

    def test_duplicate_write_retries_and_keeps_both_attempts(self):
        def make(payload, index):
            plan = self.plan(payload)
            if index == 1:
                plan["writes"].append(copy.deepcopy(plan["writes"][0]))
            return plan
        meta = self.run_ai(make)
        self.assertEqual(meta["attemptCount"], 2)
        audit = json.loads(next((self.root / "audit").glob("*.json")).read_text(encoding="utf-8"))
        self.assertIn("duplicate", audit["attempts"][0]["validationError"])
        self.assertEqual(len(audit["attempts"][0]["plan"]["writes"]), 7)
        self.assertEqual(len(audit["finalPlan"]["writes"]), 6)

    def test_empty_source_reference_is_returned_to_model_not_fixed_by_program(self):
        def make(payload, index):
            plan = self.plan(payload)
            if index == 1:
                plan["writes"][1]["source"]["location"] = "Source!A4"
            return plan
        meta = self.run_ai(make)
        self.assertEqual(meta["attemptCount"], 2)
        audit = json.loads(next((self.root / "audit").glob("*.json")).read_text(encoding="utf-8"))
        self.assertEqual(audit["attempts"][0]["plan"]["writes"][1]["source"]["location"], "Source!A4")
        self.assertEqual(audit["finalPlan"]["writes"][1]["value"], "201")

    def test_repeated_invalid_plan_retains_failed_audit_and_no_workbook(self):
        def make(payload, index):
            plan = self.plan(payload)
            plan["writes"][1]["targetCell"] = "E2"
            plan["writes"][1]["semanticLabel"] = "Total"
            return plan
        with self.assertRaisesRegex(ValidationError, "formula or non-empty"):
            self.run_ai(make)
        self.assertFalse(self.output.exists())
        audit = json.loads(next((self.root / "audit").glob("*.json")).read_text(encoding="utf-8"))
        self.assertEqual(audit["status"], "failed")
        self.assertEqual(len(audit["attempts"]), 2)

    def test_duplicate_names_do_not_guess_special_employee(self):
        self.run_ai()
        wb = load_workbook(self.output)
        wb["Payroll-L"]["A3"] = "Bob"
        wb.save(self.output)
        wb.close()
        manifest = prepare_independent_template(self.template, self.root / "blank.xlsx")
        result = copy_independent_special_cells(self.code, self.output,
                    [{"sheet": "Payroll-L", "row": 3, "col": 4}], manifest)
        self.assertEqual(result["records"], [])
        self.assertEqual(result["issues"][0]["code"], "CODE_SPECIAL_UNRESOLVED")

    def test_date_metadata_is_not_employee_and_notes_survive(self):
        wb = load_workbook(self.template)
        ws = wb["Payroll-L"]
        ws.insert_rows(1, 6)
        ws["A2"], ws["B2"], ws["C2"] = "Period", "2026/03/01", "2026/03/31"
        ws["E8"], ws["E9"] = "=SUM(B8:D8)", "=SUM(B9:D9)"
        wb.save(self.template)
        wb.close()
        manifest = prepare_independent_template(self.template, self.root / "blank.xlsx")
        self.assertEqual(manifest["employeeRows"], [8, 9])
        self.assertEqual(manifest["metadataInputs"], {"B2": "date", "C2": "date"})
        wb = load_workbook(self.root / "blank.xlsx")
        self.addCleanup(wb.close)
        self.assertEqual(wb["Payroll-L"]["A11"].value, "Notes: retain this fixed instruction")

    def test_pdf_page_existence_validated_without_rewriting_values(self):
        manifest = prepare_independent_template(self.template, self.root / "blank.xlsx")
        fake_payload = {"input": [{"content": [{"text": json.dumps({"template": manifest})}]}]}
        plan = self.plan(fake_payload)
        pdf = self.root / "source.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        with pdf.open("wb") as stream:
            writer.write(stream)
        for write in plan["writes"]:
            write["source"].update(location="Payroll item", page=1)
        original = copy.deepcopy(plan)
        validate_independent_plan(plan, manifest, [{"fileId": "source-1"}], source_paths={"source-1": pdf})
        self.assertEqual(plan, original)
        plan["writes"][0]["source"]["page"] = 2
        with self.assertRaisesRegex(ValidationError, "existing page"):
            validate_independent_plan(plan, manifest, [{"fileId": "source-1"}], source_paths={"source-1": pdf})

    def test_code_special_formula_moves_with_employee_and_ordinary_formula_is_not_copied(self):
        self.run_ai()
        wb = load_workbook(self.code)
        wb["Payroll-L"]["D2"] = "=B2*0.1"
        wb["Payroll-L"]["C2"] = "=9999"
        wb.save(self.code)
        wb.close()
        manifest = prepare_independent_template(self.template, self.root / "blank.xlsx")
        result = copy_independent_special_cells(self.code, self.output,
                    [{"sheet": "Payroll-L", "row": 2, "col": 4}], manifest)
        self.assertEqual(len(result["records"]), 1)
        wb = load_workbook(self.output)
        self.addCleanup(wb.close)
        self.assertEqual(wb["Payroll-L"]["D3"].value, "=B3*0.1")
        self.assertEqual(wb["Payroll-L"]["C3"].value, 10)

    def test_mirrored_path_labels_prefer_code_source_column(self):
        """India-L left/right Gross Salary blocks share pathLabel; use CODE col."""
        code = self.root / "mirror-code.xlsx"
        ai = self.root / "mirror-ai.xlsx"
        for path, basic_left, basic_right in ((code, 133667, None), (ai, None, None)):
            wb = Workbook()
            ws = wb.active
            ws.title = "India-L"
            ws["B3"], ws["G3"], ws["S3"] = "Gross Salary", "Gross Salary", "Gross Salary"
            ws["B4"], ws["G4"], ws["S4"] = "Employee Name", "Basic salary", "Basic salary"
            ws["B10"], ws["G10"], ws["S10"] = "Rekha Raivyas", basic_left, basic_right
            wb.save(path)
            wb.close()
        manifest = prepare_independent_template(
            ai, self.root / "mirror-blank.xlsx",
            layout={"headerRow": 4, "dataStartRow": 10, "dataEndRow": 10, "nameColumns": [2]},
        )
        # prepare wrote a blank copy; overlay onto AI workbook with the name present.
        import shutil
        shutil.copy2(ai, self.root / "mirror-target.xlsx")
        target = self.root / "mirror-target.xlsx"
        result = copy_independent_special_cells(
            code, target,
            [{"sheet": "India-L", "row": 10, "col": 7, "kind": "indiaSalarySplit", "label": "Basic salary"}],
            manifest,
        )
        self.assertEqual(len(result["records"]), 1, result["issues"])
        self.assertEqual(result["records"][0]["targetCell"], "G10")
        wb = load_workbook(target)
        self.addCleanup(wb.close)
        self.assertEqual(wb["India-L"]["G10"].value, 133667)
        self.assertIsNone(wb["India-L"]["S10"].value)

    def test_configured_layout_protects_fixed_cells_without_country_defaults(self):
        wb = load_workbook(self.template)
        wb["Payroll-L"]["A1"] = "Person"
        wb["Payroll-L"]["C2"] = "Fixed instruction"
        wb.save(self.template)
        wb.close()
        manifest = prepare_independent_template(self.template, self.root / "blank.xlsx", layout={
            "headerRow": 1, "dataStartRow": 2, "dataEndRow": 3, "nameColumns": [1], "protectedCells": ["C2"]})
        self.assertEqual(manifest["employeeRows"], [2, 3])
        wb = load_workbook(self.root / "blank.xlsx")
        self.addCleanup(wb.close)
        self.assertEqual(wb["Payroll-L"]["C2"].value, "Fixed instruction")

    def test_top_summary_identifies_blank_detail_slots_without_using_summary_as_employee(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Changing-L"
        ws.append(["Employee Name", "Basic Salary", "Additional Expenses"])
        ws["B2"] = "=SUM(B3:B6)"
        ws["C2"] = "=SUM(C3:C6)"
        ws["A8"] = "Notes: fixed instructions"
        wb.save(self.template)
        wb.close()
        manifest = prepare_independent_template(self.template, self.root / "blank.xlsx")
        self.assertEqual(manifest["employeeRows"], [3, 4, 5, 6])
        wb = load_workbook(self.root / "blank.xlsx")
        self.addCleanup(wb.close)
        self.assertEqual(wb["Changing-L"]["B2"].value, "=SUM(B3:B6)")
        self.assertEqual(wb["Changing-L"]["A8"].value, "Notes: fixed instructions")

    def test_invalid_source_references_reported_together(self):
        manifest = prepare_independent_template(self.template, self.root / "blank.xlsx")
        payload = {"input": [{"content": [{"text": json.dumps({"template": manifest})}]}]}
        plan = self.plan(payload)
        plan["writes"][1]["source"]["location"] = "Source!A99"
        plan["writes"][2]["source"]["location"] = "Source!B99"
        before = copy.deepcopy(plan)
        with self.assertRaises(ValidationError) as caught:
            validate_independent_plan(plan, manifest, [{"fileId": "source-1"}], source_paths={"source-1": self.source})
        self.assertIn("B2:", str(caught.exception))
        self.assertIn("C2:", str(caught.exception))
        self.assertEqual(plan, before)

    def test_vertical_label_amount_clears_amounts_keeps_labels_and_accepts_plan(self):
        """UK/EOR single-person -L is a vertical form, not a row-per-employee table."""
        from ai_collection.independent import write_independent_template_copy

        wb = Workbook()
        ws = wb.active
        ws.title = "UK-L"
        ws["A3"] = "Sample Person Salary Calculation for FY 26-27"
        ws["B3"] = "=TODAY()"
        ws["A5"], ws["B5"] = "Details", "Amount in GBP"
        ws["A6"], ws["B6"] = "Days Paid", "Full Month"
        ws["A7"], ws["B7"] = "Gross Salary", 3120
        ws["A8"], ws["B8"] = "Holiday Pay", 0
        ws["A10"], ws["B10"] = "ER' NIC", 421.1
        ws["A12"], ws["B12"] = "Gross Salary + Er NIC", "=B7+B10+B8"
        ws["A14"], ws["B14"] = "PAYE (Estimated)", 10
        ws["D24"] = 1.407
        wb.save(self.template)
        wb.close()

        blank = self.root / "uk-blank.xlsx"
        manifest = prepare_independent_template(self.template, blank)
        self.assertEqual(manifest["layout"], "vertical_label_amount")
        self.assertEqual(manifest["employeeNameCell"], "A3")
        self.assertEqual(manifest["employeeNameFy"], "26-27")
        self.assertIn(7, manifest["employeeRows"])
        self.assertIn("B7", {f["amountCell"] for f in manifest["rowFields"]})
        out_wb = load_workbook(blank)
        self.addCleanup(out_wb.close)
        self.assertIsNone(out_wb["UK-L"]["A3"].value)
        self.assertIsNone(out_wb["UK-L"]["B7"].value)
        self.assertEqual(out_wb["UK-L"]["A7"].value, "Gross Salary")
        self.assertEqual(out_wb["UK-L"]["B12"].value, "=B7+B10+B8")
        self.assertEqual(out_wb["UK-L"]["B3"].value, "=TODAY()")

        plan = {
            "planVersion": 4,
            "templateSha256": manifest["templateSha256"],
            "sheetName": "UK-L",
            "model": "fixture",
            "automaticWriteEnabled": False,
            "issues": [],
            "writes": [
                {"targetCell": "A3", "semanticLabel": "Employee Name", "sourceLabel": "Name",
                 "valueType": "text", "value": "Alice Example", "confidence": 0.9,
                 "reason": "Person name from the bill title.",
                 "source": {"fileId": "source-1", "location": "Source!A2", "page": None,
                            "rawText": "Alice Example"}},
                {"targetCell": "B7", "semanticLabel": "Gross Salary", "sourceLabel": "Gross Pay",
                 "valueType": "decimal", "value": "3000", "confidence": 0.9,
                 "reason": "Gross pay maps to Gross Salary.",
                 "source": {"fileId": "source-1", "location": "Source!B2", "page": None,
                            "rawText": "Gross Pay 3000"}},
            ],
        }
        # Source sheet needs cited cells occupied.
        src = Workbook()
        src.active.title = "Source"
        src.active["A2"], src.active["B2"] = "Alice Example", 3000
        src.save(self.source)
        src.close()
        validate_independent_plan(plan, manifest, [{"fileId": "source-1"}],
                                  source_paths={"source-1": self.source})
        written = write_independent_template_copy(
            blank, self.output, plan, [{"fileId": "source-1"}],
            source_paths={"source-1": self.source}, manifest=manifest)
        self.assertEqual(written["writeCount"], 2)
        result = load_workbook(self.output)
        self.addCleanup(result.close)
        self.assertEqual(result["UK-L"]["A3"].value,
                         "Alice Example Salary Calculation for FY 26-27")
        self.assertEqual(result["UK-L"]["B7"].value, 3000)

    def test_vertical_layout_flag_required_path_also_works(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "UK-L"
        ws["A3"] = "Bob Salary Calculation for FY 25-26"
        ws["A5"], ws["B5"] = "Details", "Amount in GBP"
        for row, label, amount in ((7, "Gross Salary", 1), (8, "Holiday Pay", 2),
                                   (9, "Car Allowance", 3), (10, "ER' NIC", 4),
                                   (11, "EE'NIC", 5)):
            ws.cell(row, 1, label)
            ws.cell(row, 2, amount)
        wb.save(self.template)
        wb.close()
        manifest = prepare_independent_template(
            self.template, self.root / "uk-flag.xlsx",
            layout={"layout": "vertical_label_amount", "labelColumn": 1, "amountColumn": 2})
        self.assertEqual(manifest["layout"], "vertical_label_amount")
        self.assertEqual(manifest["amountColumn"], 2)

    def test_uk_master_placeholder_title_with_explicit_vertical_layout(self):
        """Region UK master uses {Employee Name}/ {FY} placeholders in A3."""
        wb = Workbook()
        ws = wb.active
        ws.title = "UK-L"
        ws["A3"] = "{Employee Name}- Salary Calculation for FY {FY}"
        ws["B3"] = "=TODAY()"
        ws["A5"], ws["B5"] = "Details", "Amount in GBP"
        for row, label, amount in ((7, "Gross Salary", 0), (8, "Holiday Pay", 0),
                                   (9, "Car Allowance", 0), (10, "ER' NIC", 0),
                                   (11, "EE'NIC", 0)):
            ws.cell(row, 1, label)
            ws.cell(row, 2, amount)
        wb.save(self.template)
        wb.close()
        blank = self.root / "uk-master-blank.xlsx"
        manifest = prepare_independent_template(
            self.template, blank,
            layout={"layout": "vertical_label_amount", "labelColumn": 1, "amountColumn": 2})
        self.assertEqual(manifest["layout"], "vertical_label_amount")
        self.assertEqual(manifest["employeeNameCell"], "A3")
        self.assertIsNone(manifest.get("employeeNameFy"))
        out = load_workbook(blank)
        self.addCleanup(out.close)
        self.assertIsNone(out["UK-L"]["A3"].value)
        self.assertIsNone(out["UK-L"]["B7"].value)
        self.assertEqual(out["UK-L"]["A7"].value, "Gross Salary")


if __name__ == "__main__":
    unittest.main()
