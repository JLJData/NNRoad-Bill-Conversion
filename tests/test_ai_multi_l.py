"""CODE controls sheet count; person packs and shared-roster tables stay distinct."""
import copy
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from ai_collection import OpenAIResponsesProvider, run_ai_comparison_workbook
from ai_collection.independent import (
    AUDIT_SHEET, PROVENANCE_SHEET, keep_ai_comparison_sheets, prepare_independent_template,
)
from bill_validation.contracts import ValidationError


class MultiLTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.template, self.code, self.source, self.output = [
            self.root / name for name in ("template.xlsx", "code.xlsx", "source.xlsx", "ai.xlsx")
        ]
        env = patch.dict(os.environ, {"AI_VALIDATION_AUDIT_DIR": str(self.root / "audit")})
        env.start()
        self.addCleanup(env.stop)
        self.prompts = []

    def run_ai(self, make_writes, *, code=True, hints=None):
        def post(url, headers, payload, timeout):
            prompt = json.loads(payload["input"][0]["content"][-1]["text"])
            # Retries append validation feedback after the original prompt.
            if "template" not in prompt:
                prompt = json.loads(payload["input"][0]["content"][-2]["text"])
            self.prompts.append(copy.deepcopy(prompt))
            manifest = prompt["template"]
            plan = {"planVersion": 4, "templateSha256": manifest["templateSha256"],
                    "sheetName": manifest["sheetName"], "model": "fixture", "automaticWriteEnabled": False,
                    "writes": make_writes(prompt, len(self.prompts)), "issues": []}
            return {"id": "fixture", "model": "fixture", "status": "completed",
                    "output": [{"type": "message", "content": [{"type": "output_text", "text": json.dumps(plan)}]}]}

        provider = OpenAIResponsesProvider(document_resolver=lambda d: self.source, api_key="fixture", http_post=post)
        return run_ai_comparison_workbook(
            original_paths=[self.source], template_path=self.template, output_path=self.output,
            code_result_path=self.code if code else None, run_id="multi-test", period="2026-09",
            currency="GBP", provider=provider, source_hints=hints or {},
        )

    @staticmethod
    def write(target, label, value, source, kind="decimal"):
        return {"targetCell": target, "semanticLabel": label, "sourceLabel": label,
                "value": str(value), "valueType": kind, "confidence": .95,
                "reason": "Original person and source table match this target field.",
                "source": {"fileId": "source-1", "location": source, "page": None,
                           "rawText": f"{label}: {value}"}}

    def read_output(self):
        wb = load_workbook(self.output)
        self.addCleanup(wb.close)
        return wb

    def uk_fixtures(self):
        def vertical(ws, person, amount):
            ws["A3"] = f"{person} Salary Calculation for FY 26-27"
            ws["A5"], ws["B5"] = "Details", "Amount in GBP"
            ws["A7"], ws["B7"] = "Gross Salary", amount
            ws["A8"], ws["B8"] = "Holiday Pay", 0

        wb = Workbook()
        wb.active.title = "UK-L"
        vertical(wb.active, "Old sample", 9999)
        wb.create_sheet("PN")
        wb.create_sheet("Extra-L")
        wb.save(self.template)
        wb.close()
        wb = Workbook()
        wb.active.title = "UK-L"
        vertical(wb.active, "Alice", 900)
        vertical(wb.create_sheet("UK-L (2)"), "Bob", 800)
        vertical(wb.create_sheet("UK-L (3)"), "Carol", 700)
        wb.save(self.code)
        wb.close()
        wb = Workbook()
        wb.active.title = "Alice"
        vertical(wb.active, "Alice", 100)
        vertical(wb.create_sheet("Bob"), "Bob", 200)
        vertical(wb.create_sheet("Carol"), "Carol", 300)
        wb.save(self.source)
        wb.close()

    def uk_writes(self, prompt, attempt):
        sheet = prompt["template"]["sheetName"]
        name, salary = {"UK-L": ("Alice", 101), "UK-L (2)": ("Bob", 202),
                        "UK-L (3)": ("Carol", 303)}[sheet]
        return [self.write("A3", "Employee Name", name, f"{name}!A3", "text"),
                self.write("B7", "Gross Salary", salary, f"{name}!B7")]

    def test_uk_three_people_clone_fill_keep_and_never_copy_code_amounts(self):
        self.uk_fixtures()
        meta = self.run_ai(self.uk_writes)
        wb = self.read_output()
        names = ["UK-L", "UK-L (2)", "UK-L (3)"]
        self.assertEqual(wb.sheetnames, names + [AUDIT_SHEET, PROVENANCE_SHEET])
        self.assertEqual(meta["lSheetNames"], names)
        self.assertEqual(meta["sheetName"], "UK-L")  # UI focus needs a real worksheet name.
        self.assertEqual(meta["multiLFillMode"], "one_person_per_sheet")
        for sheet, person, salary in zip(names, ["Alice", "Bob", "Carol"], [101, 202, 303]):
            self.assertTrue(wb[sheet]["A3"].value.startswith(person + " Salary"))
            self.assertEqual(wb[sheet]["B7"].value, salary)
        for prompt in self.prompts:
            instructions = " ".join(prompt["profileInstructions"])
            self.assertIn(prompt["template"]["codeSheetIdentityHint"], instructions)
            self.assertNotIn("Old sample", json.dumps(prompt))

    def test_wrong_person_title_retries_instead_of_filling_another_person(self):
        self.uk_fixtures()

        def writes(prompt, attempt):
            items = self.uk_writes(prompt, attempt)
            if attempt == 1:
                items[0]["value"] = "Bob"
                items[0]["source"]["location"] = "Bob!A3"
            return items

        meta = self.run_ai(writes)
        self.assertEqual(meta["attemptCount"], 4)
        self.assertTrue(self.read_output()["UK-L"]["A3"].value.startswith("Alice Salary"))

    def test_correct_name_with_another_persons_amount_reference_retries(self):
        self.uk_fixtures()

        def writes(prompt, attempt):
            items = self.uk_writes(prompt, attempt)
            if attempt == 1:
                items[1]["source"]["location"] = "Bob!B7"
            return items

        self.assertEqual(self.run_ai(writes)["attemptCount"], 4)

    def test_another_persons_source_quotation_retries(self):
        self.uk_fixtures()

        def writes(prompt, attempt):
            items = self.uk_writes(prompt, attempt)
            if attempt == 1:
                items[1]["source"]["rawText"] = "Bob Gross Salary 200"
            return items

        self.assertEqual(self.run_ai(writes)["attemptCount"], 4)

    def china_fixtures(self):
        wb = Workbook()
        wb.active.title = "China-L"
        for ws, salary in [(wb.active, 999), (wb.create_sheet("China-L (2)"), 888)]:
            ws.append(["Period", "2026-09"])
            ws["A4"], ws["B4"] = "Employee Name", "Basic Salary"
            ws["A5"], ws["B5"] = "Alice", salary
            ws["A6"], ws["B6"] = "Bob", salary
        wb.create_sheet("Unused-L (2)")
        wb.create_sheet("PN")
        wb.save(self.template)
        del wb["Unused-L (2)"]
        wb.save(self.code)
        wb.close()
        wb = Workbook()
        wb.active.title = "S-Payslip"
        wb.active.append(["Name", "Salary"])
        wb.active.append(["Alice", 100])
        wb.active.append(["Bob", 200])
        ws = wb.create_sheet("S-Payroll Report")
        ws.append(["Name", "Salary"])
        ws.append(["Bob", 220])
        ws.append(["Alice", 110])
        wb.save(self.source)
        wb.close()

    def china_writes(self, prompt, attempt):
        sheet = prompt["template"]["sheetName"]
        source = "S-Payslip" if sheet == "China-L" else "S-Payroll Report"
        rows = [("Alice", 100), ("Bob", 200)] if sheet == "China-L" else [("Bob", 220), ("Alice", 110)]
        return [item for target_row, source_row, (person, salary) in zip((5, 6), (2, 3), rows) for item in (
            self.write(f"A{target_row}", "Employee Name", person, f"'{source}'!A{source_row}", "text"),
            self.write(f"B{target_row}", "Basic Salary", salary, f"'{source}'!B{source_row}"),
        )]

    def test_china_two_tables_keep_same_roster_and_source_specific_values(self):
        self.china_fixtures()

        def writes(prompt, attempt):
            items = self.china_writes(prompt, attempt)
            if attempt == 1:
                items[1]["source"]["location"] = "'S-Payroll Report'!B2"
            return items

        meta = self.run_ai(writes, hints={"engineId": "china_hrone_payment_notice"})
        wb = self.read_output()
        self.assertEqual(meta["multiLFillMode"], "same_roster")
        self.assertEqual(meta["attemptCount"], 3)
        self.assertEqual(wb.sheetnames[:2], ["China-L", "China-L (2)"])
        self.assertNotIn("Unused-L (2)", wb.sheetnames)
        for sheet, names, values in [("China-L", ["Alice", "Bob"], [100, 200]),
                                     ("China-L (2)", ["Bob", "Alice"], [220, 110])]:
            self.assertEqual([wb[sheet][f"A{row}"].value for row in (5, 6)], names)
            self.assertEqual([wb[sheet][f"B{row}"].value for row in (5, 6)], values)
            self.assertIsNone(wb[sheet]["A7"].value)

    def test_no_code_or_single_code_retains_only_last_exact_l(self):
        for has_code in (False, True):
            with self.subTest(code=has_code):
                case = self.root / str(has_code)
                case.mkdir()
                self.output = case / "ai.xlsx"
                with patch.dict(os.environ, {"AI_VALIDATION_AUDIT_DIR": str(case / "audit")}):
                    wb = Workbook()
                    wb.active.title = "First-L"
                    for ws in [wb.active, wb.create_sheet("Payroll-L"), wb.create_sheet("Payroll-L (2)")]:
                        ws.append(["Employee Name", "Basic Salary"])
                        ws.append(["Sample", 999])
                    wb.save(self.template)
                    del wb["First-L"]
                    del wb["Payroll-L (2)"]
                    wb.save(self.code)
                    wb.close()
                    wb = Workbook()
                    wb.active.title = "Source"
                    wb.active.append(["Name", "Salary"])
                    wb.active.append(["Alice", 100])
                    wb.save(self.source)
                    wb.close()
                    meta = self.run_ai(lambda p, n: [
                        self.write("A2", "Employee Name", "Alice", "Source!A2", "text"),
                        self.write("B2", "Basic Salary", 101, "Source!B2"),
                    ], code=has_code)
                    self.assertEqual(meta["lSheetNames"], ["Payroll-L"])
                    self.assertEqual(self.read_output().sheetnames, ["Payroll-L", AUDIT_SHEET, PROVENANCE_SHEET])

    def test_retention_rejects_missing_selected_sheet(self):
        self.uk_fixtures()
        with self.assertRaisesRegex(ValidationError, "missing a CODE"):
            keep_ai_comparison_sheets(self.template, self.code)

    def test_second_sheet_without_identity_header_is_rejected_without_borrowing_layout(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Payroll-L"
        ws.append(["Employee Name", "Salary"])
        ws.append(["Alice", 100])
        second = wb.create_sheet("Payroll-L (2)")
        second.append(["Role", "Salary"])
        second.append(["Manager", 200])
        wb.save(self.template)
        wb.close()
        before = self.template.read_bytes()
        output = self.root / "invalid-second.xlsx"
        with self.assertRaisesRegex(ValidationError, r"Payroll-L \(2\).*employee header"):
            prepare_independent_template(
                self.template, output, l_sheet_names=["Payroll-L", "Payroll-L (2)"],
            )
        self.assertFalse(output.exists())
        self.assertEqual(self.template.read_bytes(), before)

    def test_each_sheet_independently_detects_its_header_name_column_and_employee_rows(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Payroll-L"
        ws.append(["Employee Name", "Salary"])
        ws.append(["Alice", 100])
        second = wb.create_sheet("Payroll-L (2)")
        second["A3"], second["B3"], second["C3"] = "Employee ID", "Employee Name", "Salary"
        second["A7"], second["B7"], second["C7"] = "B001", "Bob", 200
        wb.save(self.template)
        wb.close()
        output = self.root / "independent-layouts.xlsx"
        manifest = prepare_independent_template(
            self.template, output, l_sheet_names=["Payroll-L", "Payroll-L (2)"],
        )
        first, second_manifest = manifest["lSheetManifests"]
        self.assertEqual((first["headerRow"], first["nameColumns"], first["employeeRows"]), (1, [1], [2]))
        self.assertEqual((second_manifest["headerRow"], second_manifest["nameColumns"], second_manifest["employeeRows"]),
                         (3, [2], [7]))
        out = load_workbook(output)
        self.addCleanup(out.close)
        self.assertIsNone(out["Payroll-L"]["A2"].value)
        self.assertIsNone(out["Payroll-L (2)"]["B7"].value)
        self.assertEqual(out["Payroll-L (2)"]["B3"].value, "Employee Name")

    def test_second_sheet_invalid_config_is_not_hidden_by_the_first_sheet(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Payroll-L"
        ws.append(["Employee ID", "Employee Name", "Salary"])
        ws.append(["A001", "Alice", 100])
        second = wb.create_sheet("Payroll-L (2)")
        second.append(["Employee Name"])
        second.append(["Bob"])
        wb.save(self.template)
        wb.close()
        output = self.root / "invalid-config.xlsx"
        with self.assertRaisesRegex(ValidationError, r"Payroll-L \(2\).*nameColumns"):
            prepare_independent_template(
                self.template, output, layout={"headerRow": 1, "nameColumns": [2]},
                l_sheet_names=["Payroll-L", "Payroll-L (2)"],
            )
        self.assertFalse(output.exists())


if __name__ == "__main__":
    unittest.main()
