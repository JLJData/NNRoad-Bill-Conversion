import copy
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from ai_collection import (
    OpenAIResponsesProvider,
    inspect_last_l_sheet,
    resolve_dynamic_template_fill_targets,
    validate_dynamic_template_fill_plan,
    validate_template_fill_plan,
    write_ai_template_copy,
    write_dynamic_ai_template_copy,
)
from bill_validation import ValidationError


class AITemplateFillTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.template = Path(self.temp.name) / "template.xlsx"
        wb = Workbook()
        wb.active.title = "Cover"
        first = wb.create_sheet("Old-L")
        first["A1"] = "old"
        target = wb.create_sheet("TW-L")
        target["A1"] = "Employee"
        target["B1"] = "Basic Salary"
        target["C1"] = "Calculated"
        target["C2"] = "=B2*2"
        wb.create_sheet("PN")
        wb.save(self.template)
        wb.close()
        self.schema = {
            "formatVersion": 0, "schemaId": "pay", "schemaVersion": "1",
            "decimalEncoding": "string", "blankIsZero": False,
            "fields": [{"field": "salary", "type": "decimal", "unit": "currency"}],
        }
        self.documents = [{"fileId": "source-1", "sha256": "1" * 64,
                           "mediaType": "application/pdf", "sourceRef": "opaque://source-1"}]
        self.collection = {
            "engine": "ai", "provider": "mock", "model": "fixture", "runId": "run-1",
            "schemaId": "pay", "schemaVersion": "1", "period": "2026-03", "currency": "TWD",
            "automaticPassEnabled": False,
            "records": [{"sourceEntity": {"employeeId": None, "name": "Alice"}, "fields": [{
                "field": "salary", "status": "found", "value": "1234.50", "currency": "TWD",
                "source": {"fileId": "source-1", "location": "Payroll!J4"}, "confidence": 0.9,
            }]}],
            "unknownItems": [], "issues": [], "readiness": "ready_for_association",
        }
        self.manifest = inspect_last_l_sheet(self.template)
        self.plan = {
            "planVersion": 1,
            "templateSha256": self.manifest["templateSha256"],
            "sheetName": "TW-L",
            "automaticWriteEnabled": False,
            "writes": [
                {"targetCell": "A2", "recordIndex": 0, "sourceKind": "entity_name", "field": None},
                {"targetCell": "B2", "recordIndex": 0, "sourceKind": "field", "field": "salary"},
            ],
            "issues": [],
        }

    def test_inspection_uses_last_l_sheet_and_reports_formula(self):
        self.assertEqual(self.manifest["sheetName"], "TW-L")
        self.assertEqual(self.manifest["headerRow"], 1)
        salary = next(item for item in self.manifest["columnContexts"] if item["columnLetter"] == "B")
        self.assertEqual(salary["primaryLabel"], "Basic Salary")
        formula = next(item for item in self.manifest["nonemptyCells"] if item["cell"] == "C2")
        self.assertEqual(formula["valueKind"], "formula")
        self.assertEqual(formula["value"], "=B2*2")

    def test_writer_only_creates_copy_and_resolves_values_from_collection(self):
        before = self.template.read_bytes()
        output = Path(self.temp.name) / "ai-result.xlsx"
        result = write_ai_template_copy(self.template, output, self.plan, self.collection,
                                        self.schema, self.documents)
        self.assertEqual(result["sheetName"], "TW-L")
        self.assertEqual(result["writeCount"], 2)
        self.assertEqual(self.template.read_bytes(), before)
        wb = load_workbook(output, data_only=False)
        try:
            self.assertEqual(wb.active.title, "TW-L")
            self.assertEqual(wb["TW-L"]["A2"].value, "Alice")
            self.assertEqual(wb["TW-L"]["B2"].value, 1234.5)
            self.assertEqual(wb["TW-L"]["C2"].value, "=B2*2")
            self.assertIsNone(wb["Old-L"]["A2"].value)
        finally:
            wb.close()

    def test_writer_rejects_formula_target_stale_plan_and_existing_output(self):
        invalid = copy.deepcopy(self.plan)
        invalid["writes"][1]["targetCell"] = "C2"
        with self.assertRaisesRegex(ValidationError, "formula"):
            write_ai_template_copy(self.template, Path(self.temp.name) / "formula.xlsx", invalid,
                                   self.collection, self.schema, self.documents)
        stale = copy.deepcopy(self.plan)
        stale["templateSha256"] = "0" * 64
        with self.assertRaisesRegex(ValidationError, "stale"):
            validate_template_fill_plan(stale, self.manifest, self.collection, self.schema, self.documents)
        existing = Path(self.temp.name) / "existing.xlsx"
        existing.write_bytes(b"keep")
        with self.assertRaisesRegex(ValidationError, "already exists"):
            write_ai_template_copy(self.template, existing, self.plan, self.collection,
                                   self.schema, self.documents)
        self.assertEqual(existing.read_bytes(), b"keep")

    def test_plan_cannot_embed_arbitrary_value_or_write_unavailable_field(self):
        injected = copy.deepcopy(self.plan)
        injected["writes"][1]["value"] = "999999"
        with self.assertRaisesRegex(ValidationError, "unsupported data"):
            validate_template_fill_plan(injected, self.manifest, self.collection, self.schema, self.documents)
        missing = copy.deepcopy(self.collection)
        missing["records"][0]["fields"][0].update(status="missing", value=None)
        with self.assertRaisesRegex(ValidationError, "unavailable"):
            write_ai_template_copy(self.template, Path(self.temp.name) / "missing.xlsx", self.plan,
                                   missing, self.schema, self.documents)

    def test_changed_sheet_name_and_moved_table_use_current_last_l_layout(self):
        wb = load_workbook(self.template)
        try:
            ws = wb["TW-L"]
            ws.title = "New Payroll-L"
            ws.move_range("A1:C2", rows=9, cols=3, translate=True)
            wb.save(self.template)
        finally:
            wb.close()
        manifest = inspect_last_l_sheet(self.template)
        self.assertEqual(manifest["sheetName"], "New Payroll-L")
        plan = copy.deepcopy(self.plan)
        plan.update(templateSha256=manifest["templateSha256"], sheetName="New Payroll-L")
        plan["writes"][0]["targetCell"] = "D11"
        plan["writes"][1]["targetCell"] = "E11"
        output = Path(self.temp.name) / "changed-layout-ai.xlsx"
        write_ai_template_copy(self.template, output, plan, self.collection, self.schema, self.documents)
        result = load_workbook(output, data_only=False)
        try:
            self.assertEqual(result["New Payroll-L"]["D11"].value, "Alice")
            self.assertEqual(result["New Payroll-L"]["E11"].value, 1234.5)
            self.assertEqual(result["New Payroll-L"]["F11"].value, "=E11*2")
        finally:
            result.close()

    def test_unambiguous_subset_is_written_when_plan_reports_review_issues(self):
        partial = copy.deepcopy(self.plan)
        partial["issues"] = [{"code": "TARGET_AMBIGUOUS", "message": "Optional field not mapped"}]
        output = Path(self.temp.name) / "partial-ai.xlsx"
        result = write_ai_template_copy(self.template, output, partial, self.collection,
                                        self.schema, self.documents)
        self.assertEqual(result["issueCount"], 1)
        self.assertTrue(output.is_file())

    def test_openai_mapping_step_uses_current_template_manifest(self):
        expected = copy.deepcopy(self.plan)
        captured = {}

        def post(url, headers, payload, timeout):
            captured["payload"] = payload
            return {"status": "completed", "output": [{
                "type": "message",
                "content": [{"type": "output_text", "text": json.dumps(expected)}],
            }]}

        provider = OpenAIResponsesProvider(document_resolver=lambda document: self.template,
                                           api_key="test-key-not-real", http_post=post)
        collection = copy.deepcopy(self.collection)
        collection.update(provider="openai", model="gpt-5.6-luna")
        result = provider.plan_template_fill(collection, self.schema, self.documents, self.manifest)
        self.assertEqual(result, expected)
        payload = captured["payload"]
        self.assertEqual(payload["text"]["format"]["name"], "ai_template_fill_plan")
        prompt = json.loads(payload["input"][0]["content"][0]["text"])
        self.assertEqual(prompt["template"]["sheetName"], "TW-L")
        self.assertEqual(prompt["aiCollection"]["engine"], "ai")
        self.assertNotIn("codeResult", prompt)

    def test_dynamic_fill_uses_current_template_without_fixed_fields(self):
        plan = {
            "planVersion": 3, "templateSha256": self.manifest["templateSha256"],
            "sheetName": "TW-L", "model": "gpt-5.6-luna", "automaticWriteEnabled": False,
            "writes": [{
                "targetCell": "B2", "semanticLabel": "Basic Salary", "sourceLabel": "Basic Salary",
                "valueType": "decimal", "value": "55.25", "confidence": 0.91,
                "source": {"fileId": "source-1", "location": "SheetX!Z99", "page": None,
                           "rawText": "Basic Salary: 55.25"},
            }], "issues": [],
        }
        output = Path(self.temp.name) / "dynamic.xlsx"
        result = write_dynamic_ai_template_copy(self.template, output, plan, self.documents)
        self.assertEqual(result["writeCount"], 1)
        wb = load_workbook(output, data_only=False)
        try:
            self.assertEqual(wb["TW-L"]["B2"].value, 55.25)
            self.assertEqual(wb["TW-L"]["C2"].value, "=B2*2")
        finally:
            wb.close()

        invalid = copy.deepcopy(plan)
        invalid["writes"][0]["source"]["fileId"] = "unknown-file"
        with self.assertRaisesRegex(ValidationError, "source evidence"):
            validate_dynamic_template_fill_plan(invalid, self.manifest, self.documents)

        wrong_column = copy.deepcopy(plan)
        wrong_column["writes"][0].update(targetCell="A2", semanticLabel="Employee")
        validate_dynamic_template_fill_plan(wrong_column, self.manifest, self.documents)
        self.assertEqual(wrong_column["writes"], [])
        self.assertEqual(wrong_column["issues"][0]["code"], "COLUMN_MISMATCH_SKIPPED")

        wrong_label = copy.deepcopy(plan)
        wrong_label["writes"][0]["semanticLabel"] = "Employee"
        with self.assertRaisesRegex(ValidationError, "does not match the target column"):
            validate_dynamic_template_fill_plan(wrong_label, self.manifest, self.documents)

    def test_openai_dynamic_prompt_uses_template_as_field_contract(self):
        expected = {
            "planVersion": 3, "templateSha256": self.manifest["templateSha256"],
            "sheetName": "TW-L", "model": "gpt-5.6-luna", "automaticWriteEnabled": False,
            "writes": [], "issues": [],
        }
        captured = {}

        def post(url, headers, payload, timeout):
            captured["payload"] = payload
            return {"status": "completed", "output": [{"type": "message", "content": [
                {"type": "output_text", "text": json.dumps(expected)},
            ]}]}

        source = Path(self.temp.name) / "source.pdf"
        source.write_bytes(b"source")
        documents = [{"fileId": "source-1", "sha256": __import__("hashlib").sha256(b"source").hexdigest(),
                      "mediaType": "application/pdf", "sourceRef": "opaque://source-1"}]
        provider = OpenAIResponsesProvider(document_resolver=lambda document: source,
                                           api_key="test-key-not-real", http_post=post)
        provider.plan_dynamic_template_fill(
            documents=documents, template_manifest=self.manifest, run_id="run-dynamic",
            period="2026-09", currency="TWD", instructions=[],
            column_mappings={"Vendor Basic": "Basic Salary"},
        )
        payload = captured["payload"]
        self.assertEqual(payload["text"]["format"]["name"], "ai_dynamic_template_fill_plan")
        prompt = json.loads(payload["input"][0]["content"][-1]["text"])
        self.assertEqual(prompt["template"]["sheetName"], "TW-L")
        self.assertNotIn("schema", prompt)
        self.assertIn("no fixed payroll field list", prompt["rules"][0])
        self.assertEqual(prompt["columnMappings"], {"Vendor Basic": "Basic Salary"})

    def test_openai_dynamic_plan_retries_after_evidence_rejection(self):
        invalid = {
            "planVersion": 3, "templateSha256": self.manifest["templateSha256"],
            "sheetName": "TW-L", "model": "gpt-5.6-luna", "automaticWriteEnabled": False,
            "writes": [{
                "targetCell": "A2", "semanticLabel": "Employee", "sourceLabel": "Basic Salary",
                "valueType": "decimal", "value": "55.25", "confidence": 0.91,
                "source": {"fileId": "source-1", "location": "SheetX!Z99", "page": None,
                           "rawText": "55.25"},
            }], "issues": [],
        }
        corrected = copy.deepcopy(invalid)
        corrected["writes"][0]["source"]["rawText"] = "Basic Salary: 55.25"
        expected = copy.deepcopy(corrected)
        expected["writes"][0].update(targetCell="B2", semanticLabel="Basic Salary")
        calls = []

        def post(url, headers, payload, timeout):
            calls.append(copy.deepcopy(payload))
            result = invalid if len(calls) == 1 else corrected
            return {"status": "completed", "output": [{"type": "message", "content": [
                {"type": "output_text", "text": json.dumps(result)},
            ]}]}

        source = Path(self.temp.name) / "retry-source.pdf"
        source.write_bytes(b"source")
        documents = [{"fileId": "source-1", "sha256": __import__("hashlib").sha256(b"source").hexdigest(),
                      "mediaType": "application/pdf", "sourceRef": "opaque://source-1"}]
        provider = OpenAIResponsesProvider(document_resolver=lambda document: source,
                                           api_key="test-key-not-real", http_post=post)
        result = provider.plan_dynamic_template_fill(
            documents=documents, template_manifest=self.manifest, run_id="run-retry",
            period="2026-09", currency="TWD", instructions=[],
        )

        self.assertEqual(result, expected)
        self.assertEqual(len(calls), 2)
        correction = json.loads(calls[1]["input"][0]["content"][-1]["text"])
        self.assertIn("source field label", correction["validationError"])
        rejected = copy.deepcopy(invalid)
        rejected["writes"][0].update(targetCell="B2", semanticLabel="Basic Salary")
        self.assertEqual(correction["rejectedPlan"], rejected)

    def test_openai_uses_trusted_xlsx_header_to_resolve_target_column(self):
        source = Path(self.temp.name) / "source.xlsx"
        source_wb = Workbook()
        source_ws = source_wb.active
        source_ws.title = "Payroll"
        source_ws["B1"] = "Basic Salary"
        source_ws["B2"] = 55.25
        source_wb.save(source)
        source_wb.close()
        document = {
            "fileId": "source-1", "sha256": __import__("hashlib").sha256(source.read_bytes()).hexdigest(),
            "mediaType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "sourceRef": "opaque://source-1",
        }
        model_plan = {
            "planVersion": 3, "templateSha256": self.manifest["templateSha256"],
            "sheetName": "TW-L", "model": "gpt-5.6-luna", "automaticWriteEnabled": False,
            "writes": [{
                "targetCell": "A2", "semanticLabel": "Employee", "sourceLabel": "Employee",
                "valueType": "decimal", "value": "55.25", "confidence": 0.91,
                "source": {"fileId": "source-1", "location": "Payroll!B2", "page": None,
                           "rawText": "Employee: 55.25"},
            }], "issues": [],
        }
        calls = []

        def post(url, headers, payload, timeout):
            calls.append(payload)
            return {"status": "completed", "output": [{"type": "message", "content": [
                {"type": "output_text", "text": json.dumps(model_plan)},
            ]}]}

        provider = OpenAIResponsesProvider(document_resolver=lambda ignored: source,
                                           api_key="test-key-not-real", http_post=post)
        result = provider.plan_dynamic_template_fill(
            documents=[document], template_manifest=self.manifest, run_id="run-trusted-source",
            period="2026-09", currency="TWD", instructions=[],
        )

        self.assertEqual(len(calls), 1)
        self.assertEqual(result["writes"][0]["sourceLabel"], "Basic Salary")
        self.assertEqual(result["writes"][0]["source"]["rawText"], "Basic Salary: 55.25")
        self.assertEqual(result["writes"][0]["targetCell"], "B2")
        self.assertEqual(result["writes"][0]["semanticLabel"], "Basic Salary")

    def test_dynamic_semantic_guard_rejects_nearby_but_wrong_column(self):
        semantic_template = Path(self.temp.name) / "semantic-template.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Dynamic-L"
        labels = ["Employee", "HI EE", "Pension EE", "HI ER", "Pension ER 1T",
                  "2G HI EE", "2G HI ER", "2G HI ER 1T", "職災 GI"]
        for column, label in enumerate(labels, 1):
            ws.cell(1, column).value = label
        wb.save(semantic_template)
        wb.close()
        manifest = inspect_last_l_sheet(semantic_template)

        def plan(target, target_label, source_label):
            return {
                "planVersion": 3, "templateSha256": manifest["templateSha256"],
                "sheetName": "Dynamic-L", "model": "fixture", "automaticWriteEnabled": False,
                "writes": [{
                    "targetCell": target, "semanticLabel": target_label,
                    "sourceLabel": source_label, "valueType": "decimal", "value": "-563",
                    "confidence": 0.9, "source": {
                        "fileId": "source-1", "location": "Payroll!AO7", "page": None,
                        "rawText": f"{source_label}: -563",
                    },
                }], "issues": [],
            }

        validate_dynamic_template_fill_plan(plan("B2", "HI EE", "HI EE"), manifest, self.documents)
        validate_dynamic_template_fill_plan(plan("G2", "2G HI ER", "2G HI ER"), manifest, self.documents)
        pension_wrong = plan("C2", "Pension EE", "HI EE")
        validate_dynamic_template_fill_plan(pension_wrong, manifest, self.documents)
        self.assertEqual(pension_wrong["writes"], [])
        self.assertEqual(pension_wrong["issues"][0]["code"], "COLUMN_MISMATCH_SKIPPED")
        tier_wrong = plan("H2", "2G HI ER 1T", "2G HI ER")
        validate_dynamic_template_fill_plan(tier_wrong, manifest, self.documents)
        self.assertEqual(tier_wrong["writes"], [])
        self.assertEqual(tier_wrong["issues"][0]["code"], "COLUMN_MISMATCH_SKIPPED")

        configured = resolve_dynamic_template_fill_targets(
            plan("C2", "Pension EE", "Pension EE"), manifest,
            column_mappings={"Pension EE": "HI EE"},
        )
        self.assertEqual(configured["writes"][0]["targetCell"], "B2")
        self.assertEqual(configured["writes"][0]["semanticLabel"], "HI EE")
        validate_dynamic_template_fill_plan(
            configured, manifest, self.documents,
            column_mappings={"Pension EE": "HI EE"},
        )

        gi = resolve_dynamic_template_fill_targets(
            plan("G3", "2G HI ER", "GI"), manifest,
        )
        self.assertEqual(gi["writes"][0]["targetCell"], "I3")
        self.assertEqual(gi["writes"][0]["semanticLabel"], "職災 GI")
        validate_dynamic_template_fill_plan(gi, manifest, self.documents)

    def test_parent_child_path_distinguishes_same_leaf_labels(self):
        path_template = Path(self.temp.name) / "path-template.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Path-L"
        # Same leaf "Hours", different parents.
        ws["A1"] = "Pay Items"
        ws["B1"] = "Leave"
        ws["A2"] = "Hours"
        ws["B2"] = "Hours"
        wb.save(path_template)
        wb.close()
        manifest = inspect_last_l_sheet(path_template)
        pay = next(item for item in manifest["columnContexts"] if item["columnLetter"] == "A")
        leave = next(item for item in manifest["columnContexts"] if item["columnLetter"] == "B")
        self.assertEqual(pay["primaryLabel"], "Hours")
        self.assertEqual(leave["primaryLabel"], "Hours")
        self.assertEqual(pay["pathLabel"], "Pay Items / Hours")
        self.assertEqual(leave["pathLabel"], "Leave / Hours")

        # Leaf-only source is ambiguous across parent paths.
        ambiguous = resolve_dynamic_template_fill_targets({
            "planVersion": 3, "templateSha256": manifest["templateSha256"],
            "sheetName": "Path-L", "model": "fixture", "automaticWriteEnabled": False,
            "writes": [{
                "targetCell": "A3", "semanticLabel": "Hours", "sourceLabel": "Hours",
                "valueType": "decimal", "value": "8", "confidence": 0.9,
                "source": {"fileId": "source-1", "location": "S!A1", "page": None,
                           "rawText": "Hours: 8"},
            }], "issues": [],
        }, manifest)
        self.assertEqual(ambiguous["writes"][0]["targetCell"], "A3")

        # Parent+child source resolves to the Pay Items column.
        resolved = resolve_dynamic_template_fill_targets({
            "planVersion": 3, "templateSha256": manifest["templateSha256"],
            "sheetName": "Path-L", "model": "fixture", "automaticWriteEnabled": False,
            "writes": [{
                "targetCell": "B3", "semanticLabel": "Hours", "sourceLabel": "Pay Items / Hours",
                "valueType": "decimal", "value": "8", "confidence": 0.9,
                "source": {"fileId": "source-1", "location": "S!A1", "page": None,
                           "rawText": "Pay Items / Hours: 8"},
            }], "issues": [],
        }, manifest)
        self.assertEqual(resolved["writes"][0]["targetCell"], "A3")
        self.assertEqual(resolved["writes"][0]["semanticLabel"], "Hours")
        validate_dynamic_template_fill_plan(resolved, manifest, self.documents)
        self.assertEqual(len(resolved["writes"]), 1)


if __name__ == "__main__":
    unittest.main()
