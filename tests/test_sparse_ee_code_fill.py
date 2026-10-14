# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from bill_convert.sparse_ee_code_fill import ensure_ee_code_fill_from_peer_row


def test_copies_fill_only_before_value_write():
    wb = Workbook()
    ws = wb.active
    ws.cell(9, 5).fill = PatternFill(fill_type="solid", fgColor="FFDEEBF7")
    ensure_ee_code_fill_from_peer_row(ws, 9, code_col=4, peer_col=5)
    ws.cell(9, 4).value = "CUS1510-0001"
    assert ws.cell(9, 4).fill.fill_type == "solid"
    assert ws.cell(9, 4).value == "CUS1510-0001"
