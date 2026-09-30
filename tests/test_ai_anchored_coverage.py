"""Regression coverage for successful AI runs that contain only anchored names."""
import copy
import hashlib
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from ai_collection import OpenAIResponsesProvider, inspect_last_l_sheet
from bill_validation import ValidationError


class AnchoredCoverageTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.source = Path(self.temp.name) / "source.xlsx"
        self.template = Path(self.temp.name) / "template.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Payroll"
        ws.append(["CN Name", "EN Name", "Basic Salary", "Service Fee"])
        for index, name in enumerate(["Bravo", "Charlie", "Alpha", "Delta"], 1):
            ws.append([None, name, index * 100, 50])
        wb.save(self.source)
        wb.close()
        wb = Workbook()
        ws = wb.active
        ws.title = "Payroll-L"
        ws.append(["CN Name", "EN Name", "Basic Salary", "Service Fee"])
        for name in ["Alpha", "Bravo", "Charlie", "Delta"]:
            ws.append([None, name])
        wb.save(self.template)
        wb.close()
        self.manifest = inspect_last_l_sheet(self.template)
        self.empty = {
            "planVersion": 3, "templateSha256": self.manifest["templateSha256"],
            "sheetName": "Payroll-L", "model": "fixture", "automaticWriteEnabled": False,
            "writes": [], "issues": [],
        }
        self.documents = [{
            "fileId": "source-1", "sha256": hashlib.sha256(self.source.read_bytes()).hexdigest(),
            "mediaType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "sourceRef": "opaque://source-1",
        }]
        self.full = copy.deepcopy(self.empty)
        for target_row, source_row in [(2, 4), (3, 2), (4, 3), (5, 5)]:
            self.full["writes"].append({
                "targetCell": f"C{target_row}", "semanticLabel": "Basic Salary",
                "sourceLabel": "Basic Salary", "valueType": "decimal",
                "value": str((source_row - 1) * 100), "confidence": 1.0,
                "source": {"fileId": "source-1", "location": f"Payroll!C{source_row}",
                           "page": None, "rawText": f"Basic Salary: {(source_row - 1) * 100}"},
            })

    def run_plans(self, plans, provenance=None):
        self.calls = []

        def post(url, headers, payload, timeout):
            self.calls.append(copy.deepcopy(payload))
            plan = plans[min(len(self.calls) - 1, len(plans) - 1)]
            return {"status": "completed", "output": [{"type": "message", "content": [
                {"type": "output_text", "text": json.dumps(plan)},
            ]}]}

        provider = OpenAIResponsesProvider(document_resolver=lambda document: self.source,
                                           api_key="test-key-not-real", http_post=post)
        return provider.plan_dynamic_template_fill(
            documents=self.documents, template_manifest=self.manifest, run_id="coverage-test",
            period="2026-08", currency="TWD", instructions=[],
            code_provenance_cells=provenance,
            code_anchored_employees=[
                {"row": row, "displayName": name, "names": [name]}
                for row, name in enumerate(["Alpha", "Bravo", "Charlie", "Delta"], 2)
            ],
        )

    def test_empty_anchored_plan_retries_with_original_facts_and_recovers(self):
        plan = self.run_plans([self.empty, self.full])
        self.assertEqual(len(self.calls), 2)
        self.assertEqual([item["value"] for item in plan["writes"]], ["300", "100", "200", "400"])
        prompt = json.loads(self.calls[0]["input"][0]["content"][-1]["text"])
        self.assertEqual(prompt["targetEmployees"][0]["targetRow"], 2)
        alpha = next(item for item in prompt["detectedSourceEmployees"] if item["names"] == ["Alpha"])
        self.assertEqual(alpha["sourceRow"], 4)
        self.assertTrue(any(fact["location"] == "Payroll!C4" and fact["value"] == "300"
                            for fact in alpha["inputFacts"]))
        self.assertFalse(any(fact["sourceLabel"] == "Service Fee" for fact in alpha["inputFacts"]))
        correction = json.loads(self.calls[1]["input"][0]["content"][-1]["text"])
        self.assertIn("no payroll data", correction["validationError"])

    def test_all_special_inputs_allow_empty_ai_plan_without_retry(self):
        reserved = [{"sheet": "Payroll-L", "row": row, "col": 3} for row in range(2, 6)]
        plan = self.run_plans([self.empty], reserved)
        self.assertEqual(plan["writes"], [])
        self.assertEqual(len(self.calls), 1)

    def test_special_writes_removed_after_retargeting_before_coverage(self):
        reserved = [{"sheet": "Payroll-L", "row": row, "col": 3} for row in range(2, 6)]
        misplaced = copy.deepcopy(self.full)
        for write in misplaced["writes"]:
            write["targetCell"] = write["targetCell"].replace("C", "D")
        plan = self.run_plans([misplaced], reserved)
        self.assertEqual(plan["writes"], [])
        self.assertEqual(len(self.calls), 1)
        self.assertTrue(any(i["code"] == "CODE_PROVENANCE_CELL_SKIPPED" for i in plan["issues"]))

    def test_special_totals_excluded_but_other_pay_fields_still_required(self):
        self.add_employee_totals()
        reserved = [{"sheet": "Payroll-L", "row": row, "col": 5} for row in range(2, 6)]
        plan = self.run_plans([self.full], reserved)
        self.assertEqual(len(plan["writes"]), 4)
        self.assertEqual(len(self.calls), 1)
        partial = copy.deepcopy(self.full)
        partial["writes"] = partial["writes"][:3]
        with self.assertRaisesRegex(ValidationError, "no payroll data.*Delta"):
            self.run_plans([partial], reserved)

    def test_same_coordinate_on_another_sheet_does_not_exempt_employee(self):
        reserved = [{"sheet": "PN", "row": row, "col": 3} for row in range(2, 6)]
        with self.assertRaisesRegex(ValidationError, "no payroll data"):
            self.run_plans([self.empty], reserved)

    def test_repeated_empty_plan_is_rejected_instead_of_successful_export(self):
        with self.assertRaisesRegex(ValidationError, "no payroll data.*Alpha.*Delta"):
            self.run_plans([self.empty])
        self.assertEqual(len(self.calls), 2)

    def test_missing_employee_data_cannot_be_hidden_by_other_employees(self):
        partial = copy.deepcopy(self.full)
        partial["writes"] = partial["writes"][:3]
        with self.assertRaisesRegex(ValidationError, "no payroll data.*Delta"):
            self.run_plans([partial])
        self.assertEqual(len(self.calls), 2)

    def test_cross_person_writes_are_filtered_then_rejected(self):
        wrong = copy.deepcopy(self.full)
        for index, item in enumerate(wrong["writes"], 2):
            item["source"]["location"] = f"Payroll!C{index}"
        with self.assertRaisesRegex(ValidationError, "no payroll data"):
            self.run_plans([wrong])

    def test_formula_only_inputs_do_not_require_ai_writes(self):
        for cell in ["C2", "C3", "C4", "C5"]:
            self.manifest["nonemptyCells"].append({
                "cell": cell, "row": int(cell[1:]), "column": 3,
                "valueKind": "formula", "value": "=0",
            })
        plan = self.run_plans([self.empty])
        self.assertEqual(plan["writes"], [])
        self.assertEqual(len(self.calls), 1)

    def test_zero_only_source_inputs_do_not_require_ai_writes(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Payroll"
        ws.append(["CN Name", "EN Name", "Basic Salary", "Service Fee"])
        for name in ["Bravo", "Charlie", "Alpha", "Delta"]:
            ws.append([None, name, 0, 50])
        wb.save(self.source)
        wb.close()
        self.documents[0]["sha256"] = hashlib.sha256(self.source.read_bytes()).hexdigest()
        plan = self.run_plans([self.empty])
        self.assertEqual(plan["writes"], [])
        self.assertEqual(len(self.calls), 1)

    def test_schema_repair_cannot_return_an_empty_success(self):
        invalid = copy.deepcopy(self.empty)
        invalid["model"] = ""
        with self.assertRaisesRegex(ValidationError, "no payroll data"):
            self.run_plans([invalid, self.empty])
        self.assertEqual(len(self.calls), 2)

    def add_employee_totals(self):
        # Total includes a service fee; skipping Service Fee must not omit Total.
        for path in [self.source, self.template]:
            wb = load_workbook(path)
            ws = wb.active
            ws["E1"] = "Total"
            if path == self.source:
                for row in range(2, 6):
                    ws.cell(row, 5, (row - 1) * 100 + 50)
            wb.save(path)
            wb.close()
        self.manifest = inspect_last_l_sheet(self.template)
        self.documents[0]["sha256"] = hashlib.sha256(self.source.read_bytes()).hexdigest()
        self.full["templateSha256"] = self.manifest["templateSha256"]
        complete = copy.deepcopy(self.full)
        for salary in self.full["writes"]:
            total = copy.deepcopy(salary)
            total["targetCell"] = salary["targetCell"].replace("C", "E")
            total["source"]["location"] = salary["source"]["location"].replace("!C", "!E")
            total["semanticLabel"] = total["sourceLabel"] = "Total"
            total["value"] = str(int(salary["value"]) + 50)
            total["source"]["rawText"] = "Total: " + total["value"]
            complete["writes"].append(total)
        return complete

    def test_salary_only_plan_retries_to_fill_each_employees_total(self):
        complete = self.add_employee_totals()
        plan = self.run_plans([self.full, complete])
        self.assertEqual(len(self.calls), 2)
        totals = {item["targetCell"]: item["value"] for item in plan["writes"]
                  if item["semanticLabel"] == "Total"}
        self.assertEqual(totals, {"E2": "350", "E3": "150", "E4": "250", "E5": "450"})
        self.assertFalse(any(item["semanticLabel"] == "Service Fee" for item in plan["writes"]))
        correction = json.loads(self.calls[1]["input"][0]["content"][-1]["text"])
        self.assertIn("Total -> E2", correction["validationError"])
        self.assertIn("Total -> E5", correction["validationError"])

    def test_repeated_missing_total_is_rejected_despite_salary_writes(self):
        self.add_employee_totals()
        with self.assertRaisesRegex(ValidationError, "omitted available source fields.*Total -> E2"):
            self.run_plans([self.full])
        self.assertEqual(len(self.calls), 2)

    def test_existing_total_formula_is_not_required_as_ai_input(self):
        self.add_employee_totals()
        for row in range(2, 6):
            self.manifest["nonemptyCells"].append({
                "cell": f"E{row}", "row": row, "column": 5,
                "valueKind": "formula", "value": f"=C{row}+D{row}",
            })
        plan = self.run_plans([self.full])
        self.assertEqual(len(self.calls), 1)
        self.assertEqual(len(plan["writes"]), 4)


if __name__ == "__main__":
    unittest.main()
