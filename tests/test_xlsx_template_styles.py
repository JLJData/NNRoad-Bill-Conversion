from copy import copy
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from xlsx_postprocess import postprocess_converted_xlsx
from xlsx_template_styles import load_conversion_template, materialize_inherited_fills


ROOT = Path(__file__).resolve().parents[1]
BLUE = PatternFill('solid', fgColor='FFDDEBF7')
GREEN = PatternFill('solid', fgColor='FFE2EFDA')
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


class TemplateFillTests(unittest.TestCase):
    def test_omitted_cells_inherit_row_before_employee_row_copy(self):
        from bill_convert.uk_layout import expand_uk_employee_rows

        wb = load_conversion_template(ROOT / 'templates/uk/template.xlsx', rich_text=True)
        try:
            ee = wb['UK EE']
            expected = copy(ee.row_dimensions[9].fill)
            self.assertEqual(ee['D9'].fill, expected)
            expand_uk_employee_rows(wb, 2, ['UK-L', 'UK-L (2)'])
            ee['D9'] = 'CUS1510-0001'
            ee['D10'] = 'CUS1510-0002'
            self.assertEqual(ee['D9'].fill, expected)
            self.assertEqual(ee['D10'].fill, expected)
        finally:
            wb.close()

    def test_explicit_white_no_fill_and_merged_cells_are_preserved(self):
        wb = Workbook()
        ws = wb.active
        ws.row_dimensions[2].fill = BLUE
        ws['B2'].fill = PatternFill('solid', fgColor='FFFFFFFF')
        ws['C2'].fill = PatternFill()
        ws['C2'].number_format = '@'
        ws.merge_cells('D2:E2')
        ws['F3'] = '=1+2'
        ws['F3'].number_format = '0.00'
        materialize_inherited_fills(wb)
        self.assertEqual(ws['A2'].fill, BLUE)
        self.assertEqual(ws['B2'].fill.fgColor.rgb, 'FFFFFFFF')
        self.assertIsNone(ws['C2'].fill.fill_type)
        self.assertEqual(ws['C2'].number_format, '@')
        self.assertIsNone(ws['E2'].fill.fill_type)
        self.assertEqual(ws['F3'].value, '=1+2')
        self.assertEqual(ws['F3'].number_format, '0.00')
        self.assertEqual(materialize_inherited_fills(wb), 0)

    def test_grouped_column_fill_row_priority_and_bounded_area(self):
        wb = Workbook()
        ws = wb.active
        ws.column_dimensions.group('B', 'D')
        ws.column_dimensions['B'].fill = GREEN
        ws.row_dimensions[2].fill = BLUE
        ws['E3'] = 12.345
        materialize_inherited_fills(wb)
        for col in 'BCD':
            self.assertEqual(ws[f'{col}1'].fill, GREEN)
            self.assertEqual(ws[f'{col}2'].fill, BLUE)
        self.assertEqual((ws.max_row, ws.max_column), (3, 5))
        self.assertEqual(ws['E3'].value, 12.345)

    def test_region_templates_keep_existing_values_formulas_and_styles(self):
        for template in (ROOT / 'templates').glob('*/template.xlsx'):
            with self.subTest(region=template.parent.name):
                wb = load_workbook(template, rich_text=True)
                try:
                    original = {
                        (ws.title, pos): (copy(cell.value), copy(cell._style))
                        for ws in wb for pos, cell in ws._cells.items()
                    }
                    materialize_inherited_fills(wb)
                    for (name, pos), (value, style) in original.items():
                        cell = wb[name]._cells[pos]
                        self.assertEqual(cell.value, value, (name, pos))
                        if style is not None:
                            self.assertEqual(cell._style, style, (name, pos))
                finally:
                    wb.close()

    def test_final_date_save_keeps_materialized_theme_and_apply_fill(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / 'result.xlsx'
            wb = Workbook()
            ws = wb.active
            ws['A1'] = '2026-10-01'
            ws['A1'].number_format = 'yyyy/m/d'
            ws['A1'].fill = BLUE
            wb.save(path)
            wb.close()
            stats = postprocess_converted_xlsx(path)
            self.assertGreater(stats['date_cells'], 0)
            with ZipFile(path) as zf:
                styles = ET.fromstring(zf.read('xl/styles.xml'))
                sheet = ET.fromstring(zf.read('xl/worksheets/sheet1.xml'))
                cell = sheet.find('s:sheetData/s:row/s:c', NS)
                xf = styles.find('s:cellXfs', NS)[int(cell.get('s'))]
                self.assertEqual(xf.get('applyFill'), '1')

    def test_uk_conversion_saves_both_employee_codes_with_blue_fill(self):
        from convert_runner import run_convert
        from profiles.uk_payroll_calc import convert as uk

        employees = [{'employee_name': name, 'amounts': {'Gross Salary': 1234.56}}
                     for name in ('Adrian Grace', 'Darren Staples')]
        directory = [{'employee_name_en': emp['employee_name'], 'employee_code': f'CUS1510-000{i}'}
                     for i, emp in enumerate(employees, 1)]
        template = ROOT / 'templates/uk/template.xlsx'
        with TemporaryDirectory() as tmp, patch.object(uk, 'parse_source', return_value=employees):
            output = Path(tmp) / 'converted.xlsx'
            result = run_convert(
                engine_id='uk_payroll_calc', source_path=template,
                output_path=output, template_path=template, employee_directory=directory,
                convert_mapping={'fxPolicy': {'mode': 'none'}, 'outstandingPayment': {'balance': 0}},
            )
            self.assertIn('outstanding_payment', result)
            saved = load_workbook(output)
            try:
                ws = saved['UK EE']
                for row in (9, 10):
                    self.assertEqual(ws.cell(row, 4).value, f'CUS1510-000{row - 8}')
                    self.assertEqual(copy(ws.cell(row, 4).fill), copy(ws.cell(row, 5).fill))
                    self.assertEqual(ws.cell(row, 4).fill.fill_type, 'solid')
            finally:
                saved.close()
            with ZipFile(output) as zf:
                styles = ET.fromstring(zf.read('xl/styles.xml'))
                sheet = ET.fromstring(zf.read('xl/worksheets/sheet3.xml'))
                for cell in sheet.findall('s:sheetData/s:row/s:c', NS):
                    if cell.get('r') in ('D9', 'D10'):
                        xf = styles.find('s:cellXfs', NS)[int(cell.get('s'))]
                        self.assertEqual(xf.get('applyFill'), '1')


if __name__ == '__main__':
    unittest.main()
