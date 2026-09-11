# -*- coding: utf-8 -*-
"""
按 mapping.lSheetCopies 把供应商横表员工行拷到母版 L 表。

不改 PN / 地区主表公式。列名默认按资格化表头（父/子）对齐，columnRename 可选。
"""
from __future__ import annotations

import re
from typing import Any

from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.worksheet import Worksheet

from bill_convert.convert_checks import check_column_rename_hits
from bill_convert.headers import list_qualified_header_cells
from convert_mapping import find_sheet_name
from xlsx_convert_utils import clean_value, norm

_TOTAL_RE = re.compile(r"^(sum|total|合计)\b", re.IGNORECASE)
MAX_EMPLOYEES = 40


def _fold(text: Any) -> str:
    return re.sub(r"\s+", " ", norm(text)).strip().lower()


def _is_formula(value: Any) -> bool:
    if isinstance(value, ArrayFormula):
        return True
    return isinstance(value, str) and value.startswith("=")


def _is_total_label(value: Any) -> bool:
    return bool(_TOTAL_RE.match(_fold(value)))


def _header_index(ws: Worksheet, header_row: int, parent_row: int | None) -> dict[str, dict[str, Any]]:
    cells = list_qualified_header_cells(ws, header_row, parent_row=parent_row)
    out: dict[str, dict[str, Any]] = {}
    for h in cells:
        item = {
            "col": int(h["col"]),
            "key": str(h["key"]),
            "child": str(h.get("child") or ""),
            "parent": str(h.get("parent") or ""),
        }
        out[str(h["key"])] = item
        folded = _fold(h["key"])
        out.setdefault(folded, item)
        child = str(h.get("child") or "")
        if child:
            out.setdefault(child, item)
            out.setdefault(_fold(child), item)
    return out


def _lookup_col(index: dict[str, dict[str, Any]], name: str) -> int | None:
    if not name:
        return None
    hit = index.get(name) or index.get(_fold(name))
    return int(hit["col"]) if hit else None


def _name_col(index: dict[str, dict[str, Any]], name_headers: list[str]) -> int | None:
    for label in name_headers:
        col = _lookup_col(index, label)
        if col:
            return col
    for fallback in ("Name 姓名", "Name", "姓名", "Name of Employee"):
        col = _lookup_col(index, fallback)
        if col:
            return col
    return None


def _read_employees(
    ws: Worksheet,
    *,
    header_row: int,
    parent_row: int | None,
    data_start: int,
    name_headers: list[str],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    qualified = list_qualified_header_cells(ws, header_row, parent_row=parent_row)
    index = _header_index(ws, header_row, parent_row)
    name_col = _name_col(index, name_headers)
    if name_col is None:
        raise ValueError(f"「{ws.title}」第 {header_row} 行未找到姓名列（{name_headers or 'Name/姓名'}）")

    employees: list[dict[str, Any]] = []
    for row in range(data_start, (ws.max_row or data_start) + 1):
        raw_name = ws.cell(row, name_col).value
        name_s = norm(raw_name)
        if not name_s:
            marker = norm(ws.cell(row, 1).value)
            if _is_total_label(marker) or _is_total_label(raw_name):
                break
            continue
        if _is_total_label(name_s) or _is_total_label(ws.cell(row, 1).value):
            break
        record: dict[str, Any] = {"_name": name_s, "_row": row}
        for h in qualified:
            val = clean_value(ws.cell(row, int(h["col"])).value)
            record[str(h["key"])] = val
            child = str(h.get("child") or "")
            if child and child not in record:
                record[child] = val
        employees.append(record)
        if len(employees) >= MAX_EMPLOYEES:
            break
    return employees, qualified


def _copy_meta_row(src: Worksheet, dst: Worksheet, spec: dict[str, Any] | None) -> None:
    if not isinstance(spec, dict):
        return
    src_row = spec.get("sourceRow") or spec.get("fromRow")
    dst_row = spec.get("targetRow") or spec.get("toRow")
    if src_row is None or dst_row is None:
        return
    src_row_i = int(src_row)
    dst_row_i = int(dst_row)
    max_col = max(int(src.max_column or 1), int(dst.max_column or 1))
    for col in range(1, max_col + 1):
        val = src.cell(src_row_i, col).value
        if val is None or val == "":
            continue
        if _is_formula(val):
            continue
        cell = dst.cell(dst_row_i, col)
        if _is_formula(cell.value):
            continue
        cell.value = val


def _clear_employee_slots(
    ws: Worksheet,
    *,
    data_start: int,
    name_col: int,
    max_col: int,
) -> int:
    """清空数据区直到合计行（不含合计行），返回可写槽位数。"""
    slots = 0
    last = (ws.max_row or data_start) + 5
    for row in range(data_start, last + 1):
        if _is_total_label(ws.cell(row, 1).value) or _is_total_label(ws.cell(row, name_col).value):
            break
        slots += 1
        if slots > MAX_EMPLOYEES:
            break
        for col in range(1, max_col + 1):
            cell = ws.cell(row, col)
            if _is_formula(cell.value):
                continue
            cell.value = None
    return max(slots, 1)


def _value_for_target(
    emp: dict[str, Any],
    tgt_key: str,
    tgt_child: str,
    target_to_source: dict[str, str],
) -> Any:
    for cand in (
        target_to_source.get(tgt_key),
        tgt_key,
        tgt_child,
        target_to_source.get(tgt_child or ""),
    ):
        if cand and cand in emp and emp.get(cand) is not None:
            return emp.get(cand)
        folded = _fold(cand) if cand else ""
        if folded:
            for k, v in emp.items():
                if k.startswith("_"):
                    continue
                if _fold(k) == folded and v is not None:
                    return v
    return None


def apply_l_sheet_copy(
    src_wb,
    dst_wb,
    copy_spec: dict[str, Any],
    *,
    column_rename: dict[str, Any] | None = None,
) -> dict[str, Any]:
    src_spec = copy_spec.get("source") if isinstance(copy_spec.get("source"), dict) else {}
    tgt_spec = copy_spec.get("target") if isinstance(copy_spec.get("target"), dict) else {}
    src_name = find_sheet_name(list(src_wb.sheetnames), src_spec)
    if not src_name:
        want = str(src_spec.get("sheet") or "")
        raise ValueError(f"源账单未找到 sheet「{want}」，现有: {list(src_wb.sheetnames)}")
    tgt_name = find_sheet_name(list(dst_wb.sheetnames), tgt_spec)
    if not tgt_name:
        want = str(tgt_spec.get("sheet") or "")
        raise ValueError(f"母版未找到 sheet「{want}」，现有: {list(dst_wb.sheetnames)}")

    src_ws: Worksheet = src_wb[src_name]
    dst_ws: Worksheet = dst_wb[tgt_name]
    src_header = int(src_spec.get("headerRow") or 7)
    src_parent = src_spec.get("parentHeaderRow")
    src_parent_i = int(src_parent) if src_parent is not None else (src_header - 1 if src_header > 1 else None)
    src_start = int(src_spec.get("dataStartRow") or (src_header + 1))
    name_headers = [str(x).strip() for x in (src_spec.get("nameHeaders") or []) if str(x).strip()]

    tgt_header = int(tgt_spec.get("headerRow") or 4)
    tgt_parent = tgt_spec.get("parentHeaderRow")
    tgt_parent_i = int(tgt_parent) if tgt_parent is not None else (tgt_header - 1 if tgt_header > 1 else None)
    tgt_start = int(tgt_spec.get("dataStartRow") or (tgt_header + 1))

    employees, src_qualified = _read_employees(
        src_ws,
        header_row=src_header,
        parent_row=src_parent_i,
        data_start=src_start,
        name_headers=name_headers,
    )
    if not employees:
        raise ValueError(f"「{src_name}」未找到有效员工行")

    rename = column_rename if isinstance(column_rename, dict) else {}
    if rename:
        sample_keys = {str(h["key"]) for h in src_qualified}
        sample_keys.update(str(h.get("child") or "") for h in src_qualified)
        check_column_rename_hits(rename, sample_keys, strict_if_configured=True)
    target_to_source = {
        str(v).strip(): str(k).strip()
        for k, v in rename.items()
        if k and v and str(k).strip() != str(v).strip()
    }

    tgt_qualified = list_qualified_header_cells(dst_ws, tgt_header, parent_row=tgt_parent_i)
    if not tgt_qualified:
        raise ValueError(f"「{tgt_name}」第 {tgt_header} 行表头为空")
    tgt_index = _header_index(dst_ws, tgt_header, tgt_parent_i)
    tgt_name_col = _name_col(tgt_index, name_headers) or int(tgt_qualified[0]["col"])
    max_col = max(int(h["col"]) for h in tgt_qualified)
    slots = _clear_employee_slots(dst_ws, data_start=tgt_start, name_col=tgt_name_col, max_col=max_col)
    if len(employees) > slots:
        raise ValueError(
            f"「{tgt_name}」员工 {len(employees)} 人超过母版数据槽位 {slots}（合计行之前）"
        )

    for idx, emp in enumerate(employees):
        row = tgt_start + idx
        for h in tgt_qualified:
            val = _value_for_target(
                emp,
                str(h["key"]),
                str(h.get("child") or ""),
                target_to_source,
            )
            if val is None:
                continue
            cell = dst_ws.cell(row, int(h["col"]))
            if _is_formula(cell.value):
                continue
            cell.value = val

    _copy_meta_row(
        src_ws,
        dst_ws,
        copy_spec.get("metaRowCopy") or copy_spec.get("durationCopy"),
    )
    return {
        "sourceSheet": src_name,
        "targetSheet": tgt_name,
        "employeeCount": len(employees),
        "employeeNames": [str(e.get("_name") or "") for e in employees],
    }


def apply_l_sheet_copies(
    src_wb,
    dst_wb,
    mapping: dict[str, Any],
) -> list[dict[str, Any]]:
    copies = mapping.get("lSheetCopies")
    if not isinstance(copies, list) or not copies:
        raise ValueError("mapping.lSheetCopies 为空")
    rename = mapping.get("columnRename") if isinstance(mapping.get("columnRename"), dict) else {}
    results: list[dict[str, Any]] = []
    for spec in copies:
        if not isinstance(spec, dict):
            continue
        results.append(apply_l_sheet_copy(src_wb, dst_wb, spec, column_rename=rename))
    if not results:
        raise ValueError("mapping.lSheetCopies 没有有效项")
    return results
