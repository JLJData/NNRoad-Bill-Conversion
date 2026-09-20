# -*- coding: utf-8 -*-
"""
Indonesia：Link Compliance 工资明细 Excel（或已是 Indonesia-L）→ Indonesia PN

用法:
  python -m profiles.indonesia_payroll_calc.convert <源.xlsx> [-o 输出.xlsx] [-t 母版.xlsx]

原则：
- PN / Indonesia / Indonesia EE 以母版公式为准；只写 Indonesia-L 数据与必要元数据。
- EE Code 必须来自员工库匹配，禁止沿用供应商账单工号。
- Cash Advance：从 Link Compliance Tax Invoice PDF 旁路抽取，写入 PN!A18/E18（见 vendor_plugins.link_compliance_cash_advance）。
"""
from __future__ import annotations

import argparse
import calendar
import re
import shutil
import sys
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.worksheet import Worksheet

from bill_convert.convert_checks import (
    check_column_rename_hits,
    merge_warnings,
    parse_cell_ref,
    sanity_check_convert_result,
)
from bill_convert.formula_copy import shift_row_formula
from bill_convert.formula_layout import sort_employees_by_code
from bill_convert.headers import list_qualified_header_cells
from convert_mapping import find_sheet_name, resolve_convert_mapping
from fx_rate import get_indonesia_pn_fx_rate
from pn_meta import PnMeta, apply_pn_meta
from profiles.tw_payroll_calc.convert import match_ee_code
from region_templates import get_region_template
from xlsx_convert_utils import coerce_datetime_for_excel, norm
from xlsx_luckysheet_compat import apply_luckysheet_compat
from xlsx_postprocess import postprocess_converted_xlsx

DEFAULT_TEMPLATE = get_region_template("Indonesia")

INDONESIA_L_SHEET = "Indonesia-L"
INDONESIA_SHEET = "Indonesia"
INDONESIA_EE_SHEET = "Indonesia EE"
PN_SHEET = "PN"

INDONESIA_L_HEADER_ROW = 7
INDONESIA_L_DATA_START = 8
INDONESIA_DATA_START = 9
INDONESIA_EE_DATA_START = 10
MAX_EMPLOYEES = 20
_DATE_FMT = "yyyy/m/d"

COL_EE_CODE = 1
COL_NAME = 2

# Indonesia-L 目标表头（写盘键）；同名自动配须与源资格化 key 完全一致
_TARGET_L_HEADERS: tuple[str, ...] = (
    "No. of EE",
    "Name of Employee",
    "Base Salary",
    "OT",
    "Salary Adjusment",
    "Bonus",
    "Other",
    "Other ",
    "Expense Reimbursment",
    "Gross Salary",
    "INCOME TAX (PPH21)",
    "JHT 2%",
    "JP 1%",
    "BPJS KESEHATAN 1%",
    "Net Salary",
    "JKK 0.24%",
    "JHT 3.7%",
    "JP 2%",
    "JKM  0.3%",
    "HEALTH INSURANCE (BPJS KESEHATAN) 4%",
    "Total Cost",
)

# 母版公式列 / EE Code：不从源表写入
_SKIP_WRITE_TARGETS = frozenset(
    {"No. of EE", "Gross Salary", "JHT 2%", "Net Salary", "Total Cost"}
)

# 缺省写 0（样例 PN：Adj/Bonus/Other/Expense）
_ZERO_FILL_KEYS = ("Salary Adjusment", "Bonus", "Other ", "Expense Reimbursment")

_ACTIVE_MAPPING: dict[str, Any] | None = None

_MONTHS = {
    "jan": 1,
    "january": 1,
    "feb": 2,
    "february": 2,
    "mar": 3,
    "march": 3,
    "apr": 4,
    "april": 4,
    "may": 5,
    "jun": 6,
    "june": 6,
    "jul": 7,
    "july": 7,
    "aug": 8,
    "august": 8,
    "sep": 9,
    "sept": 9,
    "september": 9,
    "oct": 10,
    "october": 10,
    "nov": 11,
    "november": 11,
    "dec": 12,
    "december": 12,
}


def _active_mapping() -> dict[str, Any]:
    return (
        _ACTIVE_MAPPING
        if isinstance(_ACTIVE_MAPPING, dict)
        else resolve_convert_mapping("indonesia_payroll_calc", None)
    )


def _norm(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\n", " ").strip()
    # 源表头常见 =UPPER("Employee Name")
    m = re.match(r'^=\s*UPPER\s*\(\s*"([^"]+)"\s*\)\s*$', text, flags=re.I)
    if m:
        return m.group(1).strip()
    m2 = re.match(r"^=\s*UPPER\s*\(\s*'([^']+)'\s*\)\s*$", text, flags=re.I)
    if m2:
        return m2.group(1).strip()
    return text


def _as_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("\xa0", "").replace("IDR", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _cell_formula_text(value: Any) -> str | None:
    if isinstance(value, ArrayFormula):
        return value.text
    if isinstance(value, str) and value.startswith("="):
        return value
    return None


def _set_cell_value(cell, value: Any) -> None:
    if isinstance(value, float):
        cell.value = round(value, 6) if abs(value) >= 1e-9 else 0
        if cell.number_format in (None, "General", ""):
            pass
        elif "yy" in str(cell.number_format).lower() or "h:" in str(cell.number_format).lower():
            cell.number_format = "General"
    elif isinstance(value, int):
        cell.value = value
        fmt = str(cell.number_format or "")
        if "yy" in fmt.lower() or "h:" in fmt.lower() or "m/d" in fmt.lower():
            cell.number_format = "General"
    else:
        cell.value = value


def _indonesia_l_layout(*, target: bool = False) -> tuple[int, int]:
    mapping = _active_mapping()
    key = "targetL" if target else "sourceEmployeeSheet"
    spec = mapping.get(key) if isinstance(mapping.get(key), dict) else {}
    if not spec and target:
        spec = (
            mapping.get("sourceEmployeeSheet")
            if isinstance(mapping.get("sourceEmployeeSheet"), dict)
            else {}
        )
    if target:
        header = int(spec.get("headerRow") or INDONESIA_L_HEADER_ROW)
        data_start = int(spec.get("dataStartRow") or INDONESIA_L_DATA_START)
    else:
        header = int(spec.get("headerRow") or 7)
        data_start = int(spec.get("dataStartRow") or 8)
    return header, data_start


def _source_parent_row(header_row: int) -> int | None:
    mapping = _active_mapping()
    spec = (
        mapping.get("sourceEmployeeSheet")
        if isinstance(mapping.get("sourceEmployeeSheet"), dict)
        else {}
    )
    raw = spec.get("parentHeaderRow")
    if raw is not None:
        try:
            return int(raw)
        except (TypeError, ValueError):
            pass
    return header_row - 1 if header_row > 1 else None


def _column_rename_map() -> dict[str, str]:
    """仅 mapping.columnRename（引擎默认已写入 ENGINE_DEFAULTS，前端可见）。"""
    raw = _active_mapping().get("columnRename") or {}
    if not isinstance(raw, dict):
        return {}
    return {norm(str(k)): str(v).strip() for k, v in raw.items() if k and str(v).strip()}


def _skip_source_headers() -> set[str]:
    raw = _active_mapping().get("skipSourceHeaders") or []
    if not isinstance(raw, list):
        return set()
    return {norm(x) for x in raw if x}


def _target_header_keys() -> set[str]:
    return {norm(x) for x in _TARGET_L_HEADERS if x}


def map_source_header(source_header: str) -> str | None:
    """
    源资格化 key → Indonesia-L 表头。
    1) columnRename 显式对照（须在默认/前端映射里，禁止静默别名）
    2) 同名：资格化 key 与目标表头完全一致（禁止只比子列名）
    """
    h = norm(source_header)
    if not h or h in _skip_source_headers():
        return None
    rename = _column_rename_map()
    if h in rename:
        return rename[h]
    if h in _target_header_keys():
        for t in _TARGET_L_HEADERS:
            if norm(t) == h:
                return t
    return None


def _meta_cell(key: str, default_row: int, default_col: int) -> tuple[int, int]:
    mapping = _active_mapping()
    meta = mapping.get("metaCells") if isinstance(mapping.get("metaCells"), dict) else {}
    return parse_cell_ref(meta.get(key), default_row=default_row, default_col=default_col)


def parse_pay_period_label(label: Any) -> tuple[int, int] | None:
    text = _norm(label)
    if not text:
        return None
    m = re.search(
        r"(jan|january|feb|february|mar|march|apr|april|may|jun|june|jul|july|"
        r"aug|august|sep|sept|september|oct|october|nov|november|dec|december)"
        r"[\s\-']*(\d{2,4})",
        text,
        flags=re.I,
    )
    if not m:
        m2 = re.search(r"(\d{1,2})\s*/\s*(\d{4})", text)
        if m2:
            return int(m2.group(2)), int(m2.group(1))
        return None
    month = _MONTHS.get(m.group(1).lower())
    if not month:
        return None
    year = int(m.group(2))
    if year < 100:
        year += 2000
    return year, month


def period_bounds_from_label(label: Any) -> tuple[date, date] | None:
    """样例 PN 使用整月起止（非 7–30 日截段）。"""
    parsed = parse_pay_period_label(label)
    if not parsed:
        return None
    year, month = parsed
    start = date(year, month, 1)
    end = date(year, month, calendar.monthrange(year, month)[1])
    return start, end


def _read_vendor_fx(ws: Worksheet) -> float | None:
    for row in range(1, min((ws.max_row or 1), 10) + 1):
        for col in range(1, min((ws.max_column or 1), 6) + 1):
            label = _norm(ws.cell(row, col).value).lower()
            if "exchange rate" in label or "1usd" in label.replace(" ", ""):
                for dc in (col + 1, col + 2, 3):
                    fx = _as_float(ws.cell(row, dc).value)
                    if fx is not None and fx > 0:
                        return fx
    return None


def _read_vendor_period_label(ws: Worksheet) -> str | None:
    for row in range(1, min((ws.max_row or 1), 8) + 1):
        for col in range(1, min((ws.max_column or 1), 6) + 1):
            text = _norm(ws.cell(row, col).value)
            if text.lower().startswith("period") or "september" in text.lower() or re.search(
                r"\b20\d{2}\b", text
            ):
                if "period" in text.lower() or parse_pay_period_label(text):
                    return text
    return None


def _indonesia_l_formula_cols(ws: Worksheet, data_start: int) -> dict[int, str]:
    out: dict[int, str] = {}
    for col in range(1, (ws.max_column or 1) + 1):
        text = _cell_formula_text(ws.cell(data_start, col).value)
        if text:
            out[col] = text
    return out


def looks_like_indonesia_l(ws: Worksheet) -> bool:
    qualified = list_qualified_header_cells(ws, INDONESIA_L_HEADER_ROW)
    keys = {norm(str(h.get("key") or "")) for h in qualified}
    return "name of employee" in keys and "base salary" in keys


def _unwrap_header_key(value: Any) -> str:
    """表头格可能是 =UPPER(\"Employee Name\")；资格化前先还原。"""
    return _norm(value)


def _qualified_source_headers(
    ws: Worksheet, header_row: int, parent_row: int | None
) -> dict[str, int]:
    """资格化表头 → 列号；并解开 UPPER() 公式表头。"""
    # 临时把 UPPER 公式写成明文再资格化，避免 key 带公式
    touched: list[tuple[int, int, Any]] = []
    rows = [header_row]
    if parent_row:
        rows.append(parent_row)
    for r in rows:
        for c in range(1, (ws.max_column or 1) + 1):
            cell = ws.cell(r, c)
            raw = cell.value
            if isinstance(raw, str) and raw.startswith("=") and "UPPER" in raw.upper():
                plain = _unwrap_header_key(raw)
                if plain and plain != raw:
                    touched.append((r, c, raw))
                    cell.value = plain
    try:
        qualified = list_qualified_header_cells(ws, header_row, parent_row=parent_row)
        return {str(h["key"]): int(h["col"]) for h in qualified if h.get("key") and h.get("col")}
    finally:
        for r, c, raw in touched:
            ws.cell(r, c).value = raw


def _parse_employees_from_sheet(
    ws: Worksheet,
    *,
    header_row: int,
    parent_row: int | None,
    data_start: int,
    warnings: list[str] | None = None,
    include_period_meta: bool = True,
) -> list[dict[str, Any]]:
    """资格化表头 + columnRename / 严格同名 → 员工行。"""
    source_headers = _qualified_source_headers(ws, header_row, parent_row)
    rename = _column_rename_map()
    if rename:
        check_column_rename_hits(rename, source_headers, strict_if_configured=False)

    name_col: int | None = None
    for src_key, col in source_headers.items():
        tgt = map_source_header(src_key)
        if tgt and norm(tgt) == norm("Name of Employee"):
            name_col = col
            break
    if name_col is None:
        src_spec = (
            _active_mapping().get("sourceEmployeeSheet")
            if isinstance(_active_mapping().get("sourceEmployeeSheet"), dict)
            else {}
        )
        for nh in src_spec.get("nameHeaders") or []:
            key = norm(nh)
            for sk, col in source_headers.items():
                if norm(sk) == key:
                    name_col = col
                    break
            if name_col is not None:
                break
    if name_col is None:
        raise ValueError(
            f"未找到员工姓名列（请在 columnRename 配置 EMPLOYEE NAME→Name of Employee）。"
            f"表头={list(source_headers.keys())[:12]}"
        )

    period_from = period_to = None
    period_label = None
    fx_rate = None
    if include_period_meta:
        period_label = _read_vendor_period_label(ws)
        bounds = period_bounds_from_label(period_label) if period_label else None
        period_from = bounds[0] if bounds else None
        period_to = bounds[1] if bounds else None
        fx_rate = _read_vendor_fx(ws)
        (pf_r, pf_c) = _meta_cell("periodFrom", 2, 3)
        (pt_r, pt_c) = _meta_cell("periodTo", 2, 5)
        (fx_r, fx_c) = _meta_cell("fxRate", 4, 3)
        if ws.cell(pf_r, pf_c).value is not None:
            period_from = period_from or ws.cell(pf_r, pf_c).value
        if ws.cell(pt_r, pt_c).value is not None:
            period_to = period_to or ws.cell(pt_r, pt_c).value
        if fx_rate is None:
            fx_rate = _as_float(ws.cell(fx_r, fx_c).value)

    employees: list[dict[str, Any]] = []
    max_row = max(ws.max_row or data_start, data_start)
    for row in range(data_start, max_row + 1):
        name = _norm(ws.cell(row, name_col).value)
        if not name or name.upper() == "TOTAL":
            continue
        if _cell_formula_text(ws.cell(row, name_col).value):
            continue
        emp: dict[str, Any] = {
            "Employee Name": name,
            "Name of Employee": name,
            "_period_from": period_from,
            "_period_to": period_to,
            "_period_label": period_label,
            "_fx_rate": fx_rate,
            "From": period_from,
            "To": period_to,
        }
        for src_key, col in source_headers.items():
            target = map_source_header(src_key)
            if not target or norm(target) in {norm(x) for x in _SKIP_WRITE_TARGETS}:
                continue
            if norm(target) == norm("Name of Employee"):
                continue
            val = ws.cell(row, col).value
            if _cell_formula_text(val) or val is None or val == "":
                continue
            num = _as_float(val)
            emp[target] = num if num is not None else val
        for zkey in _ZERO_FILL_KEYS:
            emp.setdefault(zkey, 0)
        employees.append(emp)
    if not employees and warnings is not None:
        warnings.append("未解析到员工行")
    return employees


def parse_link_compliance_employees(ws: Worksheet, warnings: list[str] | None = None) -> list[dict[str, Any]]:
    header_row, data_start = _indonesia_l_layout(target=False)
    return _parse_employees_from_sheet(
        ws,
        header_row=header_row,
        parent_row=_source_parent_row(header_row),
        data_start=data_start,
        warnings=warnings,
        include_period_meta=True,
    )


def parse_indonesia_l_employees(ws: Worksheet, warnings: list[str] | None = None) -> list[dict[str, Any]]:
    header_row, data_start = _indonesia_l_layout(target=True)
    return _parse_employees_from_sheet(
        ws,
        header_row=header_row,
        parent_row=None,
        data_start=data_start,
        warnings=warnings,
        include_period_meta=True,
    )


def _enrich_employees_from_data_only(
    source_path: Path,
    employees: list[dict[str, Any]],
    sheet_name: str,
) -> list[dict[str, Any]]:
    """用 Excel 缓存值补全公式列（如 BPJS 4%）。"""
    if not employees:
        return employees
    try:
        wb_val = load_workbook(source_path, data_only=True)
    except Exception:
        return employees
    try:
        if sheet_name not in wb_val.sheetnames:
            return employees
        ws = wb_val[sheet_name]
        header_row, data_start = _indonesia_l_layout(target=False)
        parent_row = _source_parent_row(header_row)
        source_headers = _qualified_source_headers(ws, header_row, parent_row)
        name_col = None
        for src_key, col in source_headers.items():
            tgt = map_source_header(src_key)
            if tgt and norm(tgt) == norm("Name of Employee"):
                name_col = col
                break
            if norm(src_key) in ("employee name", "name of employee"):
                name_col = col
                break
        if name_col is None:
            return employees
        fx_rate = _read_vendor_fx(ws)
        idx = 0
        max_row = max(ws.max_row or data_start, data_start)
        for row in range(data_start, max_row + 1):
            name = _norm(ws.cell(row, name_col).value)
            if not name or name.upper() == "TOTAL":
                continue
            if idx >= len(employees):
                break
            emp = employees[idx]
            if fx_rate is not None:
                emp["_fx_rate"] = fx_rate
            for src_key, col in source_headers.items():
                target = map_source_header(src_key)
                if not target or norm(target) in {norm(x) for x in _SKIP_WRITE_TARGETS}:
                    continue
                if _as_float(emp.get(target)) is not None:
                    continue
                num = _as_float(ws.cell(row, col).value)
                if num is not None:
                    emp[target] = num
            idx += 1
        return employees
    finally:
        wb_val.close()


def parse_source_workbook(source_path: Path, warnings: list[str] | None = None) -> list[dict[str, Any]]:
    wb = load_workbook(source_path, data_only=False)
    try:
        if INDONESIA_L_SHEET in wb.sheetnames and looks_like_indonesia_l(wb[INDONESIA_L_SHEET]):
            return parse_indonesia_l_employees(wb[INDONESIA_L_SHEET], warnings)

        mapping = _active_mapping()
        src_spec = (
            mapping.get("sourceEmployeeSheet")
            if isinstance(mapping.get("sourceEmployeeSheet"), dict)
            else {}
        )
        sheet_name = find_sheet_name(list(wb.sheetnames), src_spec)
        if not sheet_name:
            sheet_name = wb.sheetnames[0] if wb.sheetnames else None
        if not sheet_name:
            raise ValueError("源工作簿无 sheet")
        employees = parse_link_compliance_employees(wb[sheet_name], warnings)
    finally:
        wb.close()

    return _enrich_employees_from_data_only(source_path, employees, sheet_name)


def write_indonesia_l_period(ws: Worksheet, employees: list[dict[str, Any]]) -> None:
    if not employees:
        return
    emp0 = employees[0]
    start = emp0.get("_period_from") or emp0.get("From")
    end = emp0.get("_period_to") or emp0.get("To")
    if start is None and end is None:
        bounds = period_bounds_from_label(emp0.get("_period_label"))
        if bounds:
            start, end = bounds
    (pf_r, pf_c) = _meta_cell("periodFrom", 2, 3)
    (pt_r, pt_c) = _meta_cell("periodTo", 2, 5)
    for (row, col), val in (((pf_r, pf_c), start), ((pt_r, pt_c), end)):
        if val is None:
            continue
        cell = ws.cell(row, col)
        dt = coerce_datetime_for_excel(val)
        if dt is not None:
            cell.value = dt
            cell.number_format = _DATE_FMT
        elif isinstance(val, str) and re.match(r"^\d{4}/\d{1,2}/\d{1,2}$", val.strip()):
            try:
                y, m, d = [int(x) for x in val.strip().split("/")]
                cell.value = datetime(y, m, d)
                cell.number_format = _DATE_FMT
            except ValueError:
                cell.value = val
        else:
            cell.value = val


def write_indonesia_l_fx(ws: Worksheet, employees: list[dict[str, Any]], fx: float | None) -> None:
    (fx_r, fx_c) = _meta_cell("fxRate", 4, 3)
    rate = fx
    if rate is None and employees:
        rate = _as_float(employees[0].get("_fx_rate"))
    if rate is None:
        return
    ws.cell(fx_r, fx_c).value = float(rate)


def write_indonesia_l(ws: Worksheet, employees: list[dict[str, Any]]) -> None:
    """只覆盖数据列，母版公式列保留/扩行复制。不写 EE Code（由员工库步骤写入）。"""
    header_row, data_start = _indonesia_l_layout(target=True)
    write_indonesia_l_period(ws, employees)

    n = len(employees)
    formula_by_col = _indonesia_l_formula_cols(ws, data_start)
    max_col = max(ws.max_column or 22, 22)

    for i in range(1, n):
        _copy_row_style_and_formula(
            ws,
            data_start,
            data_start + i,
            max_col=max_col,
            l_from=data_start,
            l_to=data_start + i,
        )

    last_keep = data_start + max(n, 1) - 1
    max_row = max(ws.max_row or data_start, data_start + MAX_EMPLOYEES)
    for row in range(data_start, max_row + 1):
        for col in range(1, max_col + 1):
            cell = ws.cell(row, col)
            if type(cell).__name__ == "MergedCell":
                continue
            if row <= last_keep and col in formula_by_col:
                continue
            cell.value = None

    for i in range(n):
        row = data_start + i
        for col, formula in formula_by_col.items():
            cur = ws.cell(row, col).value
            if _cell_formula_text(cur):
                continue
            ws.cell(row, col).value = shift_row_formula(
                formula,
                data_start,
                row,
                target_l_from=data_start,
                target_l_to=row,
                target_l_sheet=INDONESIA_L_SHEET,
            )

    headers = {
        str(h["key"]): int(h["col"])
        for h in list_qualified_header_cells(ws, header_row)
        if h.get("key") and h.get("col")
    }
    # 兼容目标表头带尾空格（Other  vs Other）
    headers_norm = {norm(k): (k, col) for k, col in headers.items()}
    name_col = COL_NAME
    if norm("Name of Employee") in headers_norm:
        name_col = headers_norm[norm("Name of Employee")][1]

    for idx, emp in enumerate(employees):
        row = data_start + idx
        name = _norm(emp.get("Employee Name") or emp.get("Name of Employee"))
        if name:
            ws.cell(row, name_col).value = name
        written: set[int] = set()
        for key, val in emp.items():
            if val is None or str(key).startswith("_"):
                continue
            if norm(key) in {norm(x) for x in _SKIP_WRITE_TARGETS} or norm(key) == norm(
                "Name of Employee"
            ):
                continue
            hit = headers_norm.get(norm(key))
            if not hit:
                continue
            _canon, col = hit
            if col in formula_by_col or col in written:
                continue
            _set_cell_value(ws.cell(row, col), val)
            written.add(col)


def _copy_row_style_and_formula(
    ws: Worksheet,
    src_row: int,
    dest_row: int,
    max_col: int = 40,
    *,
    l_from: int | None = None,
    l_to: int | None = None,
) -> None:
    for c in range(1, max_col + 1):
        src = ws.cell(src_row, c)
        dest = ws.cell(dest_row, c)
        if type(dest).__name__ == "MergedCell" or type(src).__name__ == "MergedCell":
            continue
        if src.has_style:
            dest.font = copy(src.font)
            dest.border = copy(src.border)
            dest.fill = copy(src.fill)
            dest.number_format = src.number_format
            dest.alignment = copy(src.alignment)
        text = _cell_formula_text(src.value)
        if text and l_from is not None and l_to is not None:
            dest.value = shift_row_formula(
                text,
                src_row,
                dest_row,
                target_l_from=l_from,
                target_l_to=l_to,
                target_l_sheet=INDONESIA_L_SHEET,
            )
        elif text:
            dest.value = shift_row_formula(text, src_row, dest_row)
        else:
            dest.value = None


def _retarget_l_refs(formula: str, l_from: int, l_to: int) -> str:
    if l_from == l_to:
        return formula
    return re.sub(
        rf"(Indonesia-L['\"]?!\$?[A-Za-z]+\$?){l_from}\b",
        rf"\g<1>{l_to}",
        formula,
        flags=re.I,
    )


def _retarget_ee_refs(formula: str, ee_from: int, ee_to: int) -> str:
    if ee_from == ee_to:
        return formula
    return re.sub(
        rf"(Indonesia EE['\"]?!\$?[A-Za-z]+\$?){ee_from}\b",
        rf"\g<1>{ee_to}",
        formula,
        flags=re.I,
    )


def _scan_indonesia_ee_layout(ws: Worksheet) -> tuple[int, int, int]:
    """返回 (data_start_row, ee_code_col, ee_name_col)。按表头定位，兼容母版列变动。"""
    code_col, name_col, header_row = 4, 5, None
    for r in range(1, 12):
        for c in range(1, 16):
            h = re.sub(r"\s+", " ", _norm(ws.cell(r, c).value)).lower()
            if h == "ee code":
                code_col = c
                header_row = r if header_row is None else min(header_row, r)
            elif h == "ee name":
                name_col = c
                header_row = r if header_row is None else min(header_row, r)
    start_from = (header_row or 3) + 1
    for r in range(start_from, 22):
        a = _norm(ws.cell(r, 1).value)
        if a and "eor" in a.lower():
            continue
        name_v = ws.cell(r, name_col).value
        code_v = ws.cell(r, code_col).value
        b_v = ws.cell(r, 2).value
        if _cell_formula_text(name_v) or _cell_formula_text(code_v) or _cell_formula_text(b_v):
            return r, code_col, name_col
        if _norm(name_v) or _norm(code_v):
            return r, code_col, name_col
    mapped = INDONESIA_EE_DATA_START
    ft = _active_mapping().get("formulaTemplates")
    block = (
        ft.get(INDONESIA_EE_SHEET)
        if isinstance(ft, dict) and isinstance(ft.get(INDONESIA_EE_SHEET), dict)
        else {}
    )
    if block.get("defaultExampleRow"):
        mapped = int(block["defaultExampleRow"])
    return mapped, code_col, name_col


def expand_indonesia_employee_rows(wb, employee_count: int) -> None:
    n = max(int(employee_count), 1)
    _, l_data_start = _indonesia_l_layout(target=True)
    if INDONESIA_SHEET in wb.sheetnames:
        main = wb[INDONESIA_SHEET]
        for i in range(1, n):
            dest = INDONESIA_DATA_START + i
            l_row = l_data_start + i
            ee_row = INDONESIA_EE_DATA_START + i
            _copy_row_style_and_formula(
                main,
                INDONESIA_DATA_START,
                dest,
                max_col=40,
                l_from=l_data_start,
                l_to=l_row,
            )
            for c in range(1, 41):
                cell = main.cell(dest, c)
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.value = _retarget_l_refs(cell.value, l_data_start, l_row)
                    cell.value = _retarget_ee_refs(cell.value, INDONESIA_EE_DATA_START, ee_row)

    if INDONESIA_EE_SHEET in wb.sheetnames:
        ee = wb[INDONESIA_EE_SHEET]
        for i in range(1, n):
            dest = INDONESIA_EE_DATA_START + i
            l_row = l_data_start + i
            _copy_row_style_and_formula(
                ee,
                INDONESIA_EE_DATA_START,
                dest,
                max_col=40,
                l_from=l_data_start,
                l_to=l_row,
            )
            for c in range(1, 41):
                cell = ee.cell(dest, c)
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.value = _retarget_l_refs(cell.value, l_data_start, l_row)


def _pn_customer_id(pn_meta: PnMeta | dict[str, Any] | None) -> str | None:
    if pn_meta is None:
        return None
    if isinstance(pn_meta, PnMeta):
        cid = (pn_meta.customer_id or "").strip()
        return cid or None
    if isinstance(pn_meta, dict):
        cid = str(pn_meta.get("customer_id") or pn_meta.get("customerId") or "").strip()
        return cid or None
    return None


def apply_indonesia_ee_codes(
    wb,
    employees: list[dict[str, Any]],
    *,
    employee_directory: list[dict[str, Any]] | None = None,
    pn_meta: PnMeta | dict[str, Any] | None = None,
) -> list[str]:
    """员工库匹配 EE Code → 直接写 Indonesia EE 表 EE Code 列；同步写 Indonesia-L!A。

    禁止使用供应商账单中的工号。
    """
    warnings: list[str] = []
    directory = list(employee_directory or [])
    _, l_data_start = _indonesia_l_layout(target=True)
    client_code = _pn_customer_id(pn_meta)

    ee_ws = wb[INDONESIA_EE_SHEET] if INDONESIA_EE_SHEET in wb.sheetnames else None
    ee_data_start, ee_code_col = INDONESIA_EE_DATA_START, 4
    if ee_ws is not None:
        ee_data_start, ee_code_col, _ = _scan_indonesia_ee_layout(ee_ws)

    for i, emp in enumerate(employees):
        # 清除可能误带的供应商工号键
        emp.pop("No. of EE", None)
        emp.pop("_ee_code", None)

        name = _norm(emp.get("Employee Name") or emp.get("Name of Employee"))
        code, warn = match_ee_code([name] if name else [], directory)
        if code:
            emp["No. of EE"] = code
            emp["_ee_code"] = code
            if INDONESIA_L_SHEET in wb.sheetnames:
                wb[INDONESIA_L_SHEET].cell(l_data_start + i, COL_EE_CODE).value = code
        if warn:
            warnings.append(f"Indonesia EE 第{i + 1}人：{warn}")

        if ee_ws is None:
            continue
        row = ee_data_start + i
        # Client Code：仅当母版该格不是公式时才写
        if client_code and not _cell_formula_text(ee_ws.cell(row, 2).value):
            ee_ws.cell(row, 2).value = client_code
        # 直接写 EE Code 列；母版公式格不覆盖；匹配不到显式清空
        if not _cell_formula_text(ee_ws.cell(row, ee_code_col).value):
            ee_ws.cell(row, ee_code_col).value = code
    return warnings


def _find_pn_row_by_label(ws: Worksheet, keyword: str, col: int = 1) -> int | None:
    key = keyword.lower()
    for row in range(1, (ws.max_row or 0) + 1):
        v = ws.cell(row, col).value
        if isinstance(v, str) and key in v.lower():
            return row
    return None


def fit_indonesia_pn_employees(wb, employee_count: int) -> dict[str, Any]:
    fx_row = _find_pn_row_by_label(wb[PN_SHEET], "FX rate") if PN_SHEET in wb.sheetnames else None
    return {"fx_row": fx_row or 28, "employee_count": employee_count}


def _resolve_pdf_profile_id(mapping: dict[str, Any] | None) -> str | None:
    if not isinstance(mapping, dict):
        return None
    for key in ("pdfProfileId", "_pdfProfileId"):
        val = mapping.get(key)
        if val is not None and str(val).strip():
            return str(val).strip()
    from bill_convert.fact_store import get_batch_facts

    batch = get_batch_facts(mapping)
    if any(str(k).startswith("link_compliance.") for k in batch):
        return "link_compliance_id"
    return None


def _ingest_cash_advance_pdfs(mapping: dict[str, Any] | None, pdf_paths: list[Path]) -> dict[str, Any]:
    """把 Tax Invoice PDF 解析为 artifactBatch，供插件写入 PN Cash Advance。"""
    from bill_convert.fact_store import set_batch_facts
    from bill_convert.vendor_plugins.link_compliance_cash_advance import (
        LinkComplianceCashAdvancePlugin,
    )

    out = dict(mapping) if isinstance(mapping, dict) else {}
    paths = [Path(p) for p in pdf_paths if p and Path(p).is_file()]
    if not paths:
        return out
    plugin = LinkComplianceCashAdvancePlugin()
    facts = plugin.parse_artifacts(paths) or {}
    warnings = facts.pop("_warnings", None) or []
    for w in warnings:
        print(f"[link-compliance] {w}")
    out = set_batch_facts(out, facts)
    out.setdefault("pdfProfileId", "link_compliance_id")
    return out


def _apply_vendor_plugins(
    wb, warnings: list[str], *, employee_count: int = 1
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    from bill_convert.fact_store import get_batch_facts
    from bill_convert.vendor_plugins.runtime import apply_vendor_plugins

    mapping = _active_mapping()
    pdf_profile_id = _resolve_pdf_profile_id(mapping)
    if not pdf_profile_id:
        return {}, []
    raw = (
        apply_vendor_plugins(
            wb,
            pdf_profile_id=pdf_profile_id,
            mapping=mapping,
            batch_facts=get_batch_facts(mapping),
            warnings=warnings,
            employee_count=employee_count,
        )
        or {}
    )
    cell_writes = raw.pop("_cell_writes", None)
    writes = cell_writes if isinstance(cell_writes, list) else []
    return raw, [x for x in writes if isinstance(x, dict)]


def apply_fx(
    wb,
    employees: list[dict[str, Any]],
    *,
    fill_fx: bool = True,
    convert_mapping: dict | None = None,
) -> tuple[float | None, str | None]:
    """优先供应商账单汇率写 Indonesia-L!C4；否则 API IDR。"""
    if not fill_fx or INDONESIA_L_SHEET not in wb.sheetnames:
        return None, None
    from fx_policy import fx_policy

    mapping = convert_mapping if isinstance(convert_mapping, dict) else _active_mapping()
    policy = fx_policy(mapping)
    mode = str(policy.get("mode") or "vendor_bill").strip().lower()
    if mode == "none":
        return None, None

    vendor_fx = None
    if employees:
        vendor_fx = _as_float(employees[0].get("_fx_rate"))
    fx = None
    source = None
    if mode in ("vendor_bill", "shared_fact") and vendor_fx is not None and vendor_fx > 0:
        fx = float(vendor_fx)
        source = "source:vendor"
    elif mode == "none":
        return None, None
    else:
        try:
            fx = float(get_indonesia_pn_fx_rate())
            source = "api:IDR"
        except Exception as exc:
            if vendor_fx is not None and vendor_fx > 0:
                fx = float(vendor_fx)
                source = f"source:vendor(fallback after api fail: {exc})"
            else:
                raise

    if fx is not None:
        write_indonesia_l_fx(wb[INDONESIA_L_SHEET], employees, fx)
    return fx, source


def convert(
    source_path: Path,
    output_path: Path,
    template_path: Path,
    *,
    pn_meta: PnMeta | dict[str, Any] | None = None,
    employee_directory: list[dict[str, Any]] | None = None,
    registry_dir: Path | None = None,
    convert_mapping: dict[str, Any] | None = None,
    fill_fx: bool = True,
    cash_advance_pdf: Path | None = None,
) -> dict[str, Any]:
    global _ACTIVE_MAPPING
    _ACTIVE_MAPPING = resolve_convert_mapping("indonesia_payroll_calc", convert_mapping)
    try:
        source_path = Path(source_path).resolve()
        output_path = Path(output_path).resolve()
        template_path = Path(template_path).resolve()
        if not source_path.is_file():
            raise FileNotFoundError(f"源文件不存在: {source_path}")
        if not template_path.is_file():
            raise FileNotFoundError(f"母版不存在: {template_path}")

        # 可选：本批附带 Tax Invoice PDF → 注入 Cash Advance 事实
        if cash_advance_pdf is not None:
            _ACTIVE_MAPPING = _ingest_cash_advance_pdfs(_ACTIVE_MAPPING, [cash_advance_pdf])
        else:
            side_paths = []
            if isinstance(_ACTIVE_MAPPING, dict):
                raw_paths = _ACTIVE_MAPPING.get("_artifactPdfPaths") or _ACTIVE_MAPPING.get(
                    "artifactPdfPaths"
                )
                if isinstance(raw_paths, list):
                    side_paths = [Path(p) for p in raw_paths if p]
                single = _ACTIVE_MAPPING.get("_cashAdvancePdf") or _ACTIVE_MAPPING.get(
                    "cashAdvancePdf"
                )
                if single:
                    side_paths.append(Path(str(single)))
            if side_paths:
                _ACTIVE_MAPPING = _ingest_cash_advance_pdfs(_ACTIVE_MAPPING, side_paths)

        parse_warnings: list[str] = []
        employees = parse_source_workbook(source_path, parse_warnings)
        if not employees:
            raise ValueError("未解析到任何印尼员工行")
        sort_employees_by_code(employees, employee_directory)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(template_path, output_path)
        wb = load_workbook(output_path)
        warnings: list[str] = list(parse_warnings)
        fx = None
        fx_source = None
        applied_pn = None
        plugin_cell_writes: list[dict[str, Any]] = []
        fact_store_updates: dict[str, Any] = {}
        try:
            if INDONESIA_L_SHEET not in wb.sheetnames:
                raise ValueError(f"母版缺少 {INDONESIA_L_SHEET}")
            write_indonesia_l(wb[INDONESIA_L_SHEET], employees)
            expand_indonesia_employee_rows(wb, len(employees))
            pn_layout = fit_indonesia_pn_employees(wb, len(employees))
            if len(employees) > 1:
                warnings.append(
                    f"Indonesia PN 多人明细行扩行暂定：已扩 Indonesia/Indonesia EE（{len(employees)} 人），请人工核对 PN"
                )
            try:
                fx, fx_source = apply_fx(
                    wb, employees, fill_fx=fill_fx, convert_mapping=_ACTIVE_MAPPING
                )
            except Exception as exc:
                warnings.append(f"写入汇率失败: {exc}")

            if pn_meta is not None:
                applied_pn = apply_pn_meta(
                    wb,
                    pn_meta,
                    registry_dir=registry_dir or output_path.parent,
                    reserve_invoice_number=True,
                )
            warnings.extend(
                apply_indonesia_ee_codes(
                    wb,
                    employees,
                    employee_directory=employee_directory,
                    pn_meta=applied_pn or pn_meta,
                )
            )

            fact_store_updates, plugin_cell_writes = _apply_vendor_plugins(
                wb, warnings, employee_count=len(employees)
            )

            apply_luckysheet_compat(wb, pn_sheet=PN_SHEET)
            wb.save(output_path)
        finally:
            wb.close()

        postprocess_converted_xlsx(output_path)
        warnings = merge_warnings(
            warnings,
            sanity_check_convert_result({"employee_count": len(employees), "warnings": warnings}),
        )
        return {
            "ok": True,
            "engine_id": "indonesia_payroll_calc",
            "region": "Indonesia",
            "output": str(output_path),
            "employee_count": len(employees),
            "employee_names": [
                _norm(e.get("Employee Name") or e.get("Name of Employee")) for e in employees
            ],
            "fx_rate": fx,
            "fx_source": fx_source,
            "warnings": warnings,
            "pn_meta": applied_pn.to_dict() if applied_pn else None,
            "pn_layout": pn_layout,
            "fact_store_updates": fact_store_updates,
            "cell_writes": plugin_cell_writes,
        }
    finally:
        _ACTIVE_MAPPING = None


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Link Compliance / Indonesia-L → Indonesia PN")
    parser.add_argument("source", type=Path)
    parser.add_argument("-o", "--output", type=Path, default=None)
    parser.add_argument("-t", "--template", type=Path, default=None)
    parser.add_argument("--no-fx", action="store_true")
    parser.add_argument(
        "--cash-advance-pdf",
        type=Path,
        default=None,
        help="Link Compliance Tax Invoice PDF（仅抽取 Cash Advance 写入 PN）",
    )
    args = parser.parse_args(argv)
    source = args.source.resolve()
    output = (args.output or source.with_name(f"PN_Indonesia_{source.stem}.xlsx")).resolve()
    template = (args.template or DEFAULT_TEMPLATE).resolve()
    result = convert(
        source,
        output,
        template,
        fill_fx=not args.no_fx,
        cash_advance_pdf=args.cash_advance_pdf.resolve() if args.cash_advance_pdf else None,
    )
    print(result)
    return 0 if result.get("ok") else 1


if __name__ == "__main__":
    sys.exit(main())
